"""PDF 안의 표를 엑셀로 옮깁니다. 여러 쪽에 이어진 표는 하나로 잇고 반복 머리글은 뺍니다.

인자(JSON 하나): {"file": "입력.pdf", "pages": "1-3,5"(생략 시 전부), "output": "출력 파일 이름(생략 가능)",
"merge": true}

결과: outputs/<이름>_표.xlsx (표마다 시트 하나 + 요약 시트). 마지막 줄에 요약 JSON을 출력합니다.
pdfplumber가 있으면 선으로 그린 표를 읽고, 없거나 표를 못 찾으면 `pdftotext -layout`(없으면 pypdf)
글자 배치를 열로 나눕니다. 스캔본(글자가 그림인 PDF)은 읽지 못합니다.
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - reported to the caller
    openpyxl = None

NUMBER = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$")
FOOTER = re.compile(
    r"^(?:-\s{0,3}\d{1,4}\s{0,3}-|\d{1,4}\s{0,3}/\s{0,3}\d{1,4}|page\s{0,3}\d{1,4}"
    r"(?:\s{0,3}of\s{0,3}\d{1,4})?|\d{1,4}\s{0,3}쪽|페이지\s{0,3}\d{1,4})$",
    re.I,
)
GAP = re.compile(r"[ \t]{2,200}")


def page_numbers(spec: str | None, count: int) -> list[int]:
    if not spec:
        return list(range(1, count + 1))
    pages: list[int] = []
    for part in str(spec).split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, _, end = part.partition("-")
            first = int(start or 1)
            last = int(end or count)
            pages.extend(range(first, min(last, count) + 1))
        else:
            pages.append(int(part))
    return [p for p in dict.fromkeys(pages) if 1 <= p <= count]


def cell_value(raw):
    if raw is None:
        return None
    text = re.sub(r"\s{1,8}", " ", str(raw)).strip()
    if not text:
        return None
    compact = text.replace(" ", "")
    if NUMBER.match(compact):
        digits = compact.lstrip("+-")
        if len(digits) > 1 and digits[0] == "0" and digits[1] != "." and "," not in digits:
            return text  # a code with a leading zero
        number = float(compact.replace(",", ""))
        return int(number) if number.is_integer() and "." not in compact else number
    return text


def looks_like_data(row: list) -> bool:
    return any(isinstance(cell_value(c), (int, float)) for c in row)


# ── extraction ──────────────────────────────────────────────────────────


def with_pdfplumber(path: Path, pages: list[int] | None) -> tuple[list[dict], int] | None:
    try:
        import pdfplumber
    except ImportError:
        return None
    tables: list[dict] = []
    with pdfplumber.open(str(path)) as pdf:
        count = len(pdf.pages)
        for number in pages or range(1, count + 1):
            if number > count:
                continue
            page = pdf.pages[number - 1]
            for table in page.extract_tables():
                rows = [[c for c in row] for row in table if any(c not in (None, "") for c in row)]
                if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
                    tables.append({"pages": [number], "rows": rows})
    return tables, count


def page_texts(path: Path, pages: list[int] | None) -> tuple[list[tuple[int, str]], int, str]:
    """Each page's text with its layout kept, and the tool that read it."""
    count = 0
    try:
        from pypdf import PdfReader
    except ImportError:
        PdfReader = None
    if PdfReader is not None:
        count = len(PdfReader(str(path)).pages)
    pdftotext = shutil.which("pdftotext")
    if pdftotext:
        if not count:
            info = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True)
            match = re.search(r"Pages:\s{1,20}(\d+)", info.stdout or "")
            count = int(match.group(1)) if match else 0
        texts = []
        for number in pages or range(1, count + 1):
            out = subprocess.run(
                [pdftotext, "-layout", "-f", str(number), "-l", str(number), str(path), "-"],
                capture_output=True, text=True, timeout=20,
            )
            texts.append((number, out.stdout))
        return texts, count, "pdftotext"
    if PdfReader is not None:
        reader = PdfReader(str(path))
        texts = []
        for number in pages or range(1, count + 1):
            page = reader.pages[number - 1]
            try:
                body = page.extract_text(extraction_mode="layout")
            except TypeError:
                body = page.extract_text()
            texts.append((number, body or ""))
        return texts, count, "pypdf"
    try:
        from PyPDF2 import PdfReader as LegacyReader
    except ImportError as exc:
        raise SystemExit(fail("PDF를 읽을 도구(pdfplumber, pdftotext, pypdf)가 없습니다.")) from exc
    reader = LegacyReader(str(path))
    count = len(reader.pages)
    texts = [(n, reader.pages[n - 1].extract_text() or "") for n in pages or range(1, count + 1)]
    return texts, count, "PyPDF2"


def split_line(line: str, width: int | None) -> list[str]:
    cells = [c for c in GAP.split(line.strip()) if c]
    if width and len(cells) < width:
        words = line.split()
        if len(words) == width:
            return words
    return cells


def from_layout(texts: list[tuple[int, str]]) -> list[dict]:
    """Runs of lines that split into the same number of columns are a table."""
    tables: list[dict] = []
    for number, body in texts:
        lines = [ln for ln in body.splitlines() if ln.strip() and not FOOTER.match(ln.strip())]
        current: dict | None = None
        for line in lines:
            width = len(current["rows"][0]) if current else None
            cells = split_line(line, width)
            if current and len(cells) == width:
                current["rows"].append(cells)
                continue
            if current and len(current["rows"]) >= 2:
                tables.append(current)
            current = {"pages": [number], "rows": [cells]} if len(cells) >= 2 else None
        if current and len(current["rows"]) >= 2:
            tables.append(current)
    return tables


# ── tidying ─────────────────────────────────────────────────────────────


def tidy(tables: list[dict], merge: bool) -> tuple[list[dict], int]:
    merged: list[dict] = []
    repeated = 0
    for table in tables:
        rows = [[(str(c).strip() if c is not None else "") for c in row] for row in table["rows"]]
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        header = rows[0] if not looks_like_data(rows[0]) else None
        body = rows[1:] if header else rows
        previous = merged[-1] if merged else None
        continues = (
            merge
            and previous is not None
            and len(previous["header"]) == width
            and previous["pages"][-1] + 1 == table["pages"][0]
            and (header == previous["header"] or header is None)
        )
        if continues:
            if header == previous["header"]:
                repeated += 1
            previous["body"].extend(body)
            previous["pages"].append(table["pages"][0])
            continue
        merged.append(
            {
                "pages": list(table["pages"]),
                "header": header or [f"열{i + 1}" for i in range(width)],
                "body": body,
            }
        )
    for table in merged:
        kept = []
        for row in table["body"]:
            if row == table["header"]:
                repeated += 1
                continue
            kept.append(row)
        table["body"] = kept
    return merged, repeated


def write_book(tables: list[dict], target: Path, method: str) -> None:
    book = openpyxl.Workbook()
    summary = book.active
    summary.title = "요약"
    bold = Font(bold=True)
    summary.append(["시트", "쪽", "행", "열", "머리글"])
    for cell in summary[1]:
        cell.font = bold
    for index, table in enumerate(tables, start=1):
        name = f"표{index}"
        sheet = book.create_sheet(name)
        sheet.append(table["header"])
        for cell in sheet[1]:
            cell.font = bold
        for row in table["body"]:
            sheet.append([cell_value(c) for c in row])
        for col, title in enumerate(table["header"], start=1):
            values = [str(r[col - 1]) for r in table["body"][:200]] + [title]
            longest = max(len(v) for v in values)
            sheet.column_dimensions[get_column_letter(col)].width = min(60, max(8, longest * 1.6 + 2))
            if any(isinstance(cell_value(r[col - 1]), int) for r in table["body"]):
                for cells in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
                    if isinstance(cells[0].value, (int, float)):
                        cells[0].number_format = "#,##0"
        sheet.freeze_panes = "A2"
        pages = table["pages"]
        span = f"{pages[0]}" if len(pages) == 1 else f"{pages[0]}-{pages[-1]}"
        summary.append([name, span, len(table["body"]), len(table["header"]), ", ".join(table["header"])])
    summary.append([])
    summary.append(["읽은 방법", method])
    summary.column_dimensions["E"].width = 60
    book.save(target)


# ── entry ───────────────────────────────────────────────────────────────


def fail(message: str) -> int:
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
    return 1


def pick_input(name: str | None) -> Path:
    if name:
        path = Path(name)
        if not path.exists():
            raise SystemExit(fail(f"입력 파일 {name} 을(를) 찾지 못했습니다."))
        return path
    found = [p for p in Path(".").iterdir() if p.is_file() and p.suffix.lower() == ".pdf"]
    if len(found) != 1:
        names = ", ".join(sorted(p.name for p in found)) or "없음"
        raise SystemExit(fail(f"file 인자로 PDF를 지정하세요. 후보: {names}"))
    return found[0]


def output_path(source: Path, wanted: str | None) -> Path:
    name = Path(wanted).name if wanted else f"{source.stem}_표.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    if name.lower().startswith("code"):
        name = "표_" + name  # the sandbox drops outputs whose name starts with "code"
    out = Path("outputs")
    out.mkdir(exist_ok=True)
    return out / name


def parse_args(argv: list[str]) -> dict:
    if len(argv) < 2:
        return {}
    raw = argv[1].strip()
    return json.loads(raw) if raw.startswith("{") else {"file": raw}


def main(argv: list[str]) -> int:
    if openpyxl is None:
        return fail("openpyxl 이 없어 실행할 수 없습니다.")
    options = parse_args(argv)
    source = pick_input(options.get("file"))
    spec = options.get("pages")
    method = "pdfplumber"
    tables: list[dict] = []
    count = 0
    plumbed = with_pdfplumber(source, None)
    if plumbed is not None:
        tables, count = plumbed
        wanted = set(page_numbers(spec, count))
        tables = [t for t in tables if t["pages"][0] in wanted]
    if not tables:
        texts, count, method = page_texts(source, None)
        wanted = page_numbers(spec, count)
        texts = [(n, body) for n, body in texts if n in wanted]
        tables = from_layout(texts)
        if not any(body.strip() for _, body in texts):
            return fail("PDF에서 글자를 읽지 못했습니다. 스캔본이면 글자 인식(OCR)이 먼저 필요합니다.")
    merged, repeated = tidy(tables, options.get("merge", True))
    if not merged:
        return fail("표를 찾지 못했습니다. 표가 그림이거나 열 구분이 없는 PDF일 수 있습니다.")
    target = output_path(source, options.get("output"))
    write_book(merged, target, method)
    summary = {
        "ok": True,
        "input": source.name,
        "pages": count,
        "method": method,
        "repeated_headers_removed": repeated,
        "outputs": [str(target)],
        "tables": [
            {
                "sheet": f"표{i}",
                "pages": t["pages"],
                "columns": t["header"],
                "rows": len(t["body"]),
                "first_row": t["body"][0] if t["body"] else [],
            }
            for i, t in enumerate(merged, start=1)
        ],
    }
    for item in summary["tables"]:
        print(f"[{item['sheet']}] 쪽 {item['pages']}, {item['rows']}행, 열: {', '.join(item['columns'])}")
    print(f"저장: {target} ({method})")
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
