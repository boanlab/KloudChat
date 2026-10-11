"""어지러운 엑셀·CSV 시트를 분석용 표로 정리합니다."""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - reported to the caller
    openpyxl = None

SCAN_ROWS = 20
MAX_LOG_ROWS = 300

TOTAL_WORDS = {
    "합계", "총계", "소계", "계", "총합", "합", "누계", "총액", "전체",
    "total", "subtotal", "sum", "grandtotal",
}
NOTE_PREFIXES = ("※", "*", "주)", "주:", "주 :", "(주)", "비고:", "참고:", "출처:", "note", "(단위", "단위:")
MONEY = {"원": 1, "천원": 1_000, "만원": 10_000, "백만원": 1_000_000, "억원": 100_000_000, "조원": 10**12}
UNIT_SUFFIX = re.compile(
    r"(천원|백만원|만원|억원|조원|원|개|명|건|회|EA|ea|kg|KG|km|톤|%)$"
)
NUMBER_CORE = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$")
DATE_YMD = re.compile(r"^(\d{4})\s{0,2}[-./]\s{0,2}(\d{1,2})\s{0,2}[-./]\s{0,2}(\d{1,2})\.?$")
DATE_COMPACT = re.compile(r"^(\d{4})(\d{2})(\d{2})$")
DATE_KOREAN = re.compile(r"^(\d{4})\s{0,2}년\s{0,2}(\d{1,2})\s{0,2}월\s{0,2}(\d{1,2})\s{0,2}일?$")
DATE_SHORT = re.compile(r"^(\d{2})[./-](\d{1,2})[./-](\d{1,2})$")
DATE_HEADER = re.compile(r"날짜|일자|일시|date|기준일|작성일|거래일|등록일|발행일|납기", re.I)
SKIP_TOTAL_HEADER = re.compile(r"번호|코드|연도|년도|^no\.?$|^id$|순번|우편", re.I)
EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d") if value.time() == dt.time() else value.isoformat(" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_number(value):
    """`(number, unit, is_percent)` or None. Leading-zero codes stay text."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value, "", False
    if not isinstance(value, str):
        return None
    raw = value.strip().replace(" ", "")
    if not raw or len(raw) > 40:
        return None
    negative = False
    if raw.startswith("(") and raw.endswith(")"):
        negative, raw = True, raw[1:-1]
    if raw[:1] in ("△", "▼", "−"):
        negative, raw = True, raw[1:]
    elif raw[:1] == "▲":
        raw = raw[1:]
    raw = raw.lstrip("₩￦$")
    unit = ""
    match = UNIT_SUFFIX.search(raw)
    if match and match.start() > 0:
        unit = match.group(1)
        raw = raw[: match.start()]
    if not NUMBER_CORE.match(raw):
        return None
    digits = raw.lstrip("+-")
    if len(digits) > 1 and digits[0] == "0" and digits[1] != "." and "," not in digits:
        return None  # 007, 0101: a code, not a number
    number = float(raw.replace(",", ""))
    if negative:
        number = -number
    if number.is_integer() and "." not in raw:
        number = int(number)
    return number, unit, unit == "%"


def _date(year: int, month: int, day: int):
    try:
        if 1900 <= year <= 2200:
            return dt.datetime(year, month, day)
    except ValueError:
        return None
    return None


def parse_date(value, *, allow_serial: bool = False):
    if isinstance(value, dt.datetime):
        return value
    if isinstance(value, dt.date):
        return dt.datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if allow_serial and 20_000 <= value <= 80_000:
            return EXCEL_EPOCH + dt.timedelta(days=float(value))
        if isinstance(value, int) or float(value).is_integer():
            match = DATE_COMPACT.match(str(int(value)))
            if match:
                return _date(*(int(g) for g in match.groups()))
        return None
    if not isinstance(value, str):
        return None
    raw = value.strip()
    if len(raw) > 30:
        return None
    for pattern in (DATE_YMD, DATE_KOREAN, DATE_COMPACT):
        match = pattern.match(raw)
        if match:
            return _date(*(int(g) for g in match.groups()))
    match = DATE_SHORT.match(raw)
    if match:
        year, month, day = (int(g) for g in match.groups())
        return _date(2000 + year, month, day)
    if allow_serial and re.fullmatch(r"\d{5}(?:\.\d{1,6})?", raw):
        return parse_date(float(raw), allow_serial=True)
    return None


def is_texty(value) -> bool:
    return (
        isinstance(value, str)
        and bool(value.strip())
        and parse_number(value) is None
        and parse_date(value) is None
    )


def read_csv(path: Path) -> tuple[dict[str, list[list]], dict]:
    data = path.read_bytes()
    encoding = "utf-8"
    for candidate in ("utf-8-sig", "utf-8", "cp949", "euc-kr", "latin-1"):
        try:
            decoded = data.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    sample = decoded[:20_000]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        delimiter = ","
    rows = [
        [cell if cell != "" else None for cell in row]
        for row in csv.reader(io.StringIO(decoded), delimiter=delimiter)
    ]
    return {sheet_title(path.stem): rows}, {"encoding": encoding, "delimiter": delimiter}


def read_xlsx(path: Path, only: str | None) -> tuple[dict[str, list[list]], dict[str, int]]:
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets: dict[str, list[list]] = {}
    for sheet in book.worksheets:
        if only and sheet.title != only:
            continue
        sheets[sheet.title] = [list(row) for row in sheet.iter_rows(values_only=True)]
    book.close()
    # Formulas without a cached value read as empty; count them so the report says so.
    missing: dict[str, int] = {}
    try:
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
        for sheet in formulas.worksheets:
            if sheet.title not in sheets:
                continue
            values = sheets[sheet.title]
            count = 0
            for r, row in enumerate(sheet.iter_rows(values_only=True)):
                for c, cell in enumerate(row):
                    if isinstance(cell, str) and cell.startswith("=") and (
                        r >= len(values) or c >= len(values[r]) or values[r][c] is None
                    ):
                        count += 1
            if count:
                missing[sheet.title] = count
        formulas.close()
    except Exception:  # noqa: BLE001 - only a report detail
        pass
    return sheets, missing


def sheet_title(name: str) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", name).strip() or "Sheet1"
    return cleaned[:31]


def trim(rows: list[list]) -> list[list]:
    width = 0
    for row in rows:
        for index in range(len(row) - 1, -1, -1):
            if not blank(row[index]):
                width = max(width, index + 1)
                break
    return [list(row[:width]) + [None] * (width - len(row[:width])) for row in rows]


def find_header(rows: list[list]) -> int:
    counts = [sum(not blank(v) for v in row) for row in rows[:SCAN_ROWS]]
    widest = max(counts or [0])
    if not widest:
        return 0
    for index, row in enumerate(rows[:SCAN_ROWS]):
        filled = counts[index]
        texty = sum(is_texty(v) for v in row)
        following = any(
            sum(not blank(v) for v in later) >= 2 for later in rows[index + 1 : index + 4]
        )
        if texty >= min(2, widest) and filled >= 0.6 * widest and texty >= 0.6 * filled and following:
            return index
    return next((i for i, n in enumerate(counts) if n == widest), 0)


def first_text(row: list) -> str:
    for value in row[:3]:
        if not blank(value):
            return text(value)
    return ""


def row_kind(row: list, header: list, width: int) -> str | None:
    filled = [v for v in row if not blank(v)]
    if not filled:
        return "blank"
    if [text(v) for v in row] == [text(v) for v in header]:
        return "header"
    lead = first_text(row)
    squeezed = re.sub(r"\s{1,8}", "", lead).lower()
    if squeezed in TOTAL_WORDS or squeezed.startswith(("합계", "소계", "총계", "총합계")):
        return "total"
    if len(filled) <= 2 and lead.lower().startswith(NOTE_PREFIXES):
        return "note"
    if len(filled) == 1 and width >= 3 and isinstance(filled[0], str) and len(filled[0]) >= 15:
        return "note"
    return None


def header_names(row: list) -> list[str]:
    names: list[str] = []
    for index, value in enumerate(row):
        name = re.sub(r"\s{1,8}", " ", text(value)).strip() or f"열{index + 1}"
        base, n = name, 2
        while name in names:
            name, n = f"{base}_{n}", n + 1
        names.append(name)
    return names


def convert_columns(names: list[str], body: list[list], log: list, sheet: str) -> dict:
    info = {"numeric": [], "date": [], "percent": [], "units": {}, "left_as_text": {}}
    for col, name in enumerate(names):
        cells = [(r, row[col]) for r, row in enumerate(body) if not blank(row[col])]
        if not cells:
            continue
        dated = [parse_date(v) for _, v in cells]
        date_like_header = bool(DATE_HEADER.search(name)) or name.endswith("일")
        plain = sum(
            d is not None and not isinstance(v, (int, float))
            for d, (_, v) in zip(dated, cells, strict=True)
        )
        native = sum(isinstance(v, (dt.date, dt.datetime)) for _, v in cells)
        if date_like_header or (plain + native) >= 0.5 * len(cells):
            serial = [parse_date(v, allow_serial=True) for _, v in cells]
            hits = sum(d is not None for d in serial)
            if hits >= 0.8 * len(cells) and (date_like_header or plain + native):
                changed = 0
                for (r, value), parsed in zip(cells, serial, strict=True):
                    if parsed is not None:
                        if not isinstance(value, dt.datetime) or value != parsed:
                            changed += 1
                        body[r][col] = parsed
                info["date"].append(name)
                if changed:
                    log.append((sheet, "날짜 변환", "", f"{name}: {changed}칸"))
                misses = [text(v) for (_, v), d in zip(cells, serial, strict=True) if d is None]
                if misses:
                    info["left_as_text"][name] = misses[:5]
                continue
        parsed = [parse_number(v) for _, v in cells]
        hits = [p for p in parsed if p is not None]
        if len(hits) < 0.8 * len(cells):
            continue
        units = {p[1] for p in hits if p[1]}
        money = units and units <= set(MONEY)
        changed = 0
        for (r, value), p in zip(cells, parsed, strict=True):
            if p is None:
                continue
            number, unit, percent = p
            if percent:
                number = number / 100
            elif money and len(units) > 1:
                number = number * MONEY.get(unit or "원", 1)
            if isinstance(number, float) and number.is_integer() and not percent:
                number = int(number)
            if number != value:
                changed += 1
            body[r][col] = number
        new_name = name
        if units == {"%"}:
            info["percent"].append(name)
        elif money and len(units) > 1:
            new_name = f"{name}(원)"
            log.append((sheet, "단위 통일", "", f"{name}: {', '.join(sorted(units))} → 원"))
        elif len(units) == 1 and "%" not in units:
            unit = next(iter(units))
            if unit not in name:
                new_name = f"{name}({unit})"
        if new_name != name:
            info["units"][new_name] = name
            names[col] = new_name
        info["numeric"].append(new_name)
        if changed:
            log.append((sheet, "숫자 변환", "", f"{new_name}: {changed}칸"))
        misses = [text(v) for (_, v), p in zip(cells, parsed, strict=True) if p is None]
        if misses:
            info["left_as_text"][new_name] = misses[:5]
            log.append((sheet, "숫자로 못 바꾼 칸", "", f"{new_name}: {', '.join(misses[:5])}"))
    return info


def clean_sheet(sheet: str, rows: list[list], options: dict, log: list) -> dict:
    rows = trim(rows)
    if not rows:
        return {"name": sheet, "empty": True}
    width = len(rows[0])
    if options.get("header_row"):
        header_at = max(0, int(options["header_row"]) - 1)
    else:
        header_at = find_header(rows)
    header = rows[header_at]
    names = header_names(header)
    removed = {"title": 0, "blank": 0, "total": 0, "note": 0, "header": 0, "duplicate": 0}
    titles = [first_text(row) for row in rows[:header_at] if not blank(first_text(row))]
    for index, row in enumerate(rows[:header_at]):
        if any(not blank(v) for v in row):
            removed["title"] += 1
            log.append((sheet, "삭제: 제목 행", index + 1, first_text(row)[:80]))
    log.append((sheet, "머리글 행", header_at + 1, ", ".join(names)))
    body: list[list] = []
    origin: list[int] = []
    for index in range(header_at + 1, len(rows)):
        row = rows[index]
        kind = row_kind(row, header, width)
        if kind == "total" and not options.get("drop_totals", True):
            kind = None
        if kind:
            removed[kind] += 1
            if kind != "blank":
                label = {"total": "합계 행", "note": "주석 행", "header": "반복 머리글"}[kind]
                cells = ", ".join(text(v) for v in row if not blank(v))
                log.append((sheet, f"삭제: {label}", index + 1, cells[:120]))
            continue
        body.append(list(row))
        origin.append(index + 1)
    info = convert_columns(names, body, log, sheet)
    # Identical rows only marked unless `dedupe`: two records may share every shown value.
    dedupe = bool(options.get("dedupe", False))
    suspected = 0
    seen: set = set()
    kept, kept_origin = [], []
    for row, at in zip(body, origin, strict=True):
        key = tuple(text(v) for v in row)
        if key in seen:
            if dedupe:
                removed["duplicate"] += 1
                log.append((sheet, "삭제: 중복 행", at, ", ".join(key)[:120]))
                continue
            suspected += 1
            log.append((sheet, "표시: 중복 의심 행(남김)", at, ", ".join(key)[:120]))
        seen.add(key)
        kept.append(row)
        kept_origin.append(at)
    body = kept
    unit_pattern = re.compile(r"단위\s{0,3}[:：]\s{0,3}[^,)]{1,10}")
    unit_note = next((m.group(0) for t in titles if (m := unit_pattern.search(t))), "")
    return {
        "name": sheet,
        "header_row": header_at + 1,
        "columns": names,
        "rows_in": len(rows) - header_at - 1,
        "rows_out": len(body),
        "removed": {k: v for k, v in removed.items() if v},
        **({"suspected_duplicates": suspected} if suspected else {}),
        "numeric_columns": info["numeric"],
        "date_columns": info["date"],
        "percent_columns": info["percent"],
        "renamed": info["units"],
        "left_as_text": info["left_as_text"],
        **({"unit_note": unit_note} if unit_note else {}),
        "_body": body,
    }


def write_book(results: list[dict], log: list, path: Path, add_totals: bool) -> dict:
    book = openpyxl.Workbook()
    book.remove(book.active)
    bold = Font(bold=True)
    totals: dict[str, dict] = {}
    for result in results:
        if result.get("empty"):
            continue
        sheet = book.create_sheet(sheet_title(result["name"]))
        names = result["columns"]
        sheet.append(names)
        for cell in sheet[1]:
            cell.font = bold
        body = result["_body"]
        for row in body:
            sheet.append(row)
        last = len(body) + 1
        for col, name in enumerate(names, start=1):
            letter = get_column_letter(col)
            fmt = None
            if name in result["date_columns"]:
                fmt = "yyyy-mm-dd"
            elif name in result["percent_columns"]:
                fmt = "0.0%"
            elif name in result["numeric_columns"]:
                values = [r[col - 1] for r in body if isinstance(r[col - 1], (int, float))]
                fmt = "#,##0" if all(float(v).is_integer() for v in values) else "#,##0.##"
            if fmt:
                for row_cells in sheet.iter_rows(min_row=2, max_row=last, min_col=col, max_col=col):
                    row_cells[0].number_format = fmt
            longest = max([len(text(name))] + [len(text(r[col - 1])) for r in body[:200]])
            sheet.column_dimensions[letter].width = min(60, max(8, longest * 1.6 + 2))
        sheet.freeze_panes = "A2"
        if add_totals and body:
            row_at = last + 1
            sheet.cell(row=row_at, column=1, value="합계").font = bold
            expected: dict[str, float] = {}
            for col, name in enumerate(names, start=1):
                if col == 1 or name not in result["numeric_columns"]:
                    continue
                if name in result["percent_columns"] or SKIP_TOTAL_HEADER.search(name):
                    continue
                letter = get_column_letter(col)
                cell = sheet.cell(row=row_at, column=col, value=f"=SUM({letter}2:{letter}{last})")
                cell.font = bold
                cell.number_format = "#,##0"
                expected[f"{letter}{row_at}"] = sum(
                    r[col - 1] for r in body if isinstance(r[col - 1], (int, float))
                )
            totals[sheet.title] = expected
    changes = book.create_sheet("변경내역")
    changes.append(["시트", "구분", "원래 행", "내용"])
    for cell in changes[1]:
        cell.font = bold
    for entry in log[:MAX_LOG_ROWS]:
        changes.append(list(entry))
    if len(log) > MAX_LOG_ROWS:
        changes.append(["", "생략", "", f"변경 {len(log) - MAX_LOG_ROWS}건 더"])
    for letter, width in zip("ABCD", (16, 18, 9, 80), strict=True):
        changes.column_dimensions[letter].width = width
    book.save(path)
    return totals


def recalc(path: Path) -> tuple[bool, str]:
    """Formula values cached by LibreOffice, so the file opens with numbers everywhere."""
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        return False, "LibreOffice가 없어 수식 계산값을 저장하지 못했습니다."
    work = Path(tempfile.mkdtemp(prefix="recalc"))
    try:
        result = subprocess.run(
            [
                office, "--headless", "--norestore",
                f"-env:UserInstallation=file://{work}/profile",
                "--convert-to", "xlsx", "--outdir", str(work / "out"), str(path.resolve()),
            ],
            capture_output=True, text=True, timeout=25,
        )
        produced = work / "out" / path.name
        if not produced.exists():
            return False, (result.stderr or result.stdout or "변환 실패").strip()[:200]
        shutil.copyfile(produced, path)
        return True, ""
    except (OSError, subprocess.SubprocessError) as exc:
        return False, str(exc)[:200]
    finally:
        shutil.rmtree(work, ignore_errors=True)


def check_totals(path: Path, totals: dict[str, dict]) -> dict:
    if not totals:
        return {}
    book = openpyxl.load_workbook(path, data_only=True)
    report = {"checked": 0, "matched": 0, "missing": 0, "mismatched": []}
    for sheet_name, cells in totals.items():
        if sheet_name not in book.sheetnames:
            continue
        sheet = book[sheet_name]
        for ref, want in cells.items():
            report["checked"] += 1
            got = sheet[ref].value
            if got is None or (isinstance(got, str) and got.startswith("=")):
                report["missing"] += 1
            elif abs(float(got) - float(want)) <= 1e-6 * max(1.0, abs(float(want))):
                report["matched"] += 1
            else:
                report["mismatched"].append(
                    {"sheet": sheet_name, "cell": ref, "value": got, "expected": want}
                )
    return report


def pick_input(name: str | None) -> Path:
    if name:
        path = Path(name)
        if not path.exists():
            raise SystemExit(fail(f"입력 파일 {name} 을(를) 찾지 못했습니다."))
        return path
    found = [
        p for p in Path(".").iterdir()
        if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm", ".csv")
    ]
    if len(found) != 1:
        names = ", ".join(sorted(p.name for p in found)) or "없음"
        raise SystemExit(fail(f"file 인자로 입력 파일을 지정하세요. 후보: {names}"))
    return found[0]


def output_path(source: Path, wanted: str | None) -> Path:
    name = Path(wanted).name if wanted else f"{source.stem}_정리.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    if name.lower().startswith("code"):
        name = "정리_" + name  # the sandbox drops outputs whose name starts with "code"
    out = Path("outputs")
    out.mkdir(exist_ok=True)
    return out / name


def fail(message: str) -> int:
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
    return 1


def parse_args(argv: list[str]) -> dict:
    if len(argv) < 2:
        return {}
    raw = argv[1].strip()
    if raw.startswith("{"):
        return json.loads(raw)
    return {"file": raw}


def main(argv: list[str]) -> int:
    if openpyxl is None:
        return fail("openpyxl 이 없어 실행할 수 없습니다.")
    options = parse_args(argv)
    source = pick_input(options.get("file"))
    log: list[tuple] = []
    missing_formulas: dict[str, int] = {}
    meta: dict = {}
    suffix = source.suffix.lower()
    if suffix == ".csv":
        sheets, meta = read_csv(source)
    elif suffix in (".xlsx", ".xlsm"):
        sheets, missing_formulas = read_xlsx(source, options.get("sheet"))
    else:
        return fail(f"{source.suffix} 형식은 지원하지 않습니다. xlsx 또는 csv로 저장해 주세요.")
    if not sheets:
        return fail(f"시트 {options.get('sheet')!r} 을(를) 찾지 못했습니다.")
    results = [clean_sheet(name, rows, options, log) for name, rows in sheets.items()]
    target = output_path(source, options.get("output"))
    add_totals = bool(options.get("add_totals"))
    totals = write_book(results, log, target, add_totals)
    recalculated, recalc_note = False, ""
    if add_totals and totals and options.get("recalc", True):
        recalculated, recalc_note = recalc(target)
    total_check = check_totals(target, totals) if recalculated else {}
    for result in results:
        result.pop("_body", None)
    summary = {
        "ok": True,
        "input": source.name,
        **({"csv": meta} if meta else {}),
        "outputs": [str(target)],
        "sheets": results,
        **({"formulas_without_values": missing_formulas} if missing_formulas else {}),
        **(
            {
                "totals_added": sum(len(v) for v in totals.values()),
                "recalculated": recalculated,
                **({"recalc_note": recalc_note} if recalc_note else {}),
                **({"totals_check": total_check} if total_check else {}),
            }
            if add_totals
            else {}
        ),
    }
    for result in results:
        if result.get("empty"):
            print(f"[{result['name']}] 빈 시트")
            continue
        removed = ", ".join(f"{k} {v}" for k, v in result["removed"].items()) or "없음"
        print(
            f"[{result['name']}] 머리글 {result['header_row']}행, 데이터 {result['rows_in']}행 → "
            f"{result['rows_out']}행 (삭제: {removed})"
        )
    print(f"저장: {target}")
    print(json.dumps(summary, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
