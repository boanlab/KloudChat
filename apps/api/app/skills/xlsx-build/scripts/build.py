"""표(직접 준 행 또는 첨부 CSV·엑셀)로 수식·합계 행·차트가 든 엑셀 파일을 만듭니다."""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - reported to the caller
    openpyxl = None

MAX_ROWS = 50_000
PREVIEW_ROWS = 30
NUMBER = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$")
PERCENT = re.compile(r"^[+-]?\d+(?:\.\d+)?%$")
DATE = re.compile(r"^(\d{4})[-./](\d{1,2})[-./](\d{1,2})\.?$")
PLACEHOLDER = re.compile(r"\{([^{}]{1,60})\}")
#: headers whose numbers are labels, never summed
LABEL_HEADER = re.compile(r"번호|코드|연도|년도|^no\.?$|^id$|순번|우편|전화", re.I)
HEAD_FILL = "EEF2F7"
SUMMED = ("#,##0", "#,##0.##")


def fail(message: str) -> int:
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
    return 1


def cell_value(value):
    """Typed cell: `1,234` → 1234, `12.5%` → 0.125, `2026-10-13` → date; `007` stays text."""
    if value is None or isinstance(value, (bool, int, float, dt.date)):
        return value
    text = str(value).strip()
    if text in ("", "-", "—"):
        return None
    if text.startswith("="):
        return text
    raw = text.replace(" ", "")
    if PERCENT.match(raw):
        return float(raw[:-1]) / 100
    if NUMBER.match(raw) and not re.match(r"^[+-]?0\d", raw):
        number = float(raw.replace(",", ""))
        return int(number) if "." not in raw else number
    match = DATE.match(raw)
    if match:
        try:
            return dt.date(int(match[1]), int(match[2]), int(match[3]))
        except ValueError:
            return text
    return text


def is_percent(value) -> bool:
    return isinstance(value, str) and bool(PERCENT.match(value.strip().replace(" ", "")))


def sheet_title(name: str) -> str:
    return (re.sub(r"[\[\]:*?/\\]", "_", str(name)).strip() or "Sheet1")[:31]


def table_of(name: str, rows: list[list]) -> dict:
    """First non-empty row is the header; fully empty rows are dropped."""
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return {"name": name, "columns": [], "rows": []}
    header = [str(c).strip() if c not in (None, "") else f"열{i + 1}"
              for i, c in enumerate(rows[0])]
    return {"name": name, "columns": header, "rows": rows[1:]}


def read_csv(path: Path) -> list[dict]:
    data = path.read_bytes()
    decoded = data.decode("latin-1")
    for encoding in ("utf-8-sig", "cp949"):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        delimiter = csv.Sniffer().sniff(decoded[:20_000], delimiters=",\t;|").delimiter
    except csv.Error:
        delimiter = ","
    return [table_of(path.stem, list(csv.reader(io.StringIO(decoded), delimiter=delimiter)))]


def read_xlsx(path: Path, only: str | None) -> list[dict]:
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = [table_of(sheet.title, [list(r) for r in sheet.iter_rows(values_only=True)])
           for sheet in book.worksheets if not only or sheet.title == only]
    book.close()
    return out


def normalised_rows(columns: list[str], rows: list) -> list[list]:
    out = []
    for row in rows[:MAX_ROWS]:
        if isinstance(row, dict):
            row = [row.get(name) for name in columns]
        row = list(row)[: len(columns)]
        out.append(row + [None] * (len(columns) - len(row)))
    return out


def substitute(formula: str, letters: dict[str, str], row: int) -> str:
    """`{열이름}` → that column's cell on `row`; `{r}` → `row`."""

    def one(match: re.Match) -> str:
        key = match.group(1).strip()
        if key == "r":
            return str(row)
        if key not in letters:
            raise KeyError(key)
        return f"{letters[key]}{row}"

    text = formula.strip()
    return PLACEHOLDER.sub(one, text if text.startswith("=") else "=" + text)


def number_format(values: list) -> str | None:
    numbers = [v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if not numbers:
        return None
    return SUMMED[0] if all(float(v).is_integer() for v in numbers) else SUMMED[1]


def column_formats(columns, body, percent_cols, formulas, formats) -> dict[str, str | None]:
    """Given formats first; else by the values; a formula column of plain sums stays `#,##0`."""
    out: dict[str, str | None] = {}
    for i, name in enumerate(columns):
        values = [r[i] for r in body]
        if name in formats:
            out[name] = formats[name]
        elif name in percent_cols:
            out[name] = "0.0%"
        elif any(isinstance(v, dt.date) for v in values):
            out[name] = "yyyy-mm-dd"
        else:
            out[name] = number_format(values)
    for name, formula in formulas.items():
        if name in formats:
            out[name] = formats[name]
            continue
        used = [k.strip() for k in PLACEHOLDER.findall(formula) if k.strip() != "r"]
        plain = "/" not in formula and used and all(out.get(k) == SUMMED[0] for k in used)
        out[name] = SUMMED[0] if plain else None
    return out


def write_sheet(book, spec: dict, defaults: dict) -> dict:
    options = {**defaults, **{k: v for k, v in spec.items() if v is not None}}
    columns = [str(c).strip() for c in options.get("columns") or []]
    if not columns:
        raise ValueError(f"시트 {spec.get('name') or '?'}에 columns(머리글)가 없습니다.")
    raw = normalised_rows(columns, list(options.get("rows") or []))
    percent_cols = {name for i, name in enumerate(columns)
                    if any(is_percent(r[i]) for r in raw[:200])}
    body = [[cell_value(v) for v in row] for row in raw]
    formulas = {str(k).strip(): str(v) for k, v in (options.get("formulas") or {}).items()}
    all_columns = columns + [name for name in formulas if name not in columns]
    letters = {name: get_column_letter(i + 1) for i, name in enumerate(all_columns)}
    formats = {str(k): str(v) for k, v in (options.get("formats") or {}).items()}
    for name, formula in formulas.items():
        for key in PLACEHOLDER.findall(formula):
            if key.strip() != "r" and key.strip() not in letters:
                raise ValueError(f"수식 {name!r}의 {{{key}}} 열이 표에 없습니다.")

    sheet = book.create_sheet(sheet_title(options.get("name") or "표"))
    bold = Font(bold=True)
    head = 1
    title = str(options.get("title") or "").strip()
    if title:
        sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        head = 3
    for col, name in enumerate(all_columns, start=1):
        cell = sheet.cell(row=head, column=col, value=name)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
    first = head + 1
    for offset, row in enumerate(body):
        at = first + offset
        for col, value in enumerate(row, start=1):
            sheet.cell(row=at, column=col, value=value)
        for name, formula in formulas.items():
            sheet[f"{letters[name]}{at}"] = substitute(formula, letters, at)
    last = first + len(body) - 1
    formats_by_col = column_formats(columns, body, percent_cols, formulas, formats)

    # Total row: SUM per summed column; a formula column repeats its formula on that row,
    # so a rate there is total ÷ total, never a sum of rates.
    total = options.get("total_row")
    total_cells: dict[str, str] = {}
    total_at = None
    if total and body:
        if isinstance(total, list):
            summed = [str(n) for n in total if str(n) in columns and str(n) not in formulas]
        else:
            summed = [name for i, name in enumerate(columns)
                      if i > 0 and name not in formulas and name not in percent_cols
                      and not LABEL_HEADER.search(name) and formats_by_col.get(name) in SUMMED]
        total_at = last + 1
        label_col = next((i + 1 for i, n in enumerate(all_columns)
                          if n not in summed and n not in formulas), 1)
        sheet.cell(row=total_at, column=label_col, value=str(options.get("total_label") or "합계"))
        for name in summed:
            letter = letters[name]
            sheet[f"{letter}{total_at}"] = f"=SUM({letter}{first}:{letter}{last})"
            total_cells[name] = f"{letter}{total_at}"
        for name, formula in formulas.items():
            sheet[f"{letters[name]}{total_at}"] = substitute(formula, letters, total_at)
            total_cells[name] = f"{letters[name]}{total_at}"
        for col in range(1, len(all_columns) + 1):
            sheet.cell(row=total_at, column=col).font = bold

    end = total_at or last
    for i, name in enumerate(all_columns, start=1):
        fmt = formats_by_col.get(name)
        if fmt:
            for (cell,) in sheet.iter_rows(min_row=first, max_row=end, min_col=i, max_col=i):
                cell.number_format = fmt
        longest = max([len(name)] + [len(str(r[i - 1] or "")) for r in body[:200] if i <= len(r)])
        sheet.column_dimensions[get_column_letter(i)].width = min(50, max(8, longest * 1.6 + 2))
    sheet.freeze_panes = sheet.cell(row=first, column=1)
    if body:
        sheet.auto_filter.ref = f"A{head}:{get_column_letter(len(all_columns))}{last}"
    charted = bool(options.get("chart")) and bool(body) and add_chart(
        sheet, options["chart"], all_columns, head, first, last)
    return {
        "name": sheet.title,
        "columns": all_columns,
        "rows": len(body),
        "header_row": head,
        "formula_columns": list(formulas),
        "total_row": total_at,
        "total_cells": total_cells,
        "chart": charted,
    }


def add_chart(sheet, spec: dict, columns: list[str], head: int, first: int, last: int) -> bool:
    x = str(spec.get("x") or columns[0])
    ys = [str(y) for y in (spec.get("y") or []) if str(y) in columns and str(y) != x]
    if x not in columns or not ys:
        raise ValueError("chart 의 x·y 는 표의 열 이름이어야 합니다.")
    kind = str(spec.get("type") or "bar").lower()
    chart = LineChart() if kind == "line" else BarChart()
    if kind == "stacked":
        chart.grouping, chart.overlap = "stacked", 100
    chart.title = str(spec.get("title") or "") or None
    chart.y_axis.title = str(spec.get("y_label") or "") or None
    chart.x_axis.title = str(spec.get("x_label") or "") or None
    chart.height, chart.width = 8, 16
    for name in ys:
        col = columns.index(name) + 1
        chart.add_data(Reference(sheet, min_col=col, min_row=head, max_row=last),
                       titles_from_data=True)
    xcol = columns.index(x) + 1
    chart.set_categories(Reference(sheet, min_col=xcol, min_row=first, max_row=last))
    sheet.add_chart(chart, f"{get_column_letter(len(columns) + 2)}{head}")
    return True


def recalc(path: Path) -> tuple[dict, str]:
    """Formula values cached by LibreOffice (formulas, formats and chart kept), so every
    viewer and the summary show numbers; without it Excel computes on open."""
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        return {}, "LibreOffice가 없어 계산값을 저장하지 못했습니다. 엑셀에서 열면 계산됩니다."
    work = Path(tempfile.mkdtemp(prefix="recalc"))
    try:
        result = subprocess.run(
            [office, "--headless", "--norestore", f"-env:UserInstallation=file://{work}/profile",
             "--convert-to", "xlsx", "--outdir", str(work / "out"), str(path.resolve())],
            capture_output=True, text=True, timeout=40,
        )
        produced = work / "out" / path.name
        if not produced.exists():
            return {}, (result.stderr or result.stdout or "계산 실패").strip()[:200]
        shutil.copyfile(produced, path)
        book = openpyxl.load_workbook(path, data_only=True)
        values = {sheet.title: {c.coordinate: c.value for row in sheet.iter_rows() for c in row
                                if c.value is not None} for sheet in book.worksheets}
        book.close()
        return values, ""
    except (OSError, subprocess.SubprocessError) as exc:
        return {}, str(exc)[:200]
    finally:
        shutil.rmtree(work, ignore_errors=True)


def output_path(wanted: str | None, stem: str) -> Path:
    name = Path(str(wanted)).name if wanted else f"{stem}.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    if name.lower().startswith("code"):
        name = "표_" + name  # the sandbox drops outputs whose name starts with "code"
    out = Path("outputs")
    out.mkdir(exist_ok=True)
    return out / name


def parse_args(argv: list[str]) -> dict:
    if len(argv) < 2:
        return {}
    raw = argv[1].strip()
    return json.loads(raw) if raw.startswith("{") else {"file": raw}


def sheet_specs(options: dict) -> tuple[list[dict], str]:
    given = [s for s in options.get("sheets") or [] if isinstance(s, dict)]
    if not options.get("file"):
        return given, "표"
    source = Path(str(options["file"]))
    if not source.exists():
        raise ValueError(f"입력 파일 {source.name} 을(를) 찾지 못했습니다.")
    suffix = source.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        tables = read_csv(source)
    elif suffix in (".xlsx", ".xlsm"):
        tables = read_xlsx(source, options.get("sheet"))
    else:
        raise ValueError(f"{source.suffix} 형식은 지원하지 않습니다. xlsx 또는 csv로 주세요.")
    by_name = {str(s.get("name")): s for s in given}
    specs = []
    for table in tables:
        extra = by_name.get(table["name"]) or (given[0] if len(given) == 1 == len(tables) else {})
        specs.append({**{k: v for k, v in extra.items() if k not in ("columns", "rows")},
                      "columns": table["columns"], "rows": table["rows"],
                      "name": extra.get("name") or table["name"]})
    return specs, source.stem


def main(argv: list[str]) -> int:
    if openpyxl is None:
        return fail("openpyxl 이 없어 실행할 수 없습니다.")
    try:
        options = parse_args(argv)
        specs, stem = sheet_specs(options)
    except ValueError as exc:
        return fail(str(exc))
    if not specs:
        return fail("만들 표가 없습니다. sheets(columns·rows) 또는 file 을 주세요.")
    defaults = {k: options[k] for k in ("formulas", "total_row", "total_label", "formats", "chart")
                if k in options}
    book = openpyxl.Workbook()
    book.remove(book.active)
    try:
        sheets = [write_sheet(book, spec, defaults) for spec in specs]
    except ValueError as exc:
        return fail(str(exc))
    target = output_path(options.get("output"), stem)
    book.save(target)

    values, note = {}, ""
    if any(s["total_cells"] or s["formula_columns"] for s in sheets) and options.get("recalc", True):
        values, note = recalc(target)
    for sheet in sheets:
        cells = values.get(sheet["name"]) or {}
        sheet["totals"] = {name: cells.get(ref) for name, ref in sheet["total_cells"].items()}
        if cells and sheet["formula_columns"]:
            # Computed formula cells of the first rows, for the answer to quote.
            first = sheet["header_row"] + 1
            letters = {n: get_column_letter(i + 1) for i, n in enumerate(sheet["columns"])}
            sheet["formula_values"] = {
                name: [cells.get(f"{letters[name]}{r}") for r in range(first, first + min(
                    sheet["rows"], PREVIEW_ROWS))]
                for name in sheet["formula_columns"]}
    summary = {
        "ok": True,
        "outputs": [str(target)],
        "sheets": sheets,
        "recalculated": bool(values),
        **({"note": note} if note else {}),
    }
    for sheet in sheets:
        extra = [f"수식 열 {', '.join(sheet['formula_columns'])}"] if sheet["formula_columns"] else []
        extra += [f"합계 {sheet['total_row']}행"] if sheet["total_row"] else []
        extra += ["차트"] if sheet["chart"] else []
        print(f"[{sheet['name']}] {sheet['rows']}행 × {len(sheet['columns'])}열"
              + (f" ({', '.join(extra)})" if extra else ""))
    print(f"저장: {target}")
    print(json.dumps(summary, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
