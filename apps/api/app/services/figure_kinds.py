"""Figures with a fixed shape (steps, cycle, house, compare, options, axes, messages): slots
the model fills, drawn to a set layout with native .pptx shapes.

A figure's stored source stays a mermaid flowchart, so the browser view, the fact check and
editing read it as before; its last line carries the slots as a mermaid comment, and the
renderer and the deck export draw from those instead of laying the flowchart out.
"""

from __future__ import annotations

import colorsys
import json
import math
import re
from typing import Any

from app.services import figure_text

#: comment line carrying the slots under the fallback flowchart
MARK = "%% kloudchat-figure: "
KINDS = ("steps", "cycle", "house", "compare", "options", "axes", "messages")

INK = "#1B2430"
SUB = "#5A6675"
NEUTRAL = "#64748B"
WHITE = "#FFFFFF"

#: shape labels for the choice; only KINDS are drawn here, the rest stay mermaid
CHOICES = {
    "steps": "one line of 3-7 phases, one after another, with no branches, no second flow "
             "and nothing looping back",
    "cycle": "3-6 stages that repeat in a loop (PDCA, a feedback cycle, a life cycle)",
    "house": "a framework: one goal on top, 2-5 parallel areas under it, and 1-3 foundations "
             "every area stands on",
    "compare": "exactly two options compared side by side on the same aspects",
    "messages": "parties sending each other messages (client and server, a handshake, "
                "a protocol)",
    "options": "three or more options compared",
    "axes": "two axes crossing (a positioning map, a 2×2 matrix)",
    "other": "anything else: an architecture, a system, a model's structure, two or more "
             "flows side by side, a flow that branches or loops back to an earlier step",
}

_CHOOSE = (
    "Pick the shape of the figure described by the user. Answer with one label from the list "
    "and nothing else.\n\n"
    + "\n".join(f"- {label}: {meaning}" for label, meaning in CHOICES.items())
)

#: the slots each kind fills, as an example
_SLOTS = {
    "steps": '{"kind": "steps", "steps": [{"name": "현황 진단", "note": "업무 흐름 조사"}, …]}',
    "cycle": '{"kind": "cycle", "center": "PDCA", "center_note": "품질 개선 순환",\n'
             ' "stages": [{"name": "계획", "en": "Plan", "note": "목표와 개선안 수립"}, …]}',
    "house": '{"kind": "house", "goal": "신뢰할 수 있는 데이터 활용",\n'
             ' "pillars": ["데이터 품질 관리", "메타데이터 관리", …],\n'
             ' "base": [{"name": "거버넌스 조직", "items": ["데이터 위원회", "데이터 관리자"]},\n'
             '          {"name": "정책 · 표준", "note": "모든 영역이 따르는 규칙"}]}',
    "compare": '{"kind": "compare", "sides": ["기존 · 수작업 월간 보고", "자동화 보고"],\n'
               ' "rows": [{"aspect": "수집", "before": "담당자가 5개 시스템에서 엑셀로 수집",\n'
               '           "after": "파이프라인이 매일 자동 수집"}, …],\n'
               ' "hot": {"row": "작성", "badge": "3일 → 즉시"}}\n'
               '`before` is the first side, `after` the second; every row fills both. `hot` '
               'names the one row that matters most; its `badge` is optional.',
    "options": '{"kind": "options", "options": ["외부 API", "자체 호스팅", "하이브리드"],\n'
               ' "rows": [{"aspect": "초기 비용", "cells": [{"text": "낮음", "mark": "good"},\n'
               '           {"text": "GPU 서버 비용 높음", "mark": "bad"}, {"text": "중간"}]}, …],\n'
               ' "hot": "하이브리드"}\n'
               '`cells` follow `options` in order, one per option; a cell the text gives no value '
               'for is "". `mark` is good, caution or bad, only where the text judges that cell. '
               '`hot` names the option the text recommends; leave it out when it recommends none.',
    "axes": '{"kind": "axes", "x": {"name": "가격", "low": "낮음", "high": "높음"},\n'
            ' "y": {"name": "기능 범위", "low": "좁음", "high": "넓음"},\n'
            ' "items": [{"name": "메신저", "x": "low", "y": "low"},\n'
            '           {"name": "우리 제품", "x": "low", "y": "high"}, …],\n'
            ' "hot": "우리 제품"}\n'
            'Each item\'s `x` and `y` is low, mid or high, where the text places it; an item the '
            'text does not place on both axes is left out. `hot` names the item the text is about '
            '(our product, the target); optional.',
    "messages": '{"kind": "messages", "parties": ["클라이언트", "서버"],\n'
                ' "messages": [{"from": "클라이언트", "to": "서버", "name": "ClientHello",\n'
                '               "note": "지원 암호 목록 · 키 공유"}, …]}\n'
                '`from` and `to` are names from `parties`; the messages are in the order they '
                'are sent.',
}

_FILL = """
You fill the slots of one figure from the user's description. The figure is a {kind}: {meaning}.

Answer with one ```json block and, under it, one caption line starting with 「그림: 」 that
says in one sentence what the figure shows. The slots, as an example:

{slots}

- Use only what the description and the text the figure goes with say: the description names
  the figure's parts, and a value it leaves out (a compare cell, a note) comes from that text.
  Never invent a note, number, name or badge; leave an optional field out instead (`note`,
  `en`, `center`, `center_note`, `items`, `hot`, `badge`, `mark`).
- A note adds what the name does not say: what is done, by whom, how long, with what result,
  taken from that text. A note that only rewords the name (「설치 작업」, 「운영 시험」,
  「구체화 과정」) is left out.
- Keep every part the description names; do not merge or drop one to fit.
- Names are short: a step, stage or pillar name about 8 characters (at most 12), a note at
  most 24, a compare or options cell at most 30, an option, axis or item name at most 12, a
  message name at most 40. Use the description's own terms.
- `en` only when the description gives the English term.
- Write the names and the caption in {language}.
"""


def choose_messages(description: str, what: str) -> list[dict[str, str]]:
    """The shape choice for a figure: one label from CHOICES."""
    return [
        {"role": "system", "content": _CHOOSE},
        {"role": "user", "content": f"그릴 것: {what}\n\n설명:\n{description.strip()}"},
    ]


def chosen(text: str) -> str | None:
    """The kind a shape choice picked, or None for one drawn as mermaid."""
    word = re.search(r"[a-z]+", (text or "").lower())
    return word.group(0) if word and word.group(0) in KINDS else None


def fill_messages(
    kind: str, description: str, language: str, material: str = ""
) -> list[dict[str, str]]:
    """The slots of one kind, from the description and the text the figure goes with."""
    system = _FILL.format(
        kind=kind, meaning=CHOICES[kind], slots=_SLOTS[kind],
        language="English" if language == "en" else "Korean",
    )
    return [
        {"role": "system", "content": system.strip()},
        {"role": "user", "content": f"설명:\n{description.strip()}" + (
            f"\n\n그림이 들어갈 글:\n{material.strip()[:4000]}" if material.strip() else "")},
    ]


# ── slots ─────────────────────────────────────────────────────────────


def _text(value: Any, limit: int) -> str:
    """One line of at most `limit` characters, cut at a space when one is near the end."""
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if len(text) <= limit:
        return text
    cut = text[:limit]
    space = cut.rfind(" ")
    return (cut[:space] if space >= limit // 2 else cut).rstrip(" ·,")


def _repeats(text: str, said: str) -> bool:
    """Whether every word of `text` is already in `said` (text the figure shows anyway)."""
    words = re.findall(r"[가-힣A-Za-z0-9%]+", text)
    return bool(words) and all(word in said for word in words)


def _named(value: Any, limit: int = 24) -> dict[str, str] | None:
    """`{"name", "note"?, "en"?}` from a dict or a bare name."""
    if isinstance(value, str):
        value = {"name": value}
    if not isinstance(value, dict):
        return None
    name = _text(value.get("name"), limit)
    if not name:
        return None
    out = {"name": name}
    for key, size in (("note", 40), ("en", 16)):
        if text := _text(value.get(key), size):
            out[key] = text
    if out.get("note") and (out["note"].replace(" ", "") in name.replace(" ", "")
                            or _repeats(out["note"], name)):
        out.pop("note")  # a note that only repeats the name says nothing
    return out


#: a count with its unit: 「3주」「120석」「2,000만 원」, 「간」「동안」 kept with it
_QUANTITY = re.compile(r"\d[\d,.]*\s{0,2}"
                       r"(?:개월|주|일|시간|분|년|명|석|회|건|단계|%|만\s?원|억\s?원|원)"
                       r"(?:간|\s{0,2}동안)?")
#: one clause of the material: sentences cut at 「,」「;」 and line ends
_CLAUSE = re.compile(r"[^,;，、\n.!?。]+")


def _stems(name: str) -> list[str]:
    return [w[:max(2, len(w) - 1)] for w in re.findall(r"[가-힣A-Za-z]{2,}", name)]


def _owns(clause: str, part: list[str], others: list[list[str]]) -> bool:
    """Whether a clause speaks of this part alone (no other part named in it)."""
    names = lambda stems: any(stem in clause for stem in stems)  # noqa: E731
    return names(part) and not any(names(o) for o in others)


def grounded(s: dict[str, Any], material: str) -> dict[str, Any]:
    """Step and stage notes kept to counts the material gives that part: a count the material
    says only of the whole (「공사는 3주 동안 철거·설치·시범 운영」) or never says is cut from
    a note, and a note left saying nothing is dropped."""
    key = {"steps": "steps", "cycle": "stages"}.get(s.get("kind", ""))
    if not key or not material.strip():
        return s
    parts = s[key]
    stems = [_stems(p["name"]) for p in parts]
    clauses = [c for c in _CLAUSE.findall(material) if c.strip()]
    for k, part in enumerate(parts):
        note = part.get("note")
        if not note:
            continue
        others = stems[:k] + stems[k + 1:]
        for m in list(_QUANTITY.finditer(note)):
            count = re.match(r"\d[\d,.]*\s{0,2}\D+?(?=간|\s{0,2}동안|$)", m.group(0))
            bare = re.sub(r"\s", "", (count or m).group(0))
            if any(bare in re.sub(r"\s", "", c) and _owns(c, stems[k], others) for c in clauses):
                continue
            note = note.replace(m.group(0), " ")
        note = re.sub(r"\s+", " ", note).strip(" ·,")
        if note != part["note"]:
            if len(note) < 2 or _repeats(note, part["name"]) or \
                    note.replace(" ", "") in part["name"].replace(" ", ""):
                part.pop("note")
            else:
                part["note"] = note
    return s


def spec(data: Any, language: str = "ko") -> dict[str, Any] | None:
    """Checked slots for one figure kind, or None when they don't make that figure; `band`
    (a figure for a slide's wide strip) kept."""
    out = _spec(data, language)
    if out is not None and isinstance(data, dict) and data.get("band") and out["kind"] in _BAND:
        out["band"] = True
    return out


def _spec(data: Any, language: str = "ko") -> dict[str, Any] | None:
    if not isinstance(data, dict) or data.get("kind") not in KINDS:
        return None
    kind = data["kind"]
    lang = "en" if language == "en" else "ko"
    if kind in ("steps", "cycle"):
        key = "steps" if kind == "steps" else "stages"
        items = [n for v in (data.get(key) or [])[: 7 if kind == "steps" else 6]
                 if (n := _named(v))]
        if len(items) < 3:
            return None
        out: dict[str, Any] = {"kind": kind, "lang": lang, key: items}
        if kind == "cycle":
            for field, size in (("center", 14), ("center_note", 20)):
                if text := _text(data.get(field), size):
                    out[field] = text
            if out.get("center_note") and _repeats(out["center_note"], out.get("center", "")):
                out.pop("center_note")
        return out
    if kind == "house":
        goal = _text(data.get("goal"), 40)
        pillars = [t for v in (data.get("pillars") or [])[:5] if (t := _text(v, 24))]
        base = []
        for row in (data.get("base") or [])[:3]:
            named = _named(row, 24)
            if named is None:
                continue
            items = [t for v in (row.get("items") or [])[:4] if (t := _text(v, 16))] \
                if isinstance(row, dict) else []
            if items:
                named["items"] = items
            base.append(named)
        if not goal or len(pillars) < 2:
            return None
        return {"kind": kind, "lang": lang, "goal": goal, "pillars": pillars, "base": base}
    if kind == "options":
        return _options(data, lang)
    if kind == "axes":
        return _axes(data, lang)
    if kind == "messages":
        return _messages(data, lang)
    sides = [t for v in (data.get("sides") or [])[:2] if (t := _text(v, 30))]
    rows = []
    for row in (data.get("rows") or [])[:6]:
        if not isinstance(row, dict):
            continue
        aspect, before, after = (_text(row.get(k), 30) for k in ("aspect", "before", "after"))
        # 「-」 or 「없음」 stands in for a cell the writer had no value for
        before, after = ("" if re.fullmatch(r"[-–—·.\s]*|없음|해당\s?없음", cell) else cell
                         for cell in (before, after))
        if aspect and before and after:
            rows.append({"aspect": aspect, "before": before, "after": after})
    if len(sides) < 2 or len(rows) < 2 or len(rows) < len(data.get("rows") or []):
        return None  # a row with one side empty is not this comparison
    out = {"kind": kind, "lang": lang, "sides": sides, "rows": rows}
    hot = data.get("hot") if isinstance(data.get("hot"), dict) else {}
    if (row := _text(hot.get("row"), 30)) and any(r["aspect"] == row for r in rows):
        out["hot"] = {"row": row}
        cells = next(r for r in rows if r["aspect"] == row)
        if (badge := _text(hot.get("badge"), 16)) and not _repeats(
                badge, cells["before"] + " " + cells["after"]):
            out["hot"]["badge"] = badge
    return out


#: 「-」 or 「없음」 stands in for a cell the writer had no value for
_EMPTY = re.compile(r"[-–—·.\s]*|없음|해당\s?없음")
_MARKS = ("good", "caution", "bad")


def _options(data: dict, lang: str) -> dict[str, Any] | None:
    names = [t for v in (data.get("options") or [])[:4] if (t := _text(v, 20))]
    if len(names) < 3 or len(set(names)) < len(names):
        return None  # two options are a compare
    rows = []
    for row in (data.get("rows") or [])[:7]:
        if not isinstance(row, dict) or not (aspect := _text(row.get("aspect"), 16)):
            continue
        cells = []
        for cell in list(row.get("cells") or [])[: len(names)]:
            cell = cell if isinstance(cell, dict) else {"text": cell}
            text = _text(cell.get("text"), 30)
            text = "" if _EMPTY.fullmatch(text) else text
            out = {"text": text}
            if text and (mark := str(cell.get("mark") or "").strip().lower()) in _MARKS:
                out["mark"] = mark
            cells.append(out)
        cells += [{"text": ""}] * (len(names) - len(cells))
        if sum(bool(c["text"]) for c in cells) >= 2:
            rows.append({"aspect": aspect, "cells": cells})
    if len(rows) < 2:
        return None
    out: dict[str, Any] = {"kind": "options", "lang": lang, "options": names, "rows": rows}
    if (hot := _text(data.get("hot"), 20)) in names:
        out["hot"] = hot
    return out


_PLACES = ("low", "mid", "high")


def _axis(value: Any) -> dict[str, str] | None:
    if isinstance(value, str):
        value = {"name": value}
    if not isinstance(value, dict) or not (name := _text(value.get("name"), 12)):
        return None
    out = {"name": name}
    for end in ("low", "high"):
        if text := _text(value.get(end), 10):
            out[end] = text
    return out


def _axes(data: dict, lang: str) -> dict[str, Any] | None:
    x, y = _axis(data.get("x")), _axis(data.get("y"))
    if x is None or y is None:
        return None
    items, seen = [], set()
    for item in (data.get("items") or [])[:8]:
        if not isinstance(item, dict) or not (name := _text(item.get("name"), 14)):
            continue
        at = [str(item.get(k) or "").strip().lower() for k in ("x", "y")]
        if name in seen or any(p not in _PLACES for p in at):
            continue  # an item the text does not place on both axes
        seen.add(name)
        items.append({"name": name, "x": at[0], "y": at[1]})
    if len(items) < 3:
        return None
    out: dict[str, Any] = {"kind": "axes", "lang": lang, "x": x, "y": y, "items": items}
    if (hot := _text(data.get("hot"), 14)) in seen:
        out["hot"] = hot
    return out


def _messages(data: dict, lang: str) -> dict[str, Any] | None:
    parties = list(dict.fromkeys(
        t for v in (data.get("parties") or [])[:4] if (t := _text(v, 14))))
    if len(parties) < 2:
        return None
    sent = []
    for message in (data.get("messages") or [])[:10]:
        if not isinstance(message, dict):
            continue
        ends = [_text(message.get(k), 14) for k in ("from", "to")]
        name = _text(message.get("name"), 40)
        if not name or ends[0] == ends[1] or any(e not in parties for e in ends):
            continue
        out = {"from": ends[0], "to": ends[1], "name": name}
        if (note := _text(message.get("note"), 30)) and not _repeats(note, name):
            out["note"] = note
        sent.append(out)
    if len(sent) < 2:
        return None
    return {"kind": "messages", "lang": lang, "parties": parties, "messages": sent}


# ── stored source ─────────────────────────────────────────────────────


def _q(text: str) -> str:
    """Label as a quoted mermaid label."""
    return '"' + text.replace('"', "'") + '"'


def _flowchart(s: dict[str, Any]) -> list[str]:
    """Plain flowchart of the same figure, for views that draw mermaid."""
    kind = s["kind"]
    if kind in ("steps", "cycle"):
        items = s["steps" if kind == "steps" else "stages"]
        lines = ["flowchart LR"]
        for k, item in enumerate(items):
            hot = ":::hot" if kind == "steps" and k == len(items) - 1 else ""
            lines.append(f"  s{k}({_q(item['name'])}){hot}")
        lines.append("  " + " --> ".join(f"s{k}" for k in range(len(items))))
        if kind == "cycle":
            lines.append(f"  s{len(items) - 1} -.-> s0")
        return lines
    if kind == "house":
        lines = ["flowchart TB", f"  g({_q(s['goal'])}):::hot"]
        for k, name in enumerate(s["pillars"]):
            lines += [f"  p{k}({_q(name)})", f"  g --> p{k}"]
        for k, row in enumerate(s["base"]):
            label = row["name"] + (": " + "·".join(row["items"]) if row.get("items") else "")
            lines.append(f"  b{k}[{_q(label)}]")
            lines += [f"  p{j} --> b{k}" for j in range(len(s["pillars"]))] if k == 0 else [
                f"  b{k - 1} --> b{k}"]
        return lines
    if kind in _FLOWS:
        return _FLOWS[kind](s)
    hot = (s.get("hot") or {}).get("row")
    lines = ["flowchart LR"]
    for side, key, field in (("old", "a", "before"), ("new", "b", "after")):
        title = s["sides"][0 if key == "a" else 1]
        lines += [f"  subgraph {side}[{_q(title)}]", "    direction TB"]
        for k, row in enumerate(s["rows"]):
            mark = ":::hot" if key == "b" and row["aspect"] == hot else ""
            lines.append(f"    {key}{k}[{_q(row['aspect'] + ': ' + (row[field] or '-'))}]{mark}")
        if len(s["rows"]) > 1:
            lines.append("    " + " ~~~ ".join(f"{key}{k}" for k in range(len(s["rows"]))))
        lines.append("  end")
    lines.append("  old -.-> new")
    return lines


def _options_flow(s: dict[str, Any]) -> list[str]:
    lines = ["flowchart LR"]
    for k, name in enumerate(s["options"]):
        lines += [f"  subgraph o{k}[{_q(name)}]", "    direction TB"]
        for j, row in enumerate(s["rows"]):
            mark = ":::hot" if j == 0 and name == s.get("hot") else ""
            cell = row["cells"][k]["text"] or "-"
            lines.append(f"    o{k}r{j}[{_q(row['aspect'] + ': ' + cell)}]{mark}")
        if len(s["rows"]) > 1:
            lines.append("    " + " ~~~ ".join(f"o{k}r{j}" for j in range(len(s["rows"]))))
        lines.append("  end")
    return lines


def _axes_flow(s: dict[str, Any]) -> list[str]:
    lines = ["flowchart TB"]
    places = {"low": "low", "mid": "mid", "high": "high"}
    for k, item in enumerate(s["items"]):
        where = (f"{s['x']['name']} {s['x'].get(item['x'], places[item['x']])} · "
                 f"{s['y']['name']} {s['y'].get(item['y'], places[item['y']])}")
        mark = ":::hot" if item["name"] == s.get("hot") else ""
        lines.append(f"  i{k}[{_q(item['name'] + ': ' + where)}]{mark}")
    return lines


def _messages_flow(s: dict[str, Any]) -> list[str]:
    lines = ["flowchart TB"]
    for k, m in enumerate(s["messages"]):
        label = f"{k + 1}. {m['from']} → {m['to']}: {m['name']}"
        if m.get("note"):
            label += f" ({m['note']})"
        lines.append(f"  m{k}[{_q(label)}]")
    lines.append("  " + " --> ".join(f"m{k}" for k in range(len(s["messages"]))))
    return lines


_FLOWS = {"options": _options_flow, "axes": _axes_flow, "messages": _messages_flow}


def to_source(s: dict[str, Any]) -> str:
    """Stored source: the fallback flowchart, then the slots as a comment line."""
    slots = json.dumps(s, ensure_ascii=False, separators=(",", ":"))
    return "\n".join([*_flowchart(s), MARK + slots])


def from_source(source: str) -> dict[str, Any] | None:
    """Slots stored under a flowchart, or None for a plain flowchart."""
    for line in reversed((source or "").splitlines()):
        if line.startswith(MARK):
            try:
                data = json.loads(line[len(MARK):])
            except ValueError:
                return None
            if not isinstance(data, dict):
                return None
            found = spec(data, str(data.get("lang") or "ko"))
            if found is not None and found.get("hot") and ":::hot" not in source:
                found.pop("hot")  # the fact check took the highlight off
            return found
    return None


# ── colour ────────────────────────────────────────────────────────────


def _code(r: float, g: float, b: float) -> str:
    return f"#{round(r):02X}{round(g):02X}{round(b):02X}"


def _hexed(colour: Any) -> str:
    from app.services import diagram_shapes

    return _code(*diagram_shapes._hex(colour))


def tint(colour: str, amount: float) -> str:
    """`colour` mixed toward white by `amount` (0-1)."""
    c = colour.lstrip("#")
    rgb = [int(c[i:i + 2], 16) for i in (0, 2, 4)]
    return _code(*(v + (255 - v) * amount for v in rgb))


#: hue steps from the accent for figures with one colour per part (blue → teal, amber, violet)
_HUES = (0.0, -0.135, 0.455, 0.115, -0.27, 0.3)


def _luminance(r: float, g: float, b: float) -> float:
    """WCAG relative luminance of 0-1 channels."""
    def channel(c: float) -> float:
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    return 0.2126 * channel(r) + 0.7152 * channel(g) + 0.0722 * channel(b)


def hues(accent: str, n: int) -> list[str]:
    """`n` distinct colours around the accent, each dark enough for white text (4.5:1)."""
    c = accent.lstrip("#")
    r, g, b = (int(c[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, lightness, s = colorsys.rgb_to_hls(r, g, b)
    out = []
    for k in range(n):
        hue, light = (h + _HUES[k % len(_HUES)]) % 1.0, min(lightness, 0.5)
        sat = max(0.55, min(s, 0.85))
        rgb = colorsys.hls_to_rgb(hue, light, sat)
        while _luminance(*rgb) > 0.183 and light > 0.15:
            light -= 0.02
            rgb = colorsys.hls_to_rgb(hue, light, sat)
        out.append(_code(*(v * 255 for v in rgb)))
    return out


# ── drawing ───────────────────────────────────────────────────────────


class Pen:
    """Shapes on a slide in figure units, scaled and centred into the figure's box."""

    def __init__(self, slide, *, left: float, top: float, width: float, height: float,
                 size: tuple[float, float], font: str) -> None:
        from pptx.util import Pt

        self.shapes = slide.shapes
        self.k = min(width / size[0], height / size[1])
        self.ox = left + (width - size[0] * self.k) / 2
        self.oy = top + (height - size[1] * self.k) / 2
        self.font = font
        self._pt = Pt

    def _x(self, x: float):
        return self._pt(self.ox + x * self.k)

    def _y(self, y: float):
        return self._pt(self.oy + y * self.k)

    def _d(self, d: float):
        return self._pt(d * self.k)

    @staticmethod
    def _rgb(colour: str):
        from pptx.dml.color import RGBColor

        return RGBColor.from_string(colour.lstrip("#").upper())

    def _paint(self, shape, fill: str | None, stroke: str | None, sw: float = 1.5,
               dash: bool = False) -> None:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.oxml.ns import qn

        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = self._rgb(fill)
        else:
            shape.fill.background()
        if stroke:
            shape.line.color.rgb = self._rgb(stroke)
            shape.line.width = self._d(sw)
            if dash:
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            shape.line.fill.background()
        # theme style brings a shadow in LibreOffice
        for style in shape._element.findall(qn("p:style")):
            shape._element.remove(style)

    def rect(self, x: float, y: float, w: float, h: float, fill: str | None = WHITE,
             stroke: str | None = None, r: float = 8, sw: float = 1.5, dash: bool = False):
        from pptx.enum.shapes import MSO_SHAPE

        kind = MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE
        shape = self.shapes.add_shape(kind, self._x(x), self._y(y), self._d(w), self._d(h))
        if r:
            shape.adjustments[0] = min(0.5, r / min(w, h))
        self._paint(shape, fill, stroke, sw, dash)
        return shape

    def circle(self, x: float, y: float, r: float, fill: str | None, stroke: str | None = None):
        from pptx.enum.shapes import MSO_SHAPE

        shape = self.shapes.add_shape(
            MSO_SHAPE.OVAL, self._x(x - r), self._y(y - r), self._d(2 * r), self._d(2 * r))
        self._paint(shape, fill, stroke)
        return shape

    def poly(self, pts: list[tuple[float, float]], fill: str | None, stroke: str | None = None,
             *, closed: bool = True, sw: float = 1.5, dash: bool = False, arrow: bool = False):
        from pptx.oxml.ns import qn

        build = self.shapes.build_freeform(self._x(pts[0][0]), self._y(pts[0][1]), scale=1.0)
        build.add_line_segments([(self._x(a), self._y(b)) for a, b in pts[1:]], close=closed)
        shape = build.convert_to_shape()
        self._paint(shape, fill, stroke, sw, dash)
        if arrow:
            line = shape._element.spPr.find(qn("a:ln"))
            line.append(line.makeelement(qn("a:tailEnd"),
                                         {"type": "triangle", "w": "med", "len": "med"}))
        return shape

    def line(self, pts: list[tuple[float, float]], colour: str, *, arrow: bool = True,
             sw: float = 1.6, dash: bool = False):
        return self.poly(pts, None, colour, closed=False, sw=sw, dash=dash, arrow=arrow)

    def _runs(self, frame, lines: list[str], size: float, bold: bool, colour: str,
              align: str) -> None:
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.oxml.ns import qn

        frame.word_wrap = False
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for n, text in enumerate(lines):
            para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
            para.alignment = {"c": PP_ALIGN.CENTER, "l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT}[align]
            run = para.add_run()
            run.text = text
            run.font.name = self.font
            run.font.size = self._d(size)
            run.font.bold = bold
            run.font.color.rgb = self._rgb(colour)
            props = run._r.get_or_add_rPr()
            props.append(props.makeelement(qn("a:ea"), {"typeface": self.font}))

    def text(self, x: float, y: float, text: str | list[str], size: float = 14,
             bold: bool = False, colour: str = INK, anchor: str = "c") -> None:
        """Lines centred on `y`, centred on `x` or starting/ending there."""
        lines = text if isinstance(text, list) else text.split("\n")
        w = max(figure_text.width(line, size, bold) for line in lines) + 8
        h = len(lines) * size * figure_text.LINE
        left = {"c": x - w / 2, "l": x, "r": x - w}[anchor]
        box = self.shapes.add_textbox(self._x(left), self._y(y - h / 2), self._d(w), self._d(h))
        self._runs(box.text_frame, lines, size, bold, colour, anchor)

    def label(self, x: float, y: float, w: float, h: float, text: str, size: float,
              bold: bool, colour: str, min_size: float = 9.0) -> None:
        """`text` fitted into `w` × `h` and centred there."""
        lines, fitted = figure_text.fit(text, w, h, size, bold, min_size=min_size)
        self.text(x + w / 2, y + h / 2, lines, fitted, bold, colour)


# ── kinds ─────────────────────────────────────────────────────────────

_STEP_W, _STEP_TIP, _STEP_H = 190.0, 26.0, 84.0


def _steps_size(s: dict[str, Any]) -> tuple[float, float]:
    n = len(s["steps"])
    notes = any(item.get("note") for item in s["steps"])
    return 40 + n * (_STEP_W - 6) + 6, 50 + _STEP_H + (92 if notes else 50)


def _steps(pen: Pen, s: dict[str, Any], accent: str) -> None:
    items = s["steps"]
    n = len(items)
    y = 30.0
    for i, item in enumerate(items):
        x = 20 + i * (_STEP_W - 6)
        last = i == n - 1
        amount = 0.0 if last else 0.78 - i * (0.62 / max(1, n - 2))
        colour = accent if last else tint(accent, amount)
        dark = amount > 0.45
        pts = [(x, y), (x + _STEP_W - _STEP_TIP, y), (x + _STEP_W, y + _STEP_H / 2),
               (x + _STEP_W - _STEP_TIP, y + _STEP_H), (x, y + _STEP_H)]
        if i:
            pts.append((x + _STEP_TIP, y + _STEP_H / 2))
        pen.poly(pts, colour, WHITE, sw=3)
        inner = x + (_STEP_TIP if i else 8)
        room = _STEP_W - _STEP_TIP - (_STEP_TIP if i else 8) - 2
        pen.label(inner, y, room, _STEP_H, item["name"], 16, True, INK if dark else WHITE, 11)
        cx = inner + room / 2
        number = f"Step {i + 1}" if s["lang"] == "en" else f"{i + 1}단계"
        pen.text(cx, y + _STEP_H + 22, number, 12, True, SUB)
        if note := item.get("note"):
            lines, size = figure_text.fit(note, _STEP_W - 20, 48, 13, min_size=10.5)
            pen.text(cx, y + _STEP_H + 58, lines, size, False, INK)


_RING_OUT, _RING_IN, _CALLOUT = 160.0, 96.0, 176.0


def _cycle_size(s: dict[str, Any]) -> tuple[float, float]:
    callouts = any(stage.get("note") for stage in s["stages"])
    return 2 * (_RING_OUT + (50 + _CALLOUT if callouts else 10)) + 20, 2 * (_RING_OUT + 50) + 40


def _cycle(pen: Pen, s: dict[str, Any], accent: str) -> None:
    stages = s["stages"]
    n = len(stages)
    w, h = _cycle_size(s)
    cx, cy = w / 2, h / 2
    span = 360.0 / n

    def at(radius: float, degrees: float) -> tuple[float, float]:
        return (cx + radius * math.cos(math.radians(degrees)),
                cy + radius * math.sin(math.radians(degrees)))

    colours = hues(accent, n)
    for k, stage in enumerate(stages):
        a0, a1 = -90 + span * k + 3, -90 + span * (k + 1) - 9
        outer = [at(_RING_OUT, a0 + (a1 - a0) * j / 24) for j in range(25)]
        inner = [at(_RING_IN, a1 - (a1 - a0) * j / 24) for j in range(25)]
        pen.poly([*outer, at((_RING_OUT + _RING_IN) / 2, a1 + 7), *inner], colours[k])
        mid = (a0 + a1) / 2
        mx, my = at((_RING_OUT + _RING_IN) / 2, mid)
        name_lines, name_size = figure_text.fit(stage["name"], 76, 40, 17, True, min_size=11)
        if stage.get("en"):
            pen.text(mx, my - 9, name_lines, name_size, True, WHITE)
            pen.text(mx, my + 12, stage["en"], 12, False, tint(colours[k], 0.75))
        else:
            pen.text(mx, my, name_lines, name_size, True, WHITE)
        note = stage.get("note")
        if not note:
            continue
        # the note outside the ring, on the stage's side
        cos = math.cos(math.radians(mid))
        side = 0 if abs(cos) < 0.3 else (1 if cos > 0 else -1)
        ex, ey = at(_RING_OUT + 34, mid)
        pen.line([at(_RING_OUT + 6, mid), (ex, ey)], tint(colours[k], 0.4), arrow=False, sw=1.2)
        pen.circle(ex, ey, 3, colours[k])
        anchor = {1: "l", -1: "r", 0: "c"}[side]
        tx = ex + side * 8
        # the note alone: the stage's name is already on the ring
        note_lines, note_size = figure_text.fit(note, _CALLOUT, 56, 14, min_size=10.5)
        block = len(note_lines) * note_size * figure_text.LINE
        if side == 0:  # above or below the ring
            middle = ey - block / 2 - 6 if my < cy else ey + block / 2 + 6
        else:
            middle = ey
        pen.text(tx, middle, note_lines, note_size, False, INK, anchor)
    pen.circle(cx, cy, _RING_IN - 14, "#F1F5F9")
    if center := s.get("center"):
        lines, size = figure_text.fit(center, 2 * (_RING_IN - 14) - 12, 40, 22, True,
                                      min_size=13, pad=4)
        pen.text(cx, cy - (8 if s.get("center_note") else 0), lines, size, True, INK)
    if note := s.get("center_note"):
        lines, size = figure_text.fit(note, 2 * (_RING_IN - 14) - 16, 30, 12, min_size=9.5, pad=2)
        pen.text(cx, cy + (16 if center else 0), lines, size, False, SUB)


def _house_size(s: dict[str, Any]) -> tuple[float, float]:
    return 960.0, 120 + 16 + 202 + 14 + 52 * len(s["base"]) + 10


def _house(pen: Pen, s: dict[str, Any], accent: str) -> None:
    deep = _deepened(accent)
    pen.poly([(480, 24), (900, 120), (60, 120)], deep)
    pen.text(480, 58, "Goal" if s["lang"] == "en" else "목표", 11, True, tint(deep, 0.6))
    pen.label(260, 70, 440, 40, s["goal"], 19, True, WHITE, 12)
    pillars = s["pillars"]
    n, gap = len(pillars), 20.0
    pw = (780 - gap * (n - 1)) / n
    for i, name in enumerate(pillars):
        x = 90 + i * (pw + gap)
        pen.rect(x, 136, pw, 18, tint(deep, 0.55), r=3)
        pen.rect(x + 12, 154, pw - 24, 170, tint(deep, 0.88), r=0)
        pen.rect(x, 324, pw, 14, tint(deep, 0.55), r=3)
        pen.label(x + 12, 154, pw - 24, 170, name, 15, True, deep, 10)
    for k, row in enumerate(s["base"]):
        y = 352 + k * 52
        last = k == len(s["base"]) - 1
        fill = deep if last and len(s["base"]) > 1 else tint(deep, 0.25)
        pen.rect(60, y, 840, 46, fill, r=6)
        if not row.get("items") and not row.get("note"):
            pen.label(60, y, 840, 46, row["name"], 15, True, WHITE, 11)
            continue
        name_lines, name_size = figure_text.fit(row["name"], 150, 46, 13, True, min_size=10, pad=4)
        pen.text(80, y + 23, name_lines, name_size, True, WHITE, "l")
        if row.get("items"):
            pen.label(240, y, 620, 46, "   ·   ".join(row["items"]), 15, True, WHITE, 11)
        else:
            pen.label(240, y, 620, 46, row["note"], 13, False, tint(deep, 0.6), 10)


def _deepened(accent: str) -> str:
    """The accent darkened for large fills under white text."""
    c = accent.lstrip("#")
    r, g, b = (int(c[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, lightness, s = colorsys.rgb_to_hls(r, g, b)
    rr, gg, bb = colorsys.hls_to_rgb(h, min(lightness, 0.32), s)
    return _code(rr * 255, gg * 255, bb * 255)


def _compare_size(s: dict[str, Any]) -> tuple[float, float]:
    return 960.0, 92 + 82 * len(s["rows"]) + 4


def _compare(pen: Pen, s: dict[str, Any], accent: str) -> None:
    old, new = NEUTRAL, accent
    pen.rect(40, 20, 360, 52, tint(old, 0.85), r=10)
    pen.label(40, 20, 360, 52, s["sides"][0], 16, True, INK, 11)
    pen.rect(560, 20, 360, 52, new, r=10)
    pen.label(560, 20, 360, 52, s["sides"][1], 16, True, WHITE, 11)
    hot = s.get("hot") or {}
    for i, row in enumerate(s["rows"]):
        y = 92 + i * 82
        is_hot = row["aspect"] == hot.get("row")
        pen.rect(40, y, 360, 66, "#F8FAFC", tint(old, 0.55))
        if row["before"]:
            pen.label(40, y, 360, 66, row["before"], 14, False, "#334155", 10)
        pen.rect(560, y, 360, 66, tint(new, 0.86) if is_hot else WHITE, new)
        if row["after"]:
            pen.label(560, y, 360, 66, row["after"], 15 if is_hot else 14, is_hot,
                      new if is_hot else INK, 10)
        pen.rect(424, y + 17, 112, 32, tint(INK, 0.9), r=16)
        pen.label(424, y + 17, 112, 32, row["aspect"], 13, True, INK, 9)
        pen.line([(402, y + 33), (420, y + 33)], tint(old, 0.4), arrow=False)
        pen.line([(538, y + 33), (556, y + 33)], new)
        if is_hot and (badge := hot.get("badge")):
            bw = figure_text.width(badge, 12, True) + 18
            pen.rect(740 - bw / 2, y - 11, bw, 22, new, r=11)
            pen.text(740, y, badge, 12, True, WHITE)


_GOOD, _CAUTION, _BAD = "#15803D", "#B45309", "#B42318"
_MARK_COLOUR = {"good": _GOOD, "caution": _CAUTION, "bad": _BAD}
_MARK_WORD = {"ko": {"good": "유리", "caution": "주의", "bad": "불리"},
              "en": {"good": "Favourable", "caution": "Caution", "bad": "Unfavourable"}}
_OPT_X, _OPT_ROW = 190.0, 56.0


def _marked(s: dict[str, Any]) -> list[str]:
    used = {c.get("mark") for row in s["rows"] for c in row["cells"]}
    return [m for m in _MARKS if m in used]


def _options_size(s: dict[str, Any]) -> tuple[float, float]:
    return 960.0, 78 + _OPT_ROW * len(s["rows"]) + (40 if _marked(s) else 8)


def _options_draw(pen: Pen, s: dict[str, Any], accent: str) -> None:
    names, rows = s["options"], s["rows"]
    m, gap = len(names), 10.0
    cw = (960 - _OPT_X - 20 - gap * (m - 1)) / m
    # the recommended option takes the accent, the others hues clear of the marks' green, amber
    # and red
    spare = [c for k, c in enumerate(hues(accent, 6)) if k in (0, 1, 3, 5)]
    hot = names.index(s["hot"]) if s.get("hot") in names else None
    order = iter(spare[1:] if hot is not None else spare)
    colours = [accent if k == hot else next(order) for k in range(m)]
    bottom = 78 + _OPT_ROW * len(rows) - 4
    for j in range(0, len(rows), 2):
        pen.rect(20, 78 + j * _OPT_ROW, 920, _OPT_ROW - 4, "#F6F8FA", r=6)
    for k, name in enumerate(names):
        x = _OPT_X + k * (cw + gap)
        if name == s.get("hot"):
            pen.rect(x, 72, cw, bottom - 72 + 6, tint(colours[k], 0.9), tint(colours[k], 0.5), r=10)
        pen.rect(x, 20, cw, 46, colours[k], r=10)
        pen.label(x + 8, 20, cw - 16, 46, name, 16, True, WHITE, 11)
    for j, row in enumerate(rows):
        y = 78 + j * _OPT_ROW
        mid = y + (_OPT_ROW - 4) / 2
        lines, size = figure_text.fit(row["aspect"], _OPT_X - 52, _OPT_ROW - 8, 14, True,
                                      min_size=10)
        pen.text(36, mid, lines, size, True, INK, "l")
        for k, cell in enumerate(row["cells"]):
            if not cell["text"]:
                continue
            x = _OPT_X + k * (cw + gap)
            left = x + 14
            if mark := cell.get("mark"):
                pen.circle(x + 18, mid, 6, _MARK_COLOUR[mark])
                left = x + 32
            lines, size = figure_text.fit(cell["text"], x + cw - 8 - left, _OPT_ROW - 8, 14,
                                          min_size=10)
            pen.text(left, mid, lines, size, False, INK, "l")
    x = 24.0
    for mark in _marked(s):
        y = bottom + 24
        pen.circle(x, y, 5, _MARK_COLOUR[mark])
        word = _MARK_WORD[s["lang"]][mark]
        pen.text(x + 10, y, word, 12, False, SUB, "l")
        x += figure_text.width(word, 12) + 36


_AX0, _AX1, _AY0, _AY1 = 130.0, 900.0, 30.0, 410.0
_AT = {"low": 0.25, "mid": 0.5, "high": 0.75}


def _axes_size(s: dict[str, Any]) -> tuple[float, float]:
    return 960.0, 480.0


def _axes_draw(pen: Pen, s: dict[str, Any], accent: str) -> None:
    mx, my = (_AX0 + _AX1) / 2, (_AY0 + _AY1) / 2
    pen.rect(_AX0, _AY0, _AX1 - _AX0, _AY1 - _AY0, "#F8FAFC", r=6)
    hot = next((i for i in s["items"] if i["name"] == s.get("hot")), None)
    if hot and hot["x"] != "mid" and hot["y"] != "mid":
        qx = _AX0 if hot["x"] == "low" else mx
        qy = _AY0 if hot["y"] == "high" else my
        pen.rect(qx, qy, mx - _AX0, my - _AY0, tint(accent, 0.93), r=0)
    pen.line([(mx, _AY0), (mx, _AY1)], tint(NEUTRAL, 0.55), arrow=False, sw=1, dash=True)
    pen.line([(_AX0, my), (_AX1, my)], tint(NEUTRAL, 0.55), arrow=False, sw=1, dash=True)
    pen.line([(_AX0, _AY1), (_AX1 + 20, _AY1)], INK)
    pen.line([(_AX0, _AY1), (_AX0, _AY0 - 14)], INK)
    x, y = s["x"], s["y"]
    pen.text(mx, _AY1 + 46, x["name"], 14, True, INK)
    if x.get("low"):
        pen.text(_AX0 + 6, _AY1 + 20, x["low"], 12, False, SUB, "l")
    if x.get("high"):
        pen.text(_AX1, _AY1 + 20, x["high"], 12, False, SUB, "r")
    if s["lang"] == "en":
        pen.text(_AX0 + 10, _AY0 - 14, y["name"], 14, True, INK, "l")
    else:  # read top to bottom beside the axis
        pen.text(70, my, list(y["name"].replace(" ", "")), 14, True, INK)
    if y.get("low"):
        pen.text(_AX0 - 12, _AY1 - 10, y["low"], 12, False, SUB, "r")
    if y.get("high"):
        pen.text(_AX0 - 12, _AY0 + 12, y["high"], 12, False, SUB, "r")
    cells: dict[tuple[str, str], list[dict[str, str]]] = {}
    for item in s["items"]:
        cells.setdefault((item["x"], item["y"]), []).append(item)
    letter = 0
    for (px, py), group in cells.items():
        cx = _AX0 + (_AX1 - _AX0) * _AT[px]
        cy = _AY1 - (_AY1 - _AY0) * _AT[py]
        spread = min(130.0, (_AX1 - _AX0) / 2 / max(1, len(group)))
        for k, item in enumerate(group):
            ix = cx + (k - (len(group) - 1) / 2) * spread
            if item is hot:
                halo = [(ix + 38 * math.cos(math.radians(a)), cy + 38 * math.sin(math.radians(a)))
                        for a in range(0, 360, 10)]
                pen.poly(halo, None, accent, sw=1.5, dash=True)
                pen.circle(ix, cy, 24, accent)
                lines, size = figure_text.fit(item["name"], spread - 8, 34, 13, True, min_size=10)
                pen.text(ix, cy + 58, lines, size, True, accent)
                continue
            pen.circle(ix, cy, 22, WHITE, NEUTRAL)
            pen.text(ix, cy, chr(ord("A") + letter), 15, True, NEUTRAL)
            letter += 1
            lines, size = figure_text.fit(item["name"], spread - 8, 34, 13, True, min_size=10)
            pen.text(ix, cy + 40, lines, size, True, INK)


_MSG_TOP = 100.0


def _messages_step(s: dict[str, Any]) -> float:
    return 62.0 if any(m.get("note") for m in s["messages"]) else 46.0


def _messages_size(s: dict[str, Any]) -> tuple[float, float]:
    return 960.0, _MSG_TOP + _messages_step(s) * (len(s["messages"]) - 1) + 40


def _messages_draw(pen: Pen, s: dict[str, Any], accent: str) -> None:
    parties = s["parties"]
    n = len(parties)
    xs = {p: 140 + k * (680 / (n - 1)) for k, p in enumerate(parties)}
    _, h = _messages_size(s)
    for party, x in xs.items():
        pen.line([(x, 58), (x, h - 6)], tint(NEUTRAL, 0.5), arrow=False, sw=1.5, dash=True)
        pen.rect(x - 80, 14, 160, 44, INK, r=8)
        pen.label(x - 76, 14, 152, 44, party, 15, True, WHITE, 10)
    step = _messages_step(s)
    for k, m in enumerate(s["messages"]):
        y = _MSG_TOP + k * step
        a, b = xs[m["from"]], xs[m["to"]]
        colour = INK if m["from"] == parties[0] else accent
        pen.line([(a, y), (b + (-4 if b > a else 4), y)], colour, sw=1.8)
        room = abs(b - a) - 24
        lines, size = figure_text.fit(m["name"], room, 20, 14, True, min_size=10)
        pen.text((a + b) / 2, y - 13, lines[:1], size, True, colour)
        if note := m.get("note"):
            lines, size = figure_text.fit(note, room, 18, 12, min_size=9.5)
            pen.text((a + b) / 2, y + 14, lines[:1], size, False, SUB)


# ── wide strips (a slide's figure band, about 5:1) ─────────────────────


def _cycle_band_size(s: dict[str, Any]) -> tuple[float, float]:
    return 960.0, 150.0


def _cycle_band(pen: Pen, s: dict[str, Any], accent: str) -> None:
    stages = s["stages"]
    n = len(stages)
    gap = 40.0
    bw = (920 - gap * (n - 1)) / n
    colours = hues(accent, n)
    for k, stage in enumerate(stages):
        x = 20 + k * (bw + gap)
        pen.rect(x, 10, bw, 78, colours[k], r=10)
        sub = stage.get("note") or stage.get("en") or ""
        if sub:
            pen.label(x, 12, bw, 36, stage["name"], 16, True, WHITE, 11)
            pen.label(x, 44, bw, 42, sub, 12, False, tint(colours[k], 0.8), 9)
        else:
            pen.label(x, 10, bw, 78, stage["name"], 17, True, WHITE, 11)
        if k:
            pen.line([(x - gap + 6, 49), (x - 6, 49)], NEUTRAL)
    first, last = 20 + bw / 2, 20 + (n - 1) * (bw + gap) + bw / 2
    pen.line([(last, 90), (last, 126), (first, 126), (first, 92)], tint(NEUTRAL, 0.2),
             dash=True)
    if center := s.get("center"):
        w = figure_text.width(center, 13, True) + 24
        pen.rect(480 - w / 2, 114, w, 24, WHITE, tint(NEUTRAL, 0.5), r=12)
        pen.text(480, 126, center, 13, True, INK)


def _house_band_size(s: dict[str, Any]) -> tuple[float, float]:
    return 960.0, 142 + 34 * len(s["base"])


def _house_band(pen: Pen, s: dict[str, Any], accent: str) -> None:
    deep = _deepened(accent)
    pen.poly([(480, 4), (900, 56), (60, 56)], deep)
    pen.label(290, 24, 380, 32, s["goal"], 15, True, WHITE, 10)
    pillars = s["pillars"]
    n, gap = len(pillars), 16.0
    pw = (800 - gap * (n - 1)) / n
    for i, name in enumerate(pillars):
        x = 80 + i * (pw + gap)
        pen.rect(x, 64, pw, 8, tint(deep, 0.55), r=2)
        pen.rect(x + 8, 72, pw - 16, 56, tint(deep, 0.88), r=0)
        pen.rect(x, 128, pw, 6, tint(deep, 0.55), r=2)
        pen.label(x + 8, 72, pw - 16, 56, name, 13, True, deep, 9)
    for k, row in enumerate(s["base"]):
        y = 142 + k * 34
        last = k == len(s["base"]) - 1
        fill = deep if last and len(s["base"]) > 1 else tint(deep, 0.25)
        pen.rect(60, y, 840, 30, fill, r=5)
        text = row["name"]
        if row.get("items"):
            text += "   " + " · ".join(row["items"])
        elif row.get("note"):
            text += "   " + row["note"]
        pen.label(60, y, 840, 30, text, 12.5, True, WHITE, 9)


def _compare_band_size(s: dict[str, Any]) -> tuple[float, float]:
    return 960.0, 40 + 36 * len(s["rows"])


def _compare_band(pen: Pen, s: dict[str, Any], accent: str) -> None:
    old, new = NEUTRAL, accent
    pen.rect(40, 0, 360, 30, tint(old, 0.85), r=8)
    pen.label(40, 0, 360, 30, s["sides"][0], 14, True, INK, 10)
    pen.rect(560, 0, 360, 30, new, r=8)
    pen.label(560, 0, 360, 30, s["sides"][1], 14, True, WHITE, 10)
    hot = s.get("hot") or {}
    for i, row in enumerate(s["rows"]):
        y = 40 + i * 36
        is_hot = row["aspect"] == hot.get("row")
        pen.rect(40, y, 360, 30, "#F8FAFC", tint(old, 0.55), r=6)
        pen.label(40, y, 360, 30, row["before"], 12.5, False, "#334155", 9)
        pen.rect(560, y, 360, 30, tint(new, 0.86) if is_hot else WHITE, new, r=6)
        after = row["after"] + (f"  ({hot['badge']})" if is_hot and hot.get("badge") else "")
        pen.label(560, y, 360, 30, after, 12.5, is_hot, new if is_hot else INK, 9)
        pen.rect(424, y + 3, 112, 24, tint(INK, 0.9), r=12)
        pen.label(424, y + 3, 112, 24, row["aspect"], 12, True, INK, 8)
        pen.line([(538, y + 15), (556, y + 15)], new)


# ── tall boxes (a slide whose figure takes the whole body) ────────────

_ROW_H, _ROW_GAP, _ROW_HEAD = 76.0, 14.0, 300.0


def _steps_rows_size(s: dict[str, Any]) -> tuple[float, float]:
    n = len(s["steps"])
    return 960.0, n * _ROW_H + (n - 1) * _ROW_GAP


def _steps_rows(pen: Pen, s: dict[str, Any], accent: str) -> None:
    """One row per step, top to bottom: a numbered arrow with the name, the note beside it."""
    items = s["steps"]
    n = len(items)
    for i, item in enumerate(items):
        y = i * (_ROW_H + _ROW_GAP)
        last = i == n - 1
        amount = 0.0 if last else 0.78 - i * (0.62 / max(1, n - 2))
        colour = accent if last else tint(accent, amount)
        ink = INK if amount > 0.45 else WHITE
        tip = 22.0
        pen.poly([(0, y), (_ROW_HEAD - tip, y), (_ROW_HEAD, y + _ROW_H / 2),
                  (_ROW_HEAD - tip, y + _ROW_H), (0, y + _ROW_H)], colour)
        number = f"Step {i + 1}" if s["lang"] == "en" else f"{i + 1}단계"
        pen.text(24, y + _ROW_H / 2, number, 15, True, ink, "l")
        pen.label(100, y, _ROW_HEAD - tip - 108, _ROW_H, item["name"], 21, True, ink, 13)
        x = _ROW_HEAD + 20
        pen.rect(x, y, 960 - x, _ROW_H, WHITE, tint(accent, 0.7), r=8, sw=1.2)
        if note := item.get("note"):
            lines, size = figure_text.fit(note, 960 - x - 48, _ROW_H - 16, 19, min_size=13)
            pen.text(x + 24, y + _ROW_H / 2, lines, size, False, INK, "l")


#: a box no wider than this for its height (a slide's whole body, about 3:1) takes the row
#: layout, when the steps carry notes; a strip under words (5:1 and wider) keeps the arrows
_TALL = 3.4
_TALL_KINDS = {"steps": (_steps_rows_size, _steps_rows)}


_DRAW = {"steps": (_steps_size, _steps), "cycle": (_cycle_size, _cycle),
         "house": (_house_size, _house), "compare": (_compare_size, _compare),
         "options": (_options_size, _options_draw), "axes": (_axes_size, _axes_draw),
         "messages": (_messages_size, _messages_draw)}
_BAND = {"cycle": (_cycle_band_size, _cycle_band), "house": (_house_band_size, _house_band),
         "compare": (_compare_band_size, _compare_band)}
#: a box this much wider than tall gets the strip layout
_WIDE = 3.2


def _layout(s: dict[str, Any], box: tuple[float, float] | None = None):
    wide = s.get("band") or (box is not None and box[0] / box[1] > _WIDE)
    if wide and s["kind"] in _BAND:
        return _BAND[s["kind"]]
    if box is not None and box[0] / box[1] < _TALL and s["kind"] in _TALL_KINDS and _noted(s):
        return _TALL_KINDS[s["kind"]]
    return _DRAW[s["kind"]]


def _noted(s: dict[str, Any]) -> bool:
    """Whether most parts carry a note (a row layout of bare names would be mostly empty)."""
    parts = s.get("steps") or []
    return bool(parts) and 2 * sum(bool(p.get("note")) for p in parts) >= len(parts)


def size(s: dict[str, Any], box: tuple[float, float] | None = None) -> tuple[float, float]:
    """The figure's own width and height in points, in the layout `box` would give it."""
    return _layout(s, box)[0](s)


def draw(slide, s: dict[str, Any], *, left: float, top: float, width: float, height: float,
         accent: Any, font: str = "", box: tuple[float, float] | None = None) -> None:
    """`s` drawn into the box at its own aspect, centred; a wide box takes the strip layout,
    a tall one the row layout. `box`: the slide box the layout is chosen for, when the
    drawing is made at the figure's own size (a picture of it)."""
    from app.services import diagram_shapes

    measure, paint = _layout(s, box or (width, height))
    pen = Pen(slide, left=left, top=top, width=width, height=height, size=measure(s),
              font=font or diagram_shapes.FONT)
    paint(pen, s, _hexed(accent))


__all__ = [
    "CHOICES", "KINDS", "MARK", "choose_messages", "chosen", "draw", "fill_messages", "grounded",
    "from_source", "size", "spec", "to_source",
]
