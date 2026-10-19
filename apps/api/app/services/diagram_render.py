"""Figure drawn with `diagram_shapes` (or `figure_kinds` for a fixed shape), PNG via sidecar;
None → caller keeps mermaid."""


from __future__ import annotations

import base64
import io
import logging

import httpx

from app.core.config import settings
from app.services import diagram_shapes, figure_kinds, flow_layout

log = logging.getLogger(__name__)

_TIMEOUT = httpx.Timeout(90.0, connect=5.0)
#: figure margin on its slide (pt)
_MARGIN = 18.0


def figure_size(graph: diagram_shapes.Graph) -> tuple[float, float]:
    """Figure box in points within a 16:9 slide; flows sized in rows for the text column."""
    if (rows := flow_layout.frame_size(graph)) is not None:
        return rows
    placed = diagram_shapes.layout(graph, 100, 100)
    layers = len({round(p.x if graph.direction == "LR" else p.y, 3) for p in placed.values()})
    across = max(
        sum(1 for q in placed.values()
            if round(q.x if graph.direction == "LR" else q.y, 3) == key)
        for key in {round(p.x if graph.direction == "LR" else p.y, 3) for p in placed.values()}
    )
    if diagram_shapes._compares(graph):
        width = min(900.0, max(560.0, 260.0 * layers))
        height = min(480.0, max(160.0, 64.0 * across + 40.0))
    elif graph.direction == "LR":
        width = min(900.0, max(420.0, 165.0 * layers))
        height = min(480.0, max(140.0, 82.0 * across + (30.0 if graph.groups else 0.0)))
    else:
        width = min(900.0, max(360.0, 190.0 * across))
        height = min(480.0, max(160.0, 88.0 * layers + (30.0 if graph.groups else 0.0)))
    return width, height


def figure_deck(source: str, accent: str = "#2563EB") -> bytes | None:
    """One-slide .pptx of the figure, or None if not drawable."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    slots = figure_kinds.from_source(source)
    graph = diagram_shapes.parse(source) if slots is None else None
    if slots is None and graph is None:
        return None
    if slots is not None:
        width, height = figure_kinds.size(slots)
        lead = 0.0
    else:
        width, height = figure_size(graph)
        lead = 8.0 if graph.groups else 0.0
    deck = Presentation()
    deck.slide_width = Pt(width + 2 * _MARGIN)
    deck.slide_height = Pt(height + 2 * _MARGIN + lead)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if slots is not None:
        figure_kinds.draw(slide, slots, left=_MARGIN, top=_MARGIN, width=width, height=height,
                          accent=accent)
    else:
        diagram_shapes.draw_pptx(
            slide, graph,
            left=_MARGIN, top=_MARGIN + lead,
            width=width, height=height, accent=accent,
        )
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


async def render_png(source: str, accent: str = "#2563EB") -> bytes | None:
    """Figure PNG from its python-pptx slide, or None."""
    base = settings.print_base_url.strip().rstrip("/")
    if not base:
        return None
    blob = figure_deck(source, accent)
    if blob is None:
        return None
    try:
        async with httpx.AsyncClient(timeout=_TIMEOUT) as client:
            response = await client.post(
                f"{base}/pptx/png", json={"pptx": base64.b64encode(blob).decode(), "dpi": 192}
            )
        response.raise_for_status()
    except httpx.HTTPError as exc:
        log.info("figure not rendered from its slide, keeping the mermaid: %s", exc)
        return None
    return _trimmed(response.content) if response.content else None


def _trimmed(png: bytes, pad: int = 24) -> bytes:
    """Picture cropped to its drawing plus `pad` px; the slide is oversized."""

    import PIL.Image
    import PIL.ImageChops

    with PIL.Image.open(io.BytesIO(png)) as raw:
        image = raw.convert("RGB")
    white = PIL.Image.new("RGB", image.size, (255, 255, 255))
    box = PIL.ImageChops.difference(image, white).getbbox()
    if not box:
        return png
    left, top, right, bottom = box
    image = image.crop((
        max(0, left - pad), max(0, top - pad),
        min(image.width, right + pad), min(image.height, bottom + pad),
    ))
    out = io.BytesIO()
    image.save(out, format="PNG", optimize=True)
    return out.getvalue()


def data_uri(png: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(png).decode()
