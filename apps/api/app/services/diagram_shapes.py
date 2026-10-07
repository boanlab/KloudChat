"""A mermaid flowchart drawn as native, editable .pptx shapes in the deck's font and accent.

Only flowcharts are read (`flowchart` / `graph`, any direction). Anything else, or a
chart past `MAX_NODES`, returns None and the caller keeps the picture.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

MAX_NODES = 18

_HEAD = re.compile(r"^\s*(?:flowchart|graph)\s+(TB|TD|BT|LR|RL)\b", re.I)
#: A node reference with an optional shape: id[..] id(..) id([..]) id((..)) id{..} id[(..)]
_NODE = re.compile(
    r"(?P<id>[A-Za-z0-9_가-힣]{1,40})\s{0,2}"
    r"(?:(?P<shape>\(\(|\(\[|\[\(|\[\[|\[|\(|\{\{|\{|>)"
    r"\s{0,2}\"?(?P<label>[^\]\)\}\"]{0,120}?)\"?\s{0,2}(?:\)\)|\]\)|\)\]|\]\]|\]|\)|\}\}|\}))?"
    r"(?:\([^()]{0,40}\))?"  # a stray 「(입력)」 after a label
    r"(?::::(?P<cls>[A-Za-z0-9_-]{1,30}))?"
)
_SUBGRAPH = re.compile(
    r"^subgraph\s+([A-Za-z0-9_가-힣]+)(?:\s{0,3}\[\s{0,3}\"?([^\]\"]{0,80})\"?\s{0,3}\])?"
)
_ARROW = re.compile(
    r"\s{0,3}(?P<arrow>-\.{1,3}->|-\.{1,3}-|==+>|--+>|---+|-->|==+)"
    r"(?:\s{0,2}\|(?P<label>[^|]{0,60})\|)?\s{0,3}"
)


@dataclass
class Node:
    id: str
    label: str
    shape: str = "box"  # box, round, diamond, circle
    hot: bool = False


@dataclass
class Edge:
    a: str
    b: str
    label: str = ""
    dashed: bool = False
    arrow: bool = True


@dataclass
class Graph:
    direction: str  # "LR" or "TB"
    nodes: dict[str, Node] = field(default_factory=dict)
    edges: list[Edge] = field(default_factory=list)
    groups: dict[str, tuple[str, list[str]]] = field(default_factory=dict)


def _shape_of(opener: str | None) -> str:
    return {
        "(": "round", "([": "round", "((": "circle", "{": "diamond", "{{": "diamond",
        "[(": "round", ">": "box", "[[": "box",
    }.get(opener or "", "box")


def _plain_label(raw: str) -> str:
    """A mermaid label as text: <br/> is a line break, other markup and entities go,
    and a line of 「key : value」 pairs keeps one space around its colon."""
    text = re.sub(r"<\s*br\s*/?\s*>", "\n", raw or "", flags=re.I)
    text = re.sub(r"<[^>]{1,40}>", "", text)
    text = text.replace("&nbsp;", " ").replace("&amp;", "&").replace("#quot;", '"')
    lines = [re.sub(r"\s{2,}", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line)


def parse(source: str) -> Graph | None:
    """The flowchart in `source`, or None when it is not one this module can draw."""
    lines = [ln.strip() for ln in (source or "").splitlines() if ln.strip()]
    if not lines or not (head := _HEAD.match(lines[0])):
        return None
    graph = Graph(direction="LR" if head.group(1).upper() in ("LR", "RL") else "TB")
    stack: list[str] = []
    hot: set[str] = set()

    def node(match: re.Match) -> str:
        nid = match.group("id")
        label = _plain_label(match.group("label") or "")
        existing = graph.nodes.get(nid)
        if existing is None:
            graph.nodes[nid] = Node(nid, label or nid, _shape_of(match.group("shape")))
        elif label:
            existing.label = label
            existing.shape = _shape_of(match.group("shape"))
        if match.group("cls") == "hot":
            hot.add(nid)
        if stack:
            members = graph.groups[stack[-1]][1]
            if nid not in members and nid not in graph.groups:
                members.append(nid)
        return nid

    for line in lines[1:]:
        if line.startswith("%%") or re.match(r"^(classDef|style|linkStyle|direction)\b", line):
            continue
        if m := re.match(r"^class\s+([\w,가-힣 ]+?)\s+hot\b", line):
            hot.update(x.strip() for x in m.group(1).split(","))
            continue
        if m := _SUBGRAPH.match(line):
            gid, label = m.group(1), (m.group(2) or m.group(1)).strip()
            graph.groups[gid] = (label, [])
            stack.append(gid)
            continue
        if line == "end":
            if stack:
                stack.pop()
            continue
        if m := re.match(r"^([A-Za-z0-9_가-힣]+):::(\w+)$", line):
            if m.group(2) == "hot":
                hot.add(m.group(1))
            continue
        # A chain: node (arrow node)*
        pos = 0
        first = _NODE.match(line, pos)
        if not first:
            continue
        current = node(first)
        pos = first.end()
        while pos < len(line):
            arrow = _ARROW.match(line, pos)
            if not arrow:
                break
            nxt = _NODE.match(line, arrow.end())
            if not nxt:
                break
            target = node(nxt)
            mark = arrow.group("arrow")
            graph.edges.append(Edge(
                current, target, _plain_label(arrow.group("label") or "").replace("\n", " "),
                dashed="." in mark, arrow=mark.endswith(">"),
            ))
            current, pos = target, nxt.end()

    for nid in hot:
        if nid in graph.nodes:
            graph.nodes[nid].hot = True
    # An edge to a group stands for its members: the first member receives it, the last sends.
    for edge in graph.edges:
        if edge.a in graph.groups and graph.groups[edge.a][1]:
            edge.a = graph.groups[edge.a][1][-1]
        if edge.b in graph.groups and graph.groups[edge.b][1]:
            edge.b = graph.groups[edge.b][1][0]
    for gid in graph.groups:
        graph.nodes.pop(gid, None)
    graph.edges = [
        e for e in graph.edges if e.a in graph.nodes and e.b in graph.nodes and e.a != e.b
    ]
    if not graph.nodes or len(graph.nodes) > MAX_NODES:
        return None
    if _compares(graph):
        graph.direction = "LR"
        # A comparison has no direction: no arrowheads, and links only between
        # counterparts on the same row.
        rows = {m: i for _g, (_l, members) in graph.groups.items() for i, m in enumerate(members)}
        side = {m: g for g, (_l, members) in graph.groups.items() for m in members}
        kept = []
        for edge in graph.edges:
            across = side.get(edge.a) != side.get(edge.b)
            if across and rows.get(edge.a) != rows.get(edge.b):
                continue
            edge.arrow = edge.arrow and not across
            kept.append(edge)
        graph.edges = kept
    return graph


@dataclass
class Placed:
    x: float  # centre, in the box's units
    y: float
    w: float
    h: float


def _compares(graph: Graph) -> bool:
    """Two or more groups holding nearly every node, with no edge inside any group: the
    groups are the sides of a comparison, not stages of one flow."""
    if len(graph.groups) < 2:
        return False
    group_of = {m: g for g, (_l, members) in graph.groups.items() for m in members}
    if len(group_of) < 0.7 * len(graph.nodes):
        return False
    return not any(
        group_of.get(e.a) is not None and group_of.get(e.a) == group_of.get(e.b)
        for e in graph.edges
    )


def layout(graph: Graph, width: float, height: float) -> dict[str, Placed]:
    """Node centres and sizes inside a `width` × `height` box: layers along the flow
    (longest path from the sources), ordered within a layer by their parents' positions."""
    ids = list(graph.nodes)
    preds: dict[str, list[str]] = {n: [] for n in ids}
    succs: dict[str, list[str]] = {n: [] for n in ids}
    for e in graph.edges:
        succs[e.a].append(e.b)
        preds[e.b].append(e.a)
    # Longest-path layers; a cycle is cut where it closes (visiting order).
    layer: dict[str, int] = {}
    visiting: set[str] = set()

    def depth(n: str) -> int:
        if n in layer:
            return layer[n]
        if n in visiting:
            return 0
        visiting.add(n)
        layer[n] = max((depth(p) + 1 for p in preds[n]), default=0)
        visiting.discard(n)
        return layer[n]

    for n in ids:
        depth(n)
    if _compares(graph):
        # A comparison: each side is a column of its own (LR layers), its items stacked in
        # source order; links between the sides are drawn across, not used for layers.
        group_index = {m: i for i, (_g, (_l, members)) in enumerate(graph.groups.items())
                       for m in members}
        for n in ids:
            layer[n] = group_index.get(n, len(graph.groups))
    # A group's members sit in consecutive layers of their own order when the source
    # put them side by side without edges: keep source order as the tie-break.
    count = max(layer.values()) + 1
    layers: list[list[str]] = [[] for _ in range(count)]
    for n in ids:
        layers[layer[n]].append(n)
    order = {n: i for i, n in enumerate(ids)}
    group_of = {m: g for g, (_label, members) in graph.groups.items() for m in members}
    group_rank = {g: i for i, g in enumerate(graph.groups)}
    for idx in range(1, count):
        position = {n: i for i, n in enumerate(layers[idx - 1])}
        layers[idx].sort(
            key=lambda n: (
                sum(position.get(p, 0) for p in preds[n]) / max(1, len(preds[n]))
                if preds[n] else len(position),
                order[n],
            )
        )
    # Members of one group sit next to each other in a layer, so its box can hold them.
    for lay in layers:
        before = {n: i for i, n in enumerate(lay)}
        lay.sort(key=lambda n: (group_rank.get(group_of.get(n, ""), -1), before[n]))
    across = max(len(lay) for lay in layers)
    flow_len = width if graph.direction == "LR" else height
    cross_len = height if graph.direction == "LR" else width
    # Spacing first, then sizes: the gap between layers carries the arrows (and their
    # labels), the gap between neighbours keeps boxes apart; what is left is the boxes.
    flow_gap = max(28.0, flow_len * 0.06) if count > 1 else 0.0
    if _compares(graph):
        # The two sides need room between their frames for the links and their labels.
        flow_gap = max(64.0, flow_len * 0.12)
    cross_gap = max(14.0, cross_len * 0.05) if across > 1 else 0.0
    cell_flow = (flow_len - flow_gap * (count - 1)) / count
    cell_cross = (cross_len - cross_gap * (across - 1)) / across
    # A box is never taller than a third of the band (LR) nor wider than its cell (TB).
    box_flow = min(cell_flow, 220.0)
    box_cross = min(cell_cross, (cross_len * 0.34) if graph.direction == "LR" else 260.0)
    placed: dict[str, Placed] = {}
    for i, lay in enumerate(layers):
        used = len(lay) * box_cross + (len(lay) - 1) * cross_gap
        start = (cross_len - used) / 2
        for j, n in enumerate(lay):
            along = i * (cell_flow + flow_gap) + cell_flow / 2
            cross = start + j * (box_cross + cross_gap) + box_cross / 2
            if graph.direction == "LR":
                w, h = box_flow, min(box_cross, 64.0)
            else:
                w, h = box_cross, min(box_flow, 56.0)
            if graph.nodes[n].shape == "circle":
                w = h = min(w, h) * 1.15
            if graph.direction == "LR":
                placed[n] = Placed(along, cross, w, h)
            else:
                placed[n] = Placed(cross, along, w, h)
    return placed


def _hex(colour: tuple[float, float, float] | str) -> tuple[int, int, int]:
    if isinstance(colour, str):
        c = colour.lstrip("#")
        return int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)
    values = tuple(colour)
    if any(v > 1 for v in values):  # already 0–255 (an RGBColor)
        return tuple(int(v) for v in values)  # type: ignore[return-value]
    return tuple(int(round(v * 255)) for v in values)  # type: ignore[return-value]


#: Body font of a figure (the app's own face); a viewer without it substitutes its sans.
FONT = "Pretendard"


_ICONS_DIR = Path(__file__).resolve().parent.parent / "assets" / "material_icons"

#: Words in a node's label and the Material Symbols icon each calls for; the first match
#: wins, so the narrower reading comes first. No match, no icon.
_ICON_WORDS: tuple[tuple[str, str], ...] = (
    (r"암호|복호|encrypt", "lock"), (r"키|인증|토큰|로그인|SSO|JWT|OAuth", "key"),
    (r"보안|방어|보호|차단|방화벽|위협|공격", "shield"), (r"정책|규정|컴플라이언스", "policy"),
    (r"검증|승인|확인|인증서", "verified_user"), (r"서버|백엔드|API|엔드포인트", "dns"),
    (r"데이터베이스|DB|저장소|스토리지|저장", "database"), (r"클라우드", "cloud"),
    (r"라우터|게이트웨이|프록시", "router"), (r"네트워크|통신|연결|망", "lan"),
    (r"모델|AI|LLM|추론|학습|신경망", "neurology"), (r"에이전트|봇|챗봇", "smart_toy"),
    (r"검색|탐지|조사|찾", "search"), (r"모니터링|관찰|감시", "monitoring"),
    (r"분석|통계|지표|KPI|측정", "analytics"), (r"성장|증가|추세|동향", "trending_up"),
    (r"로그|기록|감사|이력", "receipt_long"), (r"경고|위험|오류|장애|이상", "warning"),
    (r"완료|결과|성공|달성", "task_alt"), (r"일정|주차|기한|마감|날짜", "calendar_month"),
    (r"시간|대기|지연", "schedule"), (r"비용|예산|매출|가격|결제|돈|원\b", "payments"),
    (r"투자|자산|저축", "savings"), (r"은행|금융", "account_balance"),
    (r"설정|구성|환경 설정", "settings"), (r"배포|출시|론칭|런칭", "rocket_launch"),
    (r"코드|개발|구현|프로그램", "code"), (r"실험|테스트|시험", "science"),
    (r"알림|푸시", "notifications"), (r"메일|이메일", "mail"), (r"채팅|대화|메시지", "chat"),
    (r"웹|브라우저|사이트", "language"), (r"기기|모바일|단말|디바이스|앱", "devices"),
    (r"파일|업로드", "upload_file"), (r"폴더|자료실", "folder"), (r"이미지|사진|그림", "image"),
    (r"입력", "input"), (r"출력", "output"), (r"필터|정제|전처리", "filter_alt"),
    (r"동기화|갱신|순환|반복", "sync"), (r"문서|보고서|계획서|기획|자료|명세", "description"),
    (r"교육|수업|강의|학교|학생", "school"), (r"연구|논문|문헌", "menu_book"),
    (r"아이디어|개념|가설", "lightbulb"), (r"목표|목적", "flag"),
    (r"판매|매장|시장", "storefront"), (r"구매|주문|장바구니", "shopping_cart"),
    (r"광고|마케팅|캠페인|홍보", "campaign"), (r"배송|물류|운송", "local_shipping"),
    (r"공장|생산|제조", "factory"), (r"하드웨어|메모리|GPU|NPU|칩", "memory"),
    (r"설계|아키텍처|구조", "architecture"), (r"엔지니어|정비|유지보수", "engineering"),
    (r"계산|수식|산정", "calculate"), (r"법|계약|소송|판결", "gavel"),
    (r"건강|안전", "health_and_safety"), (r"약|복용|처방", "medication"),
    (r"여행|항공|비행", "flight"), (r"집|주택|전세|임대", "home"),
    (r"팀|조직|그룹|고객사", "group"), (r"사용자|고객|사람|직원|관리자|교사|학부모", "person"),
    (r"허브|통합|중앙", "hub"),
)
_ICON_RES = [(re.compile(p, re.I), name) for p, name in _ICON_WORDS]


def icon_for(label: str) -> str | None:
    """The Material Symbols icon a node's label calls for, or None."""
    for pattern, name in _ICON_RES:
        if pattern.search(label or "") and (_ICONS_DIR / f"{name}.svg").exists():
            return name
    return None


def _icon_blobs(name: str, rgb: tuple[int, int, int]) -> tuple[bytes, bytes]:
    """`(svg, png)` of the icon in one solid colour: the SVG is what PowerPoint shows,
    the PNG the raster fallback it requires beside it."""
    import io

    import PIL.Image

    colour = "#{:02X}{:02X}{:02X}".format(*rgb)
    svg = (_ICONS_DIR / f"{name}.svg").read_text()
    svg = svg.replace("<svg ", f'<svg fill="{colour}" ', 1).encode()
    with PIL.Image.open(_ICONS_DIR / f"{name}.png") as raw:
        alpha = raw.convert("RGBA").getchannel("A")
    tinted = PIL.Image.new("RGBA", alpha.size, (*rgb, 255))
    tinted.putalpha(alpha)
    buffer = io.BytesIO()
    tinted.save(buffer, format="PNG")
    return svg, buffer.getvalue()


def _add_svg_picture(slide, svg: bytes, png: bytes, left, top, size) -> None:
    """A picture whose blip is the PNG with the SVG attached (`asvg:svgBlip`), as
    PowerPoint itself writes an inserted SVG: vector where read, raster elsewhere."""
    import io

    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml.ns import qn

    picture = slide.shapes.add_picture(io.BytesIO(png), left, top, size, size)
    package = slide.part.package
    existing = {str(part.partname) for part in package.iter_parts()}
    index = 1
    while f"/ppt/media/icon{index}.svg" in existing:
        index += 1
    svg_part = Part(PackURI(f"/ppt/media/icon{index}.svg"), "image/svg+xml", package, svg)
    rid = slide.part.relate_to(svg_part, RT.IMAGE)
    blip = picture._element.find(".//" + qn("a:blip"))
    ext_list = blip.makeelement(qn("a:extLst"), {})
    ext = ext_list.makeelement(qn("a:ext"), {"uri": "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"})
    svg_blip = ext.makeelement(
        "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip",
        {qn("r:embed"): rid},
        nsmap={"asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main"},
    )
    ext.append(svg_blip)
    ext_list.append(ext)
    blip.append(ext_list)


def _tighten(text: str) -> str:
    """「AI 보안」, 「소재: 사내」 without the space at a Hangul–Latin boundary."""
    text = re.sub(r"(?<=[A-Za-z0-9:)\]%])[ ](?=[가-힣])", "", text)
    text = re.sub(r"(?<=[가-힣])[ ](?=[A-Za-z0-9(\[:])", "", text)
    return re.sub(r"(?<=[A-Za-z0-9)\]]) (?=:)", "", text)


def _sides(a: Placed, b: Placed, lr: bool) -> tuple[str, str]:
    """Which side a line leaves `a` by and enters `b` by: along the flow between layers,
    across it between neighbours in one layer."""
    if lr:
        if abs(a.x - b.x) < 1:
            return ("bottom", "top") if b.y > a.y else ("top", "bottom")
        return ("right", "left") if b.x > a.x else ("left", "right")
    if abs(a.y - b.y) < 1:
        return ("right", "left") if b.x > a.x else ("left", "right")
    return ("bottom", "top") if b.y > a.y else ("top", "bottom")


_STEP = 4.0  # grid pitch in points
_OUT = {"right": (1, 0), "left": (-1, 0), "bottom": (0, 1), "top": (0, -1)}


def _route(
    a_box, a_side: str, b_box, b_side: str, *, obstacles, frames, used, bounds, owner: str = ""
) -> list[tuple[float, float]]:
    """An orthogonal A* path between the middles of two box sides, clear of other boxes.

    Bends and shared runs cost extra; falls back to an elbow. Taken cells go into `used`.
    """
    import heapq

    def mid(box, side):
        x0, y0, x1, y1 = box
        return {"right": (x1, (y0 + y1) / 2), "left": (x0, (y0 + y1) / 2),
                "bottom": ((x0 + x1) / 2, y1), "top": ((x0 + x1) / 2, y0)}[side]

    bx0, by0, bx1, by1 = bounds
    cols, rows = int((bx1 - bx0) / _STEP) + 1, int((by1 - by0) / _STEP) + 1
    cell = lambda x, y: (round((x - bx0) / _STEP), round((y - by0) / _STEP))  # noqa: E731
    point = lambda c: (bx0 + c[0] * _STEP, by0 + c[1] * _STEP)  # noqa: E731
    blocked: set[tuple[int, int]] = set()
    for x0, y0, x1, y1 in [*obstacles, a_box, b_box]:
        c0, c1 = cell(x0 - 6, y0 - 6), cell(x1 + 6, y1 + 6)
        blocked.update((i, j) for i in range(c0[0], c1[0] + 1) for j in range(c0[1], c1[1] + 1))
    for x0, y0, x1, _ in frames:  # the group's title strip
        c0, c1 = cell(x0 + 4, y0), cell(x1 - 4, y0 + 18)
        blocked.update((i, j) for i in range(c0[0], c1[0] + 1) for j in range(c0[1], c1[1] + 1))
    outline: set[tuple[int, int]] = set()
    for x0, y0, x1, y1 in frames:
        c0, c1 = cell(x0, y0), cell(x1, y1)
        outline.update((i, c0[1]) for i in range(c0[0], c1[0] + 1))
        outline.update((i, c1[1]) for i in range(c0[0], c1[0] + 1))
        outline.update((c0[0], j) for j in range(c0[1], c1[1] + 1))
        outline.update((c1[0], j) for j in range(c0[1], c1[1] + 1))

    sx, sy = mid(a_box, a_side)
    gx, gy = mid(b_box, b_side)
    # A short stub straight out of each side, so the line meets the box square on.
    so, go = _OUT[a_side], _OUT[b_side]
    start = cell(sx + so[0] * 10, sy + so[1] * 10)
    goal = cell(gx + go[0] * 10, gy + go[1] * 10)
    for c in (start, goal):
        blocked.discard(c)
    # The stubs' own cells are free even inside the 6-point margin.
    for (x, y), (dx, dy) in (((sx, sy), so), ((gx, gy), go)):
        for k in range(0, 4):
            blocked.discard(cell(x + dx * k * _STEP / 1.5, y + dy * k * _STEP / 1.5))

    def h(c):
        return abs(c[0] - goal[0]) + abs(c[1] - goal[1])

    first = so
    frontier = [(h(start), 0.0, start, first)]
    best = {(start, first): 0.0}
    came: dict = {}
    found = None
    while frontier:
        _, cost, c, d = heapq.heappop(frontier)
        if c == goal:
            found = (c, d)
            break
        if cost > best.get((c, d), float("inf")):
            continue
        for nd in ((1, 0), (-1, 0), (0, 1), (0, -1)):
            if nd == (-d[0], -d[1]):
                continue
            n = (c[0] + nd[0], c[1] + nd[1])
            if not (0 <= n[0] < cols and 0 <= n[1] < rows) or n in blocked:
                continue
            # Lines from one box share their trunk; another box's line is avoided.
            shared = used.get(n)
            step = 1.0 + (8.0 if nd != d else 0.0)
            step += (0.0 if owner in shared else 3.0) if shared else 0.0
            step += 2.0 if n in outline else 0.0
            new = cost + step
            if new < best.get((n, nd), float("inf")):
                best[(n, nd)] = new
                came[(n, nd)] = (c, d)
                heapq.heappush(frontier, (new + h(n), new, n, nd))
    if found is None:
        # No way round: an elbow through the middle.
        mxp = (sx + gx) / 2
        return [(sx, sy), (mxp, sy), (mxp, gy), (gx, gy)] if a_side in ("left", "right") else [
            (sx, sy), (sx, (sy + gy) / 2), (gx, (sy + gy) / 2), (gx, gy)]
    cells = [found[0]]
    state = found
    while state in came:
        state = came[state]
        cells.append(state[0])
    cells.reverse()
    for c in cells:
        used.setdefault(c, set()).add(owner)
    mids = [point(c) for c in cells]
    # The first and last straight runs sit exactly on the sides' middles (the grid is
    # 4 points coarse), so the line meets each box square on, with no jog.
    horizontal_a, horizontal_b = a_side in ("left", "right"), b_side in ("left", "right")
    k = 0
    while k < len(cells) and (cells[k][1] if horizontal_a else cells[k][0]) == (
        cells[0][1] if horizontal_a else cells[0][0]
    ):
        mids[k] = (mids[k][0], sy) if horizontal_a else (sx, mids[k][1])
        k += 1
    k = len(cells) - 1
    while k >= 0 and (cells[k][1] if horizontal_b else cells[k][0]) == (
        cells[-1][1] if horizontal_b else cells[-1][0]
    ):
        mids[k] = (mids[k][0], gy) if horizontal_b else (gx, mids[k][1])
        k -= 1
    pts = [(sx, sy), *mids, (gx, gy)]
    corners = [pts[0]]
    for prev, here, nxt in zip(pts, pts[1:], pts[2:], strict=False):
        straight = (abs(prev[0] - here[0]) < 0.01 and abs(here[0] - nxt[0]) < 0.01) or (
            abs(prev[1] - here[1]) < 0.01 and abs(here[1] - nxt[1]) < 0.01)
        if not straight:
            corners.append(here)
    corners.append(pts[-1])
    return corners


def _tint(rgb: tuple[int, int, int], amount: float) -> tuple[int, int, int]:
    """`rgb` mixed toward white by `amount` (0 keeps it, 1 is white)."""
    return tuple(int(round(c + (255 - c) * amount)) for c in rgb)  # type: ignore[return-value]


def _font_size(label: str, width: float, height: float) -> int:
    """The largest size from 14 pt down to 8 pt at which `label` wraps inside the box."""
    for size in range(14, 7, -1):
        per_line = max(1, int((width - 8) / (size * 1.0)))  # a Hangul glyph is about 1 em
        lines = sum(-(-max(1, len(part)) // per_line) for part in label.split("\n"))
        if lines * size * 1.25 <= height - 6:
            return size
    return 8


def draw_pptx(
    slide,
    graph: Graph,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    accent: tuple[float, float, float] | str,
    font: str = "",
    pt=None,
    tight: bool = False,
) -> None:
    """Native shapes and routed lines for `graph` inside the box, in points (960 × 540 slide)."""
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    pt = pt or Pt
    font = font or FONT
    primary_rgb = _hex(accent)
    # Palette: primary (the deck's accent) for the emphasised step, secondary (a pale
    # tint of it) for the other boxes, neutral tones for groups, lines and text.
    primary = RGBColor(*primary_rgb)
    secondary = RGBColor(*_tint(primary_rgb, 0.86))
    group_line = RGBColor(*_tint(primary_rgb, 0.6))
    text_dark = RGBColor(0x1F, 0x29, 0x37)
    text_light = RGBColor(0xFF, 0xFF, 0xFF)
    line_grey = RGBColor(0x94, 0xA3, 0xB8)
    label_grey = RGBColor(0x47, 0x55, 0x69)
    placed = layout(graph, width, height)

    def flat(shape) -> None:
        """No shadow, glow or bevel: an empty effect list overrides the theme's."""
        sppr = shape._element.spPr
        for old in sppr.findall(qn("a:effectLst")):
            sppr.remove(old)
        # The theme style reference brings its own effects (LibreOffice draws a shadow
        # from it); every colour here is set directly, so the reference goes.
        for style in shape._element.findall(qn("p:style")):
            shape._element.remove(style)
        effects = sppr.makeelement(qn("a:effectLst"), {})
        # The effect list follows the fill and line in spPr's schema order.
        line = sppr.find(qn("a:ln"))
        if line is not None:
            line.addnext(effects)
        else:
            sppr.append(effects)

    def set_text(shape, text: str, size: int, colour, *, bold: bool = False,
                 anchor=MSO_ANCHOR.MIDDLE) -> None:
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = pt(4)
        frame.margin_top = frame.margin_bottom = pt(2)
        # A label's own line breaks (mermaid's <br/>) are paragraphs, each centred.
        for n, line in enumerate(text.split("\n")):
            para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
            para.alignment = PP_ALIGN.CENTER
            run = para.add_run()
            # LibreOffice adds its own gap where Hangul meets Latin letters, digits or
            # punctuation; the written space there would make it a double one.
            run.text = _tighten(line) if tight else line
            run.font.name = font
            run.font.size = pt(size)
            run.font.bold = bold
            run.font.color.rgb = colour
            # East Asian glyphs take the font from <a:ea>, not <a:latin>.
            rpr = run._r.get_or_add_rPr()
            ea = rpr.find(qn("a:ea"))
            if ea is None:
                ea = rpr.makeelement(qn("a:ea"), {})
                rpr.append(ea)
            ea.set("typeface", font)

    # Group backgrounds first, so they sit under their members.
    pad, title_room = 8.0, 16.0
    frames: list[tuple[float, float, float, float]] = []
    for label, members in graph.groups.values():
        inside = [placed[m] for m in members if m in placed]
        if not inside:
            continue
        x0 = min(left + p.x - p.w / 2 for p in inside) - pad
        y0 = min(top + p.y - p.h / 2 for p in inside) - pad - title_room
        x1 = max(left + p.x + p.w / 2 for p in inside) + pad
        y1 = max(top + p.y + p.h / 2 for p in inside) + pad
        others = [placed[n] for n in placed if n not in members]
        if any(
            x0 < left + o.x + o.w / 2 and left + o.x - o.w / 2 < x1
            and y0 < top + o.y + o.h / 2 and top + o.y - o.h / 2 < y1
            for o in others
        ):
            continue  # it would cover a shape outside it: no box rather than a wrong one
        x0, y0 = max(x0, left - pad), max(y0, top - pad - title_room)
        frames.append((x0, y0, x1, y1))
        g = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, pt(x0), pt(y0), pt(x1 - x0), pt(y1 - y0)
        )
        g.adjustments[0] = 0.06
        # Flat on white: the group is a thin dashed outline on the white ground.
        g.fill.solid()
        g.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        g.line.color.rgb = group_line
        g.line.width = pt(0.75)
        g.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        flat(g)
        set_text(g, label, 9, label_grey, bold=True, anchor=MSO_ANCHOR.TOP)

    shapes = {}
    kinds = {
        "box": MSO_SHAPE.ROUNDED_RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE,
        "diamond": MSO_SHAPE.DIAMOND, "circle": MSO_SHAPE.OVAL,
    }
    for nid, node in graph.nodes.items():
        p = placed[nid]
        x, y = left + p.x - p.w / 2, top + p.y - p.h / 2
        s = slide.shapes.add_shape(
            kinds.get(node.shape, MSO_SHAPE.ROUNDED_RECTANGLE), pt(x), pt(y), pt(p.w), pt(p.h)
        )
        if node.shape in ("box", "round"):
            s.adjustments[0] = 0.12 if node.shape == "box" else 0.3
        s.fill.solid()
        s.fill.fore_color.rgb = primary if node.hot else secondary
        s.line.fill.background()  # no outline: the fill carries the shape
        flat(s)
        # A diamond's text area is its inner half: size the text for that.
        text_w = p.w * (0.6 if node.shape == "diamond" else 0.9 if node.shape == "circle" else 1)
        text_h = p.h * (0.6 if node.shape == "diamond" else 0.9 if node.shape == "circle" else 1)
        icon = icon_for(node.label) if node.shape in ("box", "round") else None
        icon_size = min(p.h * 0.42, 22.0)
        if icon and p.w - icon_size - 16 >= 48:
            # Icon on the left, text beside it: the text frame starts after the icon.
            text_w -= icon_size + 6
        else:
            icon = None
        set_text(
            s, node.label, _font_size(node.label, text_w, text_h),
            text_light if node.hot else text_dark, bold=node.hot,
        )
        if icon:
            s.text_frame.margin_left = pt(8 + icon_size + 4)
            ink = (0xFF, 0xFF, 0xFF) if node.hot else primary_rgb
            svg, png = _icon_blobs(icon, ink)
            _add_svg_picture(
                slide, svg, png, pt(x + 8), pt(y + (p.h - icon_size) / 2), pt(icon_size)
            )
        shapes[nid] = s

    # Open polylines routed around boxes; a connector would run straight through them.

    lr = graph.direction == "LR"
    boxes = {
        nid: (left + q.x - q.w / 2, top + q.y - q.h / 2,
              left + q.x + q.w / 2, top + q.y + q.h / 2)
        for nid, q in placed.items()
    }
    used: dict[tuple[int, int], set[str]] = {}
    placed_labels: list[tuple[float, float, float, float]] = []
    for edge in graph.edges:
        a_box, b_box = boxes[edge.a], boxes[edge.b]
        a, b = placed[edge.a], placed[edge.b]
        sa, sb = _sides(a, b, lr)
        path = _route(
            a_box, sa, b_box, sb,
            obstacles=[box for nid, box in boxes.items() if nid not in (edge.a, edge.b)],
            frames=frames, used=used, owner=edge.a,
            bounds=(left - 30, top - 40, left + width + 30, top + height + 40),
        )
        builder = slide.shapes.build_freeform(pt(path[0][0]), pt(path[0][1]), scale=1.0)
        builder.add_line_segments([(pt(x), pt(y)) for x, y in path[1:]], close=False)
        line = builder.convert_to_shape()
        line.fill.background()
        line.line.color.rgb = line_grey
        line.line.width = pt(1.25)
        flat(line)
        if edge.dashed:
            line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if edge.arrow:
            ln = line.line._get_or_add_ln()
            head = {"type": "triangle", "w": "med", "len": "med"}
            ln.append(ln.makeelement(qn("a:tailEnd"), head))
        if edge.label:
            # Beside a run of the line, longest runs first, at the first spot that covers
            # no box and no label already placed.
            tw = min(150.0, 9.0 * len(edge.label) + 12)
            runs = sorted(
                zip(path, path[1:], strict=False),
                key=lambda s: -(abs(s[1][0] - s[0][0]) + abs(s[1][1] - s[0][1])),
            )
            spots = []
            for (x0, y0), (x1, y1) in runs:
                mx, my = (x0 + x1) / 2, (y0 + y1) / 2
                if abs(y1 - y0) < 1:  # horizontal: above, then below
                    spots += [(mx - tw / 2, my - 16), (mx - tw / 2, my + 3)]
                else:  # vertical: right, then left
                    spots += [(mx + 4, my - 7), (mx - tw - 4, my - 7)]
            def clear(x: float, y: float, tw: float = tw) -> bool:
                return not any(
                    x < bx1 and x + tw > bx0 and y < by1 and y + 14 > by0
                    for bx0, by0, bx1, by1 in [*boxes.values(), *placed_labels]
                )
            x, y = next((spot for spot in spots if clear(*spot)), spots[0])
            placed_labels.append((x, y, x + tw, y + 14))
            tb = slide.shapes.add_textbox(pt(x), pt(y), pt(tw), pt(14))
            set_text(tb, edge.label, 8, label_grey)
