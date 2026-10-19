"""A structure (a cluster's nodes, a service's tiers, a short figure that forks and joins) is laid
out by ELK in the printer and drawn as shapes: each group's frame holds its own members, the link
back is the one the writer drew back, and without the printer the built-in layouts draw it.

The printer's answers are recorded in tests/fixtures/elk (regenerate by posting
`elk_layout._requests` of these figures to the printer's /layout)."""

from __future__ import annotations

import asyncio
import io
import json
from pathlib import Path

import httpx
import pytest
from pptx import Presentation
from pptx.util import Pt

from app.services import diagram_render, elk_layout, flow_layout
from app.services import diagram_shapes as ds

_RECORDED = Path(__file__).parent / "fixtures" / "elk"

_CLUSTER = """flowchart LR
    user[사용자] -->|kubectl| api
    subgraph cp["컨트롤 플레인"]
        direction TB
        api(API 서버)
        etcd[(etcd)]
        sched(스케줄러)
        api -.-> etcd
        api -.-> sched
    end
    subgraph n1["워커 노드 1"]
        direction TB
        kl1(kubelet)
        cr1(컨테이너 런타임)
        p11(파드 1a)
        p12(파드 1b)
        kl1 -.-> cr1
        cr1 -.-> p11
        cr1 -.-> p12
    end
    subgraph n2["워커 노드 2"]
        direction TB
        kl2(kubelet)
        cr2(컨테이너 런타임)
        kl2 -.-> cr2
    end
    lb(로드밸런서) --> svc(서비스)
    svc -.->|라우팅| kl1
    svc -.->|라우팅| kl2
    kl1 -->|통신| api
    kl2 -->|통신| api"""

_DISTILL = """flowchart LR
  batch[입력 배치] --> t(교사 모델)
  batch --> s(학생 모델):::hot
  t -->|T로 부드럽게| pt[소프트 확률]
  s --> ps[학생 확률]
  pt --> kl{KL 발산}
  ps --> kl
  lbl[정답 라벨] --> ce{교차 엔트로피}
  ps --> ce
  kl -->|λ·| L[전체 손실]
  ce -->|+| L
  L -.->|역전파| s"""

_CHECK = """flowchart LR
  sum[요약문] --> dec(원자적 주장 분해) --> list[주장 목록]
  list --> num(수치 정합성 검사기)
  list --> ret(원문 문장 검색)
  ret --> nli(NLI 함의 판정)
  nli --> gate{불확실성 게이트}
  gate -.->|불확실| ret
  gate --> conf(충돌 확인)
  conf -.->|충돌 시 NEI| score(주장별 점수)
  num --> score
  score --> avg(가중 평균) --> out[전체 사실 일관성]:::hot"""


@pytest.fixture(autouse=True)
def _no_layouts_left_over():
    """Layouts kept for one test are not drawn by another."""
    yield
    elk_layout._CACHE.clear()


def _printer(monkeypatch, sent: list | None = None, fail: bool = False) -> None:
    """The printer answering /layout with the recorded layouts of the graph sent."""
    recorded = {name: json.loads((_RECORDED / f"{name}.json").read_text())
                for name in ("cluster", "distill")}

    def handler(request: httpx.Request) -> httpx.Response:
        if fail:
            return httpx.Response(503)
        graphs = json.loads(request.content)["graphs"]
        if sent is not None:
            sent.extend(graphs)
        ids = {c["id"] for c in graphs[0]["children"]}
        laid = recorded["cluster" if "g:cp" in ids else "distill"]
        return httpx.Response(200, json={"graphs": laid[:len(graphs)]})

    real = httpx.AsyncClient

    def client(**kwargs):
        return real(transport=httpx.MockTransport(handler), **kwargs)

    monkeypatch.setattr(elk_layout.httpx, "AsyncClient", client)
    monkeypatch.setattr(elk_layout.settings, "print_base_url", "http://print.test")
    elk_layout._CACHE.clear()


def test_a_cluster_and_a_short_figure_are_drawn_by_elk_but_flows_are_not():
    assert elk_layout.applies(ds.parse(_CLUSTER))
    assert elk_layout.applies(ds.parse(_DISTILL))
    assert not elk_layout.applies(ds.parse(_CHECK)) and flow_layout.applies(ds.parse(_CHECK))
    phases = ds.parse("""flowchart LR
  subgraph prep["준비"]
    a(수집) --> b(정제)
  end
  subgraph train["학습"]
    c(사전 학습) --> d(미세 조정)
  end
  prep --> train""")
    assert not elk_layout.applies(phases) and flow_layout.applies(phases)


def test_each_group_frame_holds_its_own_members_only(monkeypatch):
    _printer(monkeypatch)
    graph = ds.parse(_CLUSTER)
    asyncio.run(elk_layout.prepare([_CLUSTER]))
    width, height = elk_layout.frame_size(graph)
    assert width <= elk_layout.FIGURE[0] + 0.5 and height <= elk_layout.FIGURE[1] + 0.5
    laid = elk_layout.place(graph, width, height)
    assert laid is not None and len(laid.frames) == 3
    titles = {title: members for title, members in graph.groups.values()}
    for title, (x0, y0, x1, y1) in laid.frames:
        for nid, (bx0, by0, bx1, by1) in laid.boxes.items():
            inside = x0 <= bx0 and bx1 <= x1 and y0 <= by0 and by1 <= y1
            assert inside == (nid in titles[title]), (title, nid)
    # every link is drawn, square, from its source's box to its target's
    for _edge, path, _label in laid.links:
        assert len(path) >= 2
        for (ax, ay), (bx, by) in zip(path, path[1:], strict=False):
            assert abs(ax - bx) < 0.01 or abs(ay - by) < 0.01


def test_the_link_back_is_the_one_the_writer_drew_back(monkeypatch):
    sent: list = []
    _printer(monkeypatch, sent)
    graph = ds.parse(_DISTILL)
    assert elk_layout.feedback(graph) == {("L", "s")}
    asyncio.run(elk_layout.prepare([_DISTILL]))
    # sent to ELK reversed, so the cycle is cut there and not at a link into the loss
    back = next(e for e in sent[0]["edges"] if e["id"] == f"e{len(graph.edges) - 1}")
    assert back["sources"] == ["s"] and back["targets"] == ["L"]
    laid = elk_layout.place(graph, *elk_layout.FIGURE)
    assert laid is not None and laid.text_pt >= elk_layout.LEGIBLE
    edge, path, _label = laid.links[-1]

    def near(point, box):
        x0, y0, x1, y1 = box
        return x0 - 1 <= point[0] <= x1 + 1 and y0 - 1 <= point[1] <= y1 + 1

    # drawn from the loss back to the student, the arrow at the student
    assert (edge.a, edge.b) == ("L", "s")
    assert near(path[0], laid.boxes["L"]) and near(path[-1], laid.boxes["s"])


def test_boxes_are_sized_to_their_labels():
    graph = ds.parse("flowchart LR\n  a[짧음] --> b[조금 더 긴 상자 이름입니다] --> c{판단}")
    height = elk_layout._height(graph, elk_layout.LINE_W)
    short = elk_layout._box(graph, "a", height, elk_layout.LINE_W)
    long = elk_layout._box(graph, "b", height, elk_layout.LINE_W)
    diamond = elk_layout._box(graph, "c", height, elk_layout.LINE_W)
    assert short[0] == elk_layout.MIN_W and long[0] > short[0] and short[1] == long[1]
    assert diamond[1] > short[1]
    # a compact layout wraps the long label onto more lines in a narrower box
    narrow = elk_layout.COMPACT[0]
    assert elk_layout._box(graph, "b", height, narrow)[0] < long[0]
    assert len(elk_layout.lines(graph.nodes["b"].label, elk_layout.TEXT_PT, False, narrow)) > 1


def test_the_figure_is_drawn_with_the_elk_layout(monkeypatch):
    _printer(monkeypatch)
    asyncio.run(elk_layout.prepare([_CLUSTER]))
    blob = diagram_render.figure_deck(_CLUSTER)
    deck = Presentation(io.BytesIO(blob))
    texts = [s.text_frame.text for s in deck.slides[0].shapes if s.has_text_frame]
    for title in ("컨트롤 플레인", "워커 노드 1", "워커 노드 2", "kubelet", "통신", "라우팅"):
        assert any(title in t for t in texts), title
    width, height = elk_layout.frame_size(ds.parse(_CLUSTER))
    assert abs(deck.slide_width - Pt(width + 2 * diagram_render._MARGIN)) < Pt(1)


def test_without_the_printer_the_built_in_layout_draws_it(monkeypatch):
    _printer(monkeypatch, fail=True)
    graph = ds.parse(_CLUSTER)
    asyncio.run(elk_layout.prepare([_CLUSTER]))
    assert elk_layout.place(graph, *elk_layout.FIGURE) is None
    assert elk_layout.frame_size(graph) is None
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    ds.draw_pptx(slide, graph, left=18, top=18, width=600, height=300, accent="#2563EB")
    assert len(slide.shapes) > len(graph.nodes)


def test_a_node_named_by_a_later_group_stays_in_the_group_that_defines_it():
    graph = ds.parse("""flowchart LR
    subgraph app["애플리케이션"]
        api(API 서버) --> ord(주문 서비스)
    end
    subgraph data["데이터"]
        db[(DB)]
        ord --> db
    end""")
    assert graph.groups["app"][1] == ["api", "ord"] and graph.groups["data"][1] == ["db"]
    # named first by a bare link, the node goes where it is defined
    moved = ds.parse("flowchart LR\n  subgraph a[\"A\"]\n    x --> y\n  end\n"
                     "  subgraph b[\"B\"]\n    y(정의)\n  end")
    assert moved.groups["a"][1] == ["x"] and moved.groups["b"][1] == ["y"]
