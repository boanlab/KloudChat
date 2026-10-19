"""Comparison columns with few, short items draw one row per item, aligned across the columns,
so the panels are not mostly empty."""

import io

from pptx import Presentation

from app.services import deck_export, deck_type, fonts

_SHORT = [{"title": "기존", "items": ["긴 테이블 120석", "차가운 형광등", "식사만"]},
          {"title": "리뉴얼", "items": ["세 구역 120석", "따뜻한 색 조명", "식사·회의·휴식"]}]
_LONG = [{"title": "A안", "items": ["긴 설명 " * 30, "짧다"]},
         {"title": "B안", "items": ["짧다", "짧다"]}]


def _rows(columns) -> float:
    pairs = [(c["title"], c["items"]) for c in columns]
    return deck_export._column_rows(pairs, fonts.korean("gothic"), 1.0, 360, 240)


def test_rows_take_the_size_their_words_allow():
    assert _rows(_SHORT) == deck_type.TYPE["columnRow"]
    assert _rows(_LONG) == 0
    assert deck_export._column_rows([("가", ["나"])], None, 1.0, 360, 240) == 0


def _runs(layout: str, columns) -> dict[str, tuple[float, float]]:
    data = deck_export.to_pptx("덱", [{"layout": "title", "title": "덱"},
                                      {"layout": layout, "title": "비교", "columns": columns}])
    slide = Presentation(io.BytesIO(data)).slides[1]
    out = {}
    for shape in slide.shapes:
        if shape.has_text_frame:
            for run in (r for p in shape.text_frame.paragraphs for r in p.runs):
                out[run.text] = (run.font.size.pt if run.font.size else 0, shape.top)
    return out


def test_the_same_point_sits_at_the_same_height_in_both_columns():
    runs = _runs("compare-2", _SHORT)
    for left, right in zip(_SHORT[0]["items"], _SHORT[1]["items"], strict=True):
        assert runs[left][0] == runs[right][0] == deck_type.TYPE["columnRow"]
        assert runs[left][1] == runs[right][1]


def test_pros_and_cons_keep_their_list():
    columns = [{"title": "장점", "items": ["빠르다", "싸다"]},
               {"title": "단점", "items": ["무겁다"]}]
    runs = _runs("pros-cons", columns)
    assert runs["빠르다"][0] != deck_type.TYPE["columnRow"]
