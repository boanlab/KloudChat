"""A lecture deck's structure, relation and flow slides get diagrams without the word 「구조도」,
an edit asking for one draws it, relations are not chains, labels follow the deck's language and
band names are fitted, never cut."""

from __future__ import annotations

import inspect
import json

import pytest

from app.services import deck, deck_export, diagram, diagrams, figure_text, illustrations

_K8S = (
    "신입 개발자 대상 1시간 강의용 슬라이드를 만들어줘. 주제는 'Kubernetes 구조와 동작 원리'. "
    "컨트롤 플레인/워커 노드 구조, 파드 생성 요청이 처리되는 흐름, 핵심 오브젝트(Deployment, "
    "Service 등)를 설명해줘."
)


def _slides() -> list[dict]:
    return [
        {"id": "s0", "title": "Kubernetes 구조와 동작 원리", "layout": "title", "body": "입문"},
        {"id": "s1", "title": "클라우드 네이티브와 컨테이너 개요", "layout": "bullets",
         "bullets": ["컨테이너는 격리된 환경에서 프로그램 실행", "Kubernetes는 오케스트레이션"]},
        {"id": "s2", "title": "컨트롤 플레인과 워커 노드 구분", "layout": "two-column",
         "bullets": ["컨트롤 플레인: kube-apiserver, etcd, kube-scheduler",
                     "워커 노드: kubelet, kube-proxy, 컨테이너 런타임"]},
        {"id": "s3", "title": "핵심 오브젝트 관계도", "layout": "process",
         "items": [["Deployment", "ReplicaSet을 통해 Pod 복제 관리"],
                   ["Service", "레이블로 Pod를 골라 접근 경로 제공"],
                   ["Pod", "컨테이너를 담은 최소 실행 단위"]]},
        {"id": "s4", "title": "파드 생성 요청 처리 흐름", "layout": "steps",
         "steps": [["요청", "kubectl apply"], ["저장", "kube-apiserver가 etcd에 저장"],
                   ["배정", "kube-scheduler가 노드 선택"], ["실행", "kubelet이 컨테이너 실행"]]},
        {"id": "s5", "title": "감사합니다", "layout": "closing", "body": "질문"},
    ]


def test_titles_and_the_request_s_own_topics_say_which_slides_are_structures_or_flows():
    assert diagrams.wanted_figure("핵심 오브젝트 관계도") == "relation"
    assert diagrams.wanted_figure("Object relationships") == "relation"
    assert diagrams.wanted_figure("클러스터 전체 구조") == "method"
    assert diagrams.wanted_figure("파드 생성 요청 처리 흐름") == "flow"
    # 「구분」 says nothing, but the request names 「컨트롤 플레인/워커 노드 구조」
    assert diagrams.wanted_figure("컨트롤 플레인과 워커 노드 구분") == ""
    assert diagrams.wanted_figure("컨트롤 플레인과 워커 노드 구분", _K8S) == "method"
    assert diagrams.wanted_figure("클라우드 네이티브와 컨테이너 개요", _K8S) == ""
    # a one-word topic (「Kubernetes 구조」) marks no slide by itself
    assert diagrams.wanted_figure("Kubernetes 소개", _K8S) == ""


def test_a_relation_or_structure_slide_is_not_planned_as_a_sequence_or_a_table():
    plan = [
        {"title": "표지", "layout": "title"},
        {"title": "핵심 오브젝트 관계도", "layout": "process"},
        {"title": "컨트롤 플레인과 워커 노드 구분", "layout": "table"},
        {"title": "도입 단계", "layout": "process"},
    ]
    out = deck.structure_as_drawable(plan, _K8S)
    assert out[1]["layout"] == "bullets" and out[2]["layout"] == "bullets"
    # a sequence titled as one keeps its shape
    assert out[3]["layout"] == "process"


@pytest.mark.asyncio
async def test_a_marked_slide_the_planner_passes_by_is_asked_again_and_a_named_diagram_is_drawn():
    prompts: list[str] = []

    async def complete(model, messages, api_key, max_tokens):
        prompts.append(messages[-1]["content"])
        return "[]", {"inputTokens": 1, "outputTokens": 1}

    parts = [("컨트롤 플레인과 워커 노드 구분", "컨트롤 플레인: kube-apiserver, etcd"),
             ("핵심 오브젝트 관계도", "Deployment — ReplicaSet으로 Pod 관리\nService — Pod 선택")]
    rows, _ = await diagrams.plan(
        parts=parts, eligible=[0, 1], request=_K8S, model="m", api_key="k",
        complete=complete, slide=True, marked={0: "method", 1: "relation"})
    assert "(구성 요소와 연결)" in prompts[0] and "(대상 사이의 관계)" in prompts[0]
    # one retry naming the slides passed by
    assert len(prompts) == 2 and "[2] 핵심 오브젝트 관계도" in prompts[1]
    # the title naming its diagram is drawn from its own words; 「구분」 is left to the planner
    assert [row.index for row in rows] == [1]
    assert rows[0].figure == "method" and rows[0].relation
    assert "Deployment" in rows[0].description and rows[0].caption == "핵심 오브젝트 관계도"


@pytest.mark.asyncio
async def test_a_relation_planned_as_a_flow_is_drawn_as_a_structure_with_named_lines():
    answer = json.dumps([{"part": 1, "figure": "flow", "caption": "오브젝트 관계",
                          "description": "Pod → Service → Deployment → StatefulSet"}],
                        ensure_ascii=False)

    async def complete(model, messages, api_key, max_tokens):
        return answer, {"inputTokens": 1, "outputTokens": 1}

    rows, _ = await diagrams.plan(
        parts=[("핵심 오브젝트 관계도", "Deployment, Service, Pod")], eligible=[0],
        request="", model="m", api_key="k", complete=complete, slide=True,
        marked={0: "relation"})
    assert rows[0].figure == "method" and rows[0].relation

    calls: list[list[dict]] = []

    async def drawn(model, messages, api_key, max_tokens=0):
        calls.append(messages)
        return ("```mermaid\nflowchart LR\n  d(Deployment) -->|관리| r(ReplicaSet)\n"
                "  r -->|관리| p(Pod)\n  s(Service) -.->|선택| p\n```\n그림: 관계",
                {"inputTokens": 1, "outputTokens": 1})

    original = diagram._complete
    diagram._complete = drawn
    try:
        source, _, _ = await diagram.draw(
            description=rows[0].description, figure="method", model="m", api_key="k",
            relation=True)
    finally:
        diagram._complete = original
    # no shape-picking call: a relation is never a fixed line of steps
    assert len(calls) == 1 and "relationship figure" in calls[0][-1]["content"]
    assert "|선택|" in source


@pytest.mark.asyncio
async def test_an_english_deck_plans_draws_and_captions_its_figures_in_english(monkeypatch):
    slides = [
        {"title": "Recommendation: Hybrid Routing Strategy", "layout": "bullets",
         "bullets": ["Route documents of 20 pages or fewer to the local 27B model",
                     "Send longer documents to the frontier cloud model"]},
        {"title": "The Critical Limitation: Long-Context Failures", "layout": "bullets",
         "bullets": ["64% vs 91% success above 20 pages on the drafting benchmark"]},
    ]
    assert deck.deck_language(slides) == "en"
    assert deck.deck_language(_slides()) == "ko"
    prompts: list[str] = []

    async def complete(model, messages, api_key, max_tokens):
        prompts.append(messages[-1]["content"])
        return json.dumps([{"part": 1, "figure": "flow", "caption": "Hybrid routing",
                            "description": "Input document, page check, local or cloud"}]), {}

    rows, _ = await diagrams.plan(
        parts=[(s["title"], "\n".join(s["bullets"])) for s in slides], eligible=[0, 1],
        request="Create a 5-slide English presentation", model="m", api_key="k",
        complete=complete, slide=True, language="en")
    assert "in English" in prompts[0] and "four English words" in prompts[0]
    assert "한국어로 구체적으로" not in prompts[0]
    assert rows[0].language == "en"

    seen: dict = {}

    async def draw_native(**kwargs):
        seen.update(kwargs)
        return "flowchart LR\n  a[Input] --> b(Local)", "Routing", {}, b"\x89PNG"

    monkeypatch.setattr(diagram, "draw_native", draw_native)
    await diagrams.make(rows[0], model="m", api_key="k", slide=True)
    assert seen["language"] == "en"

    asked: list[str] = []

    async def pictures(model, messages, api_key, max_tokens):
        asked.append(messages[-1]["content"])
        return "[]", {}

    await illustrations.plan(parts=[("a", "b")], eligible=[0], request="", model="m",
                             api_key="k", complete=pictures, slide=True, language="en")
    assert "short English caption" in asked[0]


@pytest.mark.asyncio
async def test_a_lecture_deck_offers_its_structure_relation_and_flow_slides_to_the_planner(
        monkeypatch):
    seen: dict = {}

    async def plan(**kwargs):
        seen.update(kwargs)
        return [], {"inputTokens": 0, "outputTokens": 0}

    monkeypatch.setattr(deck.diagrams, "plan", plan)
    async for _ in deck._draw_figures(_slides(), request=_K8S, model="m", api_key="k",
                                      usage={"inputTokens": 0, "outputTokens": 0},
                                      wrap=lambda p: [{"role": "user", "content": p}]):
        pass
    # the process and steps slides open because their titles are a relation and a flow
    assert seen["eligible"] == [1, 2, 3, 4]
    assert seen["marked"] == {2: "method", 3: "relation", 4: "flow"}
    assert seen["language"] == "ko"
    # a pattern's items reach the planner
    assert "Deployment" in seen["parts"][3][1]

    offered: dict = {}

    async def pictures(**kwargs):
        offered.update(kwargs)
        return [], {"inputTokens": 0, "outputTokens": 0}

    monkeypatch.setattr(deck.illustrations, "plan", pictures)
    structure = {"id": "x", "title": "클러스터 전체 구조", "layout": "bullets", "bullets": ["a"]}
    async for _ in deck._draw_pictures([*_slides(), structure], image_model={}, request=_K8S,
                                       model="m", api_key="k",
                                       usage={"inputTokens": 0, "outputTokens": 0},
                                       wrap=None):
        pass
    # an illustration never stands in for a structure
    assert offered["eligible"] == [1]


def test_an_edit_asking_for_a_diagram_is_told_from_one_removing_it():
    ask = "4장을 컨트롤 플레인과 워커 노드의 구조도(그림)로 바꿔줘. 왼쪽에 컨트롤 플레인"
    assert deck.asks_slide_figure(ask)
    assert deck.asks_slide_figure("이 장을 흐름도로 그려 줘")
    assert deck.asks_slide_figure("관계를 그림으로 보여줘")
    assert not deck.asks_slide_figure("구조도는 빼 줘")
    assert not deck.asks_slide_figure("2장을 더 짧게 써 줘")
    slide = {"title": "핵심 오브젝트 관계도", "layout": "process"}
    assert deck.asked_figure_kind(ask, slide) == "method"
    assert deck.asked_figure_kind("흐름도로 바꿔 줘", slide) == "flow"
    assert deck.asked_figure_kind("그림으로 그려 줘", slide) == "relation"
    # the deck edit draws it instead of rewriting the words
    from app.pipeline import revise_stages

    source = inspect.getsource(revise_stages.Rewrite.run)
    assert "asks_slide_figure(asked)" in source and "figure_slide(" in source


@pytest.mark.asyncio
async def test_a_table_slide_asked_to_be_a_structure_diagram_gets_the_figure(monkeypatch):
    table = {"id": "t", "title": "컨트롤 플레인과 워커 노드 구분", "layout": "table",
             "accent": "#1f6feb",
             "rows": [["구분", "컨트롤 플레인", "워커 노드"], ["역할", "결정 및 조정", "실제 실행"],
                      ["핵심 구성", "kube-apiserver, etcd", "kubelet, kube-proxy"]]}
    planned_seen: list = []

    async def make(planned, *, model, api_key, slide, accent=""):
        planned_seen.append(planned)
        return {"figure": planned.figure, "description": planned.description,
                "source": "flowchart LR\n  subgraph cp[컨트롤 플레인]\n    a(kube-apiserver)\n"
                          "  end\n  subgraph wn[워커 노드]\n    k(kubelet)\n  end\n  a --> k",
                "caption": planned.caption, "key": "k"}, {"inputTokens": 3}, b"\x89PNG"

    monkeypatch.setattr(deck.diagrams, "make", make)
    written, usage = await deck.figure_slide(
        request=_K8S, slides=[table], target_id="t",
        instruction="4장을 컨트롤 플레인과 워커 노드의 구조도(그림)로 바꿔줘", model="m",
        api_key="k")
    assert planned_seen[0].figure == "method" and "kube-apiserver" in planned_seen[0].description
    assert written["diagram"]["source"].startswith("flowchart LR")
    assert written["image"]["diagram"] is True
    # the table's cells are words under the figure or in the notes, never a table under it
    assert written["layout"] == "bullets" and "rows" not in written
    assert "결정 및 조정" in " ".join([*written.get("bullets", []), written.get("notes", "")])
    assert usage == {"inputTokens": 3}


def test_a_band_s_long_name_is_kept_whole_and_fitted_to_its_label():
    pairs = deck._clean_pairs([["kube-controller-manager", "복제, 엔드포인트 등 컨트롤러"],
                               ["Container Runtime", "컨테이너 실행"]], "bands")
    assert [p[0] for p in pairs] == ["kube-controller-manager", "Container Runtime"]
    names = [p[0] for p in pairs] + ["kube-apiserver"]
    for width in (816.0, 380.0):
        label, fitted = deck_export._band_labels(names, 16.0, width, 50.0)
        assert 96.0 <= label <= width / 3
        for name, (lines, size) in zip(names, fitted, strict=True):
            # every letter kept, each line inside the label
            assert "".join(lines).replace(" ", "") == name.replace(" ", "")
            assert all(figure_text.width(line, size, bold=True) <= label - 16 for line in lines)


def test_the_exported_bands_carry_the_whole_name():
    from io import BytesIO

    from pptx import Presentation

    slide = {"id": "b", "title": "컨트롤 플레인 핵심 요소", "layout": "bands",
             "bands": [["kube-apiserver", "요청 수신, 인증, 검증"],
                       ["kube-controller-manager", "복제, 엔드포인트 등 다양한 컨트롤러"]]}
    pptx = Presentation(BytesIO(deck_export.to_pptx("덱", [slide])))
    text = " ".join(shape.text_frame.text for page in pptx.slides for shape in page.shapes
                    if shape.has_text_frame)
    assert "kube-controller-manager" in text.replace("\n", "")
    assert deck_export.to_pdf("덱", [slide]).startswith(b"%PDF")
