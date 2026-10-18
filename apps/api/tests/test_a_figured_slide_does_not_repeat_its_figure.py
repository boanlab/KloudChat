"""A slide with a drawn figure keeps only the words the figure does not say; a figure that says
it all takes the whole body, laid out for that box."""

import io
import json

from pptx import Presentation

from app.services import deck, deck_export, diagram_render, figure_kinds, quality_gate

STEPS = {"kind": "steps", "lang": "ko", "steps": [
    {"name": "철거", "note": "긴 테이블 120석과 형광등 철거"},
    {"name": "배선·조명 교체", "note": "형광등을 따뜻한 색 조명으로 바꾼다"},
    {"name": "가구 배치", "note": "창가 카페 좌석·공용 테이블·1인 좌석 설치"},
    {"name": "시범 운영", "note": "직원 이용 후 의견 수렴"},
]}
SOURCE = ("flowchart LR\n  s0(\"철거\")\n" + figure_kinds.MARK
          + json.dumps(STEPS, ensure_ascii=False))


def _slide(bullets: list[str]) -> dict:
    return {"layout": "bullets", "title": "공사 단계", "bullets": bullets, "notes": "",
            "diagram": {"source": SOURCE, "caption": "공사 단계"}}


def test_bullets_that_list_the_steps_again_go_to_the_notes():
    slide = _slide(["철거: 기존 시설을 제거합니다", "설치: 배선과 조명을 교체합니다",
                    "구체: 가구를 배치합니다", "시범: 운영을 시험합니다"])
    deck._words_under_figure(slide)
    assert slide["bullets"] == [] and deck._figure_alone(slide)
    assert "철거: 기존 시설을 제거합니다" in slide["notes"]


def test_a_line_the_figure_does_not_say_stays_under_it():
    fact = "공사는 여름 휴가 3주 동안 하고 예산은 [금액]이다"
    slide = _slide(["철거: 기존 시설을 제거합니다", fact])
    deck._words_under_figure(slide)
    assert fact in slide["bullets"]


def test_cards_that_add_details_stay():
    tree = ("flowchart TB\n root(세 가지 공간)\n a[창가 카페 좌석]\n b[중앙 회의 테이블]\n"
            " c[안쪽 휴식 1인 좌석]\n root-->a\n root-->b\n root-->c")
    cards = [["카페", "창가에 햇빛이 드는 2~4인 좌석"],
             ["회의", "중앙에 큰 공용 테이블과 화이트보드"],
             ["휴식", "안쪽에 조용히 혼자 쉬는 1인 좌석"]]
    slide = {"layout": "cards", "title": "세 가지 공간 컨셉", "cards": [list(c) for c in cards],
             "diagram": {"source": tree}}
    deck._words_under_figure(slide)
    assert slide["cards"] == cards
    # …and a free-drawn box per card the cards already name is not drawn at all
    plain = {"layout": "cards", "title": "세 가지 공간 컨셉", "cards": cards}
    assert deck._figure_adds_nothing(plain, {"source": tree})
    other = "flowchart LR\n a[예약 시스템]-->b[출입]"
    assert not deck._figure_adds_nothing(plain, {"source": other})
    assert not deck._figure_adds_nothing(plain, {"source": SOURCE})  # a fixed shape stays


def test_steps_alone_on_a_slide_are_rows_and_a_strip_keeps_its_arrows():
    assert figure_kinds._layout(STEPS, deck.FIGURE_BODY)[1] is figure_kinds._steps_rows
    assert figure_kinds._layout(STEPS, (816, 120))[1] is figure_kinds._steps
    assert figure_kinds._layout(STEPS)[1] is figure_kinds._steps  # a report's figure
    bare = {**STEPS, "steps": [{"name": s["name"]} for s in STEPS["steps"]]}
    assert figure_kinds._layout(bare, deck.FIGURE_BODY)[1] is figure_kinds._steps
    # the picture the panel shows is drawn for the same box
    blob = diagram_render.figure_deck(SOURCE, "#1e3a8a", box=deck.FIGURE_BODY)
    texts = [s.text_frame.text for s in Presentation(io.BytesIO(blob)).slides[0].shapes
             if s.has_text_frame]
    assert "1단계" in texts and "직원 이용 후 의견 수렴" in texts


def test_a_figure_alone_fills_the_body_without_a_caption_repeating_the_title():
    slide = _slide([])
    slide["image"] = {"src": "data:image/png;base64,iVBORw0KGgo=", "caption": "공사 단계",
                      "diagram": True, "renderer": "pptx", "size": "full"}
    out = Presentation(io.BytesIO(deck_export.to_pptx("덱", [slide])))
    texts = [s.text_frame.text for s in out.slides[0].shapes if s.has_text_frame]
    assert texts.count("공사 단계") == 1  # the title, no caption under the figure
    rows = [s for s in out.slides[0].shapes if s.has_text_frame and "단계" in s.text_frame.text
            and s.text_frame.text != "공사 단계"]
    assert len(rows) == 4
    assert deck_export._unless_title("공사 단계 흐름도", {"title": "공사 단계"}) == ""
    assert deck_export._unless_title("공사 단계", {"title": "공사 4단계"}) == ""
    assert deck_export._unless_title("현재와 리뉴얼 뒤", {"title": "공사 단계"})
    # and a figure that carries its own words is not a finding
    codes = [f["code"] for f in quality_gate.deck_findings([slide])]
    assert "figure_without_words" not in codes
