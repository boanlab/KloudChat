"""Script skills: packages on disk, schemas computed by code, the sandbox tools, their outputs."""

from __future__ import annotations

import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import pytest

from app.schemas.workspace import SkillOut
from app.services import agent, file_schema, sandbox, skill_packages, starter
from app.services.tools import sandbox_tools
from app.services.tools.base import Tool, ToolContext, ToolResult

SKILLS = Path(skill_packages.__file__).resolve().parent.parent / "skills"


# ── fixtures built with the stdlib ──────────────────────────────────────


def _column(index: int) -> str:
    name = ""
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        name = chr(65 + rest) + name
    return name


_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_PKG = "http://schemas.openxmlformats.org/package/2006"
_DOC = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def make_xlsx(
    sheets: dict[str, list[list]], *, date_cells: frozenset[tuple[str, int, int]] = frozenset()
) -> bytes:
    """A minimal workbook: inline strings, numbers, and date-styled serials."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            f'<?xml version="1.0"?><Types xmlns="{_PKG}/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            + "".join(
                f'<Override PartName="/xl/worksheets/sheet{n}.xml" ContentType="application/'
                'vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                for n in range(1, len(sheets) + 1)
            )
            + "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            f'<?xml version="1.0"?><Relationships xmlns="{_PKG}/relationships">'
            f'<Relationship Id="rId1" Type="{_DOC}/officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>",
        )
        entries, rels = [], []
        for number, (name, rows) in enumerate(sheets.items(), start=1):
            entries.append(f'<sheet name="{escape(name)}" sheetId="{number}" r:id="rId{number}"/>')
            rels.append(
                f'<Relationship Id="rId{number}" Type="{_DOC}/worksheet" '
                f'Target="worksheets/sheet{number}.xml"/>'
            )
            body = []
            for r, row in enumerate(rows, start=1):
                cells = []
                for c, value in enumerate(row):
                    ref = f"{_column(c)}{r}"
                    if value is None:
                        continue
                    if isinstance(value, str) and value.startswith("="):
                        cells.append(f'<c r="{ref}"><f>{escape(value[1:])}</f></c>')
                    elif isinstance(value, str):
                        inline = f"<is><t>{escape(value)}</t></is>"
                        cells.append(f'<c r="{ref}" t="inlineStr">{inline}</c>')
                    else:
                        style = ' s="1"' if (name, r, c) in date_cells else ""
                        cells.append(f'<c r="{ref}"{style}><v>{value}</v></c>')
                body.append(f'<row r="{r}">{"".join(cells)}</row>')
            archive.writestr(
                f"xl/worksheets/sheet{number}.xml",
                f'<?xml version="1.0"?><worksheet xmlns="{_MAIN}">'
                f"<sheetData>{''.join(body)}</sheetData></worksheet>",
            )
        archive.writestr(
            "xl/workbook.xml",
            f'<?xml version="1.0"?><workbook xmlns="{_MAIN}" xmlns:r="{_DOC}">'
            f"<sheets>{''.join(entries)}</sheets></workbook>",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            f'<?xml version="1.0"?><Relationships xmlns="{_PKG}/relationships">'
            f"{''.join(rels)}</Relationships>",
        )
        archive.writestr(
            "xl/styles.xml",
            f'<?xml version="1.0"?><styleSheet xmlns="{_MAIN}">'
            '<cellXfs count="2"><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs></styleSheet>',
        )
    return buffer.getvalue()


MESSY_ROWS = [
    ["2026년 3분기 지점별 실적 보고", None, None, None],
    ["(단위: 원, 작성: 기획팀)", None, None, None],
    [None, None, None, None],
    ["지점", "방문객", "매출", "비용"],
    ["강남점", "28,887", "349,275,012", "270,171,128"],
    ["홍대점", "26,455", "436,905,959", "228,859,795"],
    ["잠실점", "81,893", "881,522,991", "96,069,098"],
    ["잠실점", "81,893", "881,522,991", "96,069,098"],
    ["합계", "137,235", "1,667,703,962", "595,100,021"],
    [None, None, None, None],
    ["※ 비용은 VAT 제외", None, None, None],
]


def make_pdf(pages: list[list[list[str]]]) -> bytes:
    """One table per page drawn as positioned text, header repeated, no ruling lines."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for rows in pages:
        pdf.setFont("HYSMyeongJo-Medium", 10)
        y = 780
        for row in rows:
            for x, cell in zip((60, 200, 330, 450), row, strict=False):
                pdf.drawString(x, y, cell)
            y -= 18
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


# ── packages ────────────────────────────────────────────────────────────


def test_every_package_loads_with_its_scripts():
    packages = skill_packages.packages()
    assert {"xlsx-clean", "pdf-table", "calc-check"} <= set(packages)
    for package in packages.values():
        assert package.name and package.description and package.when_to_use and package.body
        assert "run_skill_script" in package.required_tools
        assert package.files[0] == "SKILL.md"
        assert all((package.root / rel).is_file() for rel in package.files)
    assert packages["xlsx-clean"].scripts.keys() == {"clean"}
    assert packages["pdf-table"].script_path("extract.py") == "skills/pdf-table/scripts/extract.py"
    assert packages["calc-check"].script_path("nope") is None


def test_the_catalogue_carries_the_package_body_and_old_keys_still_resolve():
    by_key = {spec["key"]: spec for spec in starter._SKILLS}
    package = skill_packages.packages()["xlsx-clean"]
    assert by_key["xlsx-clean"]["body"] == package.body
    assert by_key["xlsx-clean"]["required_tools"] == ["run_skill_script"]
    assert "calculation-unit-check" not in by_key
    assert skill_packages.package_for("calculation-unit-check").key == "calc-check"
    assert skill_packages.find("「엑셀 정리」").key == "xlsx-clean"
    assert skill_packages.script_names("calculation-unit-check") == ["check"]
    assert skill_packages.script_names("citation") == []


def test_front_matter_is_required_and_strict(tmp_path):
    with pytest.raises(skill_packages.PackageError):
        skill_packages.split_front_matter("no front matter")
    package = tmp_path / "demo"
    (package / "scripts").mkdir(parents=True)
    head = "---\nname: 데모\ndescription: d\nwhen_to_use: w\n"
    (package / "SKILL.md").write_text(head + "extra: 1\n---\n본문")
    with pytest.raises(skill_packages.PackageError, match="unknown fields extra"):
        skill_packages.load(package)
    (package / "SKILL.md").write_text(head + "---\n본문\n")
    (package / "scripts" / "go.py").write_text('"""한 줄 설명.\n\n자세히."""\n')
    loaded = skill_packages.load(package)
    assert loaded.body == "본문" and loaded.scripts == {"go": "한 줄 설명."}
    assert loaded.kinds == ("chat",) and loaded.version == "1.0.0"


def test_a_skill_row_lists_its_package_scripts():
    from app.models.workspace import Skill

    row = Skill(owner_id="u", name="엑셀 정리", catalog_key="xlsx-clean", body="b")
    assert SkillOut.of(row).scripts == ["clean"]
    assert SkillOut.of(Skill(owner_id="u", name="mine", body="b")).scripts == []


# ── schemas ─────────────────────────────────────────────────────────────


def test_xlsx_schema_finds_the_header_and_shows_rows_as_json():
    data = make_xlsx({"실적": MESSY_ROWS, "날짜": [["일자", "값"], [46083, 1]]},
                     date_cells={("날짜", 2, 0)})
    text = file_schema.describe("지점실적.xlsx", data)
    assert "머리글 추정: 4행" in text
    assert '["지점", "방문객", "매출", "비용"]' in text
    assert '["강남점", "28,887", "349,275,012", "270,171,128"]' in text
    assert "※ 비용은 VAT 제외" in text  # last rows, so notes and totals are visible
    assert " | " not in text
    assert '["2026-03-02", 1]' in text  # a date-styled serial reads as a date


def test_csv_schema_reports_encoding_and_delimiter():
    data = "거래일;금액\n2026.4.21;4000\n46083;4300\n".encode("cp949")
    text = file_schema.describe("거래.csv", data)
    assert "cp949" in text and "sep=';'" in text
    assert '["2026.4.21", "4000"]' in text


def test_pdf_docx_pptx_schemas():
    pdf = make_pdf([[["거래일", "구분", "금액"], ["2026-09-01", "입금", "286,000"]]])
    text = file_schema.describe("내역.pdf", pdf)
    assert "PDF 1쪽" in text and "2026-09-01" in text

    import docx
    import pptx

    document = docx.Document()
    document.add_paragraph("첫 문단")
    document.add_table(rows=1, cols=2)
    buffer = io.BytesIO()
    document.save(buffer)
    assert "표 1개" in file_schema.describe("a.docx", buffer.getvalue())

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "제목 슬라이드"
    buffer = io.BytesIO()
    deck.save(buffer)
    text = file_schema.describe("b.pptx", buffer.getvalue())
    assert "슬라이드 1장" in text and "제목 슬라이드" in text


def test_output_checks_open_the_file():
    data = make_xlsx({"정리": [["a", "b"], [1, 2], [3, "=A3*2"]]})
    note = file_schema.check_output("x.xlsx", data)
    assert "시트 1개" in note and "정리 2행" in note and "계산값 없는 수식 1칸" in note
    assert file_schema.check_output("bad.xlsx", b"nope").startswith("열리지 않음")
    assert "1쪽" in file_schema.check_output("p.pdf", make_pdf([[["a", "b"]]]))


# ── scripts, run standalone ─────────────────────────────────────────────


def _run_script(path: Path, args: dict, cwd: Path) -> dict:
    done = subprocess.run(
        [sys.executable, str(path), json.dumps(args, ensure_ascii=False)],
        cwd=cwd, capture_output=True, text=True, timeout=60,
    )
    assert done.returncode == 0, done.stderr + done.stdout
    return json.loads(done.stdout.strip().splitlines()[-1])


def test_calc_check_script_flags_the_wrong_figure(tmp_path):
    summary = _run_script(
        SKILLS / "calc-check" / "scripts" / "check.py",
        {
            "items": [
                {"label": "부가세", "expr": "1,250,000 * 0.1", "expected": 125000},
                {"label": "증가율", "kind": "change", "old": 120, "new": 150, "expected": 25},
                {"label": "반올림", "expr": "2/3*100", "expected": 66.67},
                {"label": "비중", "kind": "share", "part": 30, "whole": 120, "expected": 30},
                {"label": "단위", "kind": "convert", "value": 1.2, "from": "억원", "to": "천원",
                 "expected": 120000},
                {"label": "0", "expr": "1/0"},
                {"label": "반올림 함수", "expr": "round(2.5)", "expected": 3},
            ]
        },
        tmp_path,
    )
    status = {item["label"]: item["status"] for item in summary["items"]}
    assert status == {
        "부가세": "match", "증가율": "match", "반올림": "rounded", "비중": "mismatch",
        "단위": "match", "0": "error", "반올림 함수": "match",
    }
    assert summary["mismatches"] == 1 and summary["errors"] == 1


def test_xlsx_clean_script_normalises_a_messy_sheet(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    (tmp_path / "지점실적.xlsx").write_bytes(make_xlsx({"Sheet1": MESSY_ROWS}))
    summary = _run_script(
        SKILLS / "xlsx-clean" / "scripts" / "clean.py",
        {"file": "지점실적.xlsx", "add_totals": True, "recalc": False},
        tmp_path,
    )
    sheet = summary["sheets"][0]
    assert sheet["header_row"] == 4
    assert sheet["rows_out"] == 3
    assert sheet["removed"] == {"title": 2, "blank": 1, "total": 1, "note": 1, "duplicate": 1}
    assert sheet["numeric_columns"] == ["방문객", "매출", "비용"]
    book = openpyxl.load_workbook(tmp_path / summary["outputs"][0])
    assert book.sheetnames == ["Sheet1", "변경내역"]
    rows = list(book["Sheet1"].iter_rows(values_only=True))
    assert rows[1] == ("강남점", 28887, 349275012, 270171128)
    assert rows[-1] == ("합계", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)")
    log = [r[1] for r in book["변경내역"].iter_rows(min_row=2, values_only=True)]
    assert "삭제: 주석 행" in log and "삭제: 중복 행" in log


def test_xlsx_clean_script_reads_dates_and_units_from_csv(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    text = (
        "거래일,금액,비율\n2026-08-22,\"1,200천원\",12.5%\n2026.4.21,3천원,50%\n"
        "20260825,￦7천원,1%\n46083,2천원,2%\n2026년 3월 9일,5천원,3%\n"
    )
    (tmp_path / "거래.csv").write_bytes(text.encode("cp949"))
    clean = SKILLS / "xlsx-clean" / "scripts" / "clean.py"
    summary = _run_script(clean, {"file": "거래.csv"}, tmp_path)
    sheet = summary["sheets"][0]
    assert summary["csv"]["encoding"] == "cp949"
    assert sheet["date_columns"] == ["거래일"]
    assert sheet["percent_columns"] == ["비율"]
    assert "금액(천원)" in sheet["numeric_columns"]
    rows = list(openpyxl.load_workbook(tmp_path / summary["outputs"][0]).worksheets[0].iter_rows(
        min_row=2, values_only=True))
    assert [r[0].strftime("%Y-%m-%d") for r in rows] == [
        "2026-08-22", "2026-04-21", "2026-08-25", "2026-03-02", "2026-03-09",
    ]
    assert [r[1] for r in rows] == [1200, 3, 7, 2, 5]
    assert rows[0][2] == pytest.approx(0.125)


def test_pdf_table_script_joins_a_table_across_pages(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    header = ["거래일", "구분", "내용", "금액"]
    first = [header] + [[f"2026-09-0{d}", "입금", "급여", f"{d * 1000:,}"] for d in range(1, 6)]
    second = [header] + [
        [f"2026-09-1{d}", "출금", "환불", f"{d * 1000 + 7:,}"] for d in range(0, 4)
    ]
    (tmp_path / "내역.pdf").write_bytes(make_pdf([first, second]))
    extract = SKILLS / "pdf-table" / "scripts" / "extract.py"
    summary = _run_script(extract, {"file": "내역.pdf"}, tmp_path)
    assert len(summary["tables"]) == 1
    table = summary["tables"][0]
    assert table["pages"] == [1, 2] and table["rows"] == 9 and table["columns"] == header
    assert summary["repeated_headers_removed"] == 1
    sheet = openpyxl.load_workbook(tmp_path / summary["outputs"][0])["표1"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert rows[0] == ("2026-09-01", "입금", "급여", 1000)
    assert rows[-1] == ("2026-09-13", "출금", "환불", 3007)


def test_scripts_report_a_missing_input_as_json(tmp_path):
    pytest.importorskip("openpyxl")
    done = subprocess.run(
        [sys.executable, str(SKILLS / "xlsx-clean" / "scripts" / "clean.py"), '{"file": "x.xlsx"}'],
        cwd=tmp_path, capture_output=True, text=True, timeout=60,
    )
    assert done.returncode == 1
    assert json.loads(done.stdout.strip().splitlines()[-1])["ok"] is False


# ── tools with a fake sandbox ───────────────────────────────────────────


class FakeSandbox:
    """Records uploads and runs; `runs` scripts each exec's result."""

    def __init__(self, runs: list[dict], blobs: dict[str, bytes] | None = None) -> None:
        self.runs = runs
        self.blobs = blobs or {}
        self.uploads: list[dict] = []
        self.execs: list[dict] = []

    async def upload(self, items, *, kind=None, read_only=False):
        session = f"up{len(self.uploads)}"
        self.uploads.append({"names": [n for n, _ in items], "kind": kind, "read_only": read_only})
        return [{"id": f"f-{name}", "name": name, "session_id": session} for name, _ in items]

    async def exec(self, code, *, lang="py", files=None, session_id=None, timeout_ms=30_000):
        self.execs.append({"code": code, "lang": lang, "files": files, "session_id": session_id})
        return sandbox.interpret({"session_id": "exec-1", **self.runs.pop(0)})

    async def download(self, session_id, file_id):
        return self.blobs[file_id]


@pytest.fixture
def chat(monkeypatch):
    """A chat holding one upload; stored outputs are captured."""
    upload = make_xlsx({"Sheet1": MESSY_ROWS})
    stored: list[dict] = []

    async def session_files(user_id, session_id):
        assert (user_id, session_id) == ("user-1", "chat-1")
        return [sandbox.InputFile("file-1", "지점실적.xlsx", len(upload), "user-1/file-1")]

    async def store_output(user_id, session_id, name, data):
        row = {"id": f"stored-{len(stored)}", "name": name, "size": len(data)}
        stored.append({**row, "user_id": user_id, "session_id": session_id})
        return row

    monkeypatch.setattr(sandbox, "session_files", session_files)
    monkeypatch.setattr(sandbox, "read_input", lambda _file: upload)
    monkeypatch.setattr(sandbox, "store_output", store_output)
    monkeypatch.setattr(sandbox, "_SKILL_REFS", {})
    return stored


def _use(monkeypatch, fake: FakeSandbox) -> None:
    async def client():
        return fake

    monkeypatch.setattr(sandbox, "client", client)


async def test_run_skill_script_uploads_runs_and_stores_outputs(monkeypatch, chat):
    cleaned = make_xlsx({"Sheet1": [["지점"], ["강남점"]], "변경내역": [["시트"]]})
    summary = {"ok": True, "outputs": ["outputs/지점실적_정리.xlsx"], "sheets": []}
    fake = FakeSandbox(
        [
            {
                "stdout": (
                    f"[Sheet1] 정리\n{json.dumps(summary, ensure_ascii=False)}\n"
                    "\n__KLOUDCHAT_EXIT__ 0\n"
                ),
                "files": [
                    {"id": "f-지점실적.xlsx", "name": "지점실적.xlsx", "session_id": "up0",
                     "inherited": True},
                    {"id": "out-1", "name": "outputs/지점실적_정리.xlsx", "session_id": "exec-1"},
                    {"id": "junk", "name": "scratch.txt", "session_id": "exec-1"},
                ],
            },
            {"stdout": "{\"ok\": true}\n__KLOUDCHAT_EXIT__ 0\n", "files": []},
        ],
        {"out-1": cleaned},
    )
    _use(monkeypatch, fake)
    ctx = ToolContext(user_id="user-1", session_id="chat-1")
    result = await sandbox_tools.run_skill_script(
        {"skill": "xlsx-clean", "script": "clean", "args": {"file": "지점실적.xlsx"}}, ctx
    )

    assert not result.failed, result.content
    assert fake.uploads[0] == {"names": ["지점실적.xlsx"], "kind": None, "read_only": False}
    assert fake.uploads[1]["kind"] == "skill" and fake.uploads[1]["read_only"] is True
    assert "skills/xlsx-clean/scripts/clean.py" in fake.uploads[1]["names"]
    run = fake.execs[0]
    assert run["lang"] == "bash" and run["session_id"] is None
    assert run["code"].startswith("python3 skills/xlsx-clean/scripts/clean.py ")
    assert "지점실적.xlsx" in run["code"]
    assert result.outputs == [
        {"id": "stored-0", "name": "지점실적_정리.xlsx", "size": len(cleaned)}
    ]
    assert chat[0]["session_id"] == "chat-1" and chat[0]["user_id"] == "user-1"
    assert "파일 id stored-0" in result.content and "엑셀 열림 · 시트 2개" in result.content
    assert "scratch" not in result.content
    assert "__KLOUDCHAT_EXIT__" not in result.content

    # The second call continues the turn's sandbox session and uploads nothing new.
    await sandbox_tools.run_skill_script(
        {"skill": "calc-check", "script": "check", "args": {}}, ctx
    )
    assert fake.execs[1]["session_id"] == "exec-1"
    assert [u["kind"] for u in fake.uploads] == [None, "skill", "skill"]


async def test_run_skill_script_refuses_unknown_scripts_and_files(monkeypatch, chat):
    _use(monkeypatch, FakeSandbox([]))
    ctx = ToolContext(user_id="user-1", session_id="chat-1")
    unknown = await sandbox_tools.run_skill_script({"skill": "citation", "script": "x"}, ctx)
    assert unknown.failed and "xlsx-clean" in unknown.content
    missing = await sandbox_tools.run_skill_script(
        {"skill": "xlsx-clean", "script": "clean", "args": {"file": "없는.xlsx"}}, ctx
    )
    assert missing.failed and "지점실적.xlsx" in missing.content


async def test_a_failing_script_is_a_failed_step(monkeypatch, chat):
    fake = FakeSandbox([{"stdout": '{"ok": false, "error": "시트 없음"}\n__KLOUDCHAT_EXIT__ 1\n'}])
    _use(monkeypatch, fake)
    result = await sandbox_tools.run_skill_script(
        {"skill": "xlsx-clean", "script": "clean", "args": {"sheet": "x"}},
        ToolContext(user_id="user-1", session_id="chat-1"),
    )
    assert result.failed and "시트 없음" in result.content


async def test_run_code_repairs_once_then_falls_back(monkeypatch, chat):
    monkeypatch.setattr(sandbox_tools.settings, "code_model", "local/qwen3.8-27b")
    monkeypatch.setattr(sandbox_tools.settings, "code_fallback_model", "local/qwen3.5-122b")
    written: list[tuple[str, list[dict]]] = []

    async def write_code(model, messages, api_key):
        written.append((model, messages))
        return f"print('attempt {len(written)}')"

    monkeypatch.setattr(sandbox_tools, "write_code", write_code)
    result_file = make_xlsx({"요약": [["월", "합계"], ["1월", 10]]})
    fake = FakeSandbox(
        [
            {"stderr": "Traceback (most recent call last):\nKeyError: '매출'"},
            {"stderr": "Traceback (most recent call last):\nValueError: bad"},
            {
                "stdout": 'done\n{"outputs": ["outputs/요약.xlsx"], "rows": 1}',
                "files": [{"id": "out-9", "name": "outputs/요약.xlsx", "session_id": "exec-1"}],
            },
        ],
        {"out-9": result_file},
    )
    _use(monkeypatch, fake)
    ctx = ToolContext(user_id="user-1", session_id="chat-1", api_key="k")
    result = await sandbox_tools.run_code({"task": "월별 매출 합계를 요약.xlsx 로"}, ctx)

    assert [model for model, _ in written] == [
        "local/qwen3.8-27b", "local/qwen3.8-27b", "local/qwen3.5-122b",
    ]
    first = written[0][1]
    assert "머리글 추정: 4행" in first[1]["content"]  # the schema reached the model
    assert "outputs/" in first[0]["content"]
    assert "KeyError: '매출'" in written[1][1][-1]["content"]
    assert "ValueError: bad" in written[2][1][-1]["content"]
    assert not result.failed
    assert result.outputs == [{"id": "stored-0", "name": "요약.xlsx", "size": len(result_file)}]
    assert '"rows": 1' in result.content


async def test_run_code_reports_a_memory_kill_as_a_limit(monkeypatch, chat):
    monkeypatch.setattr(sandbox_tools.settings, "code_fallback_model", "")
    prompts: list[list[dict]] = []

    async def write_code(model, messages, api_key):
        prompts.append(messages)
        return "print(1)"

    monkeypatch.setattr(sandbox_tools, "write_code", write_code)
    oom = {"stderr": "Invalid response from REPL: delimiter not found"}
    _use(monkeypatch, FakeSandbox([dict(oom), dict(oom)]))
    result = await sandbox_tools.run_code(
        {"task": "큰 파일 집계"}, ToolContext(user_id="user-1", session_id="chat-1")
    )
    assert result.failed
    assert sandbox.MEMORY_LIMIT in result.content
    assert "delimiter" not in result.content
    assert "메모리 한도" in prompts[1][-1]["content"] and "chunksize" in prompts[1][-1]["content"]


def test_a_strict_local_turn_writes_code_on_strict_local_models(monkeypatch):
    monkeypatch.setattr(sandbox_tools.settings, "code_model", "local/qwen3.8-27b")
    monkeypatch.setattr(sandbox_tools.settings, "code_fallback_model", "openai/gpt-x")
    assert sandbox_tools.code_models(False) == ["local/qwen3.8-27b", "openai/gpt-x"]
    assert sandbox_tools.code_models(True) == ["strict-local/qwen3.8-27b"]


async def test_execute_code_mounts_the_files_it_names(monkeypatch, chat):
    fake = FakeSandbox([{"stdout": "4"}, {"stdout": "2"}])
    _use(monkeypatch, fake)
    ctx = ToolContext(user_id="user-1", session_id="chat-1")
    named = await sandbox_tools.execute_code({"code": "pd.read_excel('지점실적.xlsx')"}, ctx)
    assert named.content.startswith("stdout:")
    assert fake.execs[0]["files"][0]["name"] == "지점실적.xlsx"
    fresh = ToolContext(user_id="user-1", session_id="chat-1")
    await sandbox_tools.execute_code({"code": "print(1+1)"}, fresh)
    assert fake.execs[1]["files"] is None


class _Stream:
    status_code = 200

    def __init__(self, lines: list[str]) -> None:
        self._lines = lines

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args) -> None:
        return None

    async def aiter_lines(self):
        for line in self._lines:
            yield line


class _Model:
    def __init__(self, scripted: list[list[str]]) -> None:
        self._scripted = scripted

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args) -> None:
        return None

    def stream(self, _method: str, _path: str, *, json: dict) -> _Stream:
        return _Stream(self._scripted.pop(0))


async def test_the_step_carries_the_output_files(monkeypatch):
    async def run(_args, _ctx):
        return ToolResult(
            content="만든 파일 1개", detail="파일 1개",
            outputs=[{"id": "f1", "name": "결과.xlsx", "size": 10}],
        )

    tool = Tool(name="run_code", description="d", parameters={"type": "object"}, run=run,
                label="파일 작업 중", title="파일 작업", wants_context=True)
    call = {"choices": [{"delta": {"tool_calls": [
        {"index": 0, "id": "c0", "function": {"name": "run_code", "arguments": '{"task": "x"}'}}
    ]}}]}
    scripted = [
        ["data: " + json.dumps(call), "data: [DONE]"],
        [
            "data: " + json.dumps({"choices": [{"delta": {"content": "만들었습니다."}}]}),
            "data: [DONE]",
        ],
    ]

    async def client(*_args, **_kwargs):
        return _Model(scripted)

    monkeypatch.setattr(agent, "_client", client)
    ctx = ToolContext(user_id="u", session_id="s", api_key="k")
    events = [e async for e in agent.run_turn("m", [{"role": "user", "content": "q"}], [tool], ctx)]
    done = [e for e in events if e["type"] == "step" and e.get("status") == "done"]
    assert done[0]["outputs"] == [{"id": "f1", "name": "결과.xlsx", "size": 10}]
    running = [e for e in events if e["type"] == "step" and e.get("status") == "running"]
    assert "outputs" not in running[0]


def test_sandbox_failures_are_named():
    killed = sandbox.interpret({"stderr": "Invalid response from REPL: delimiter not found"})
    assert killed.memory_exceeded and killed.crashed and killed.stderr == sandbox.MEMORY_LIMIT
    slow = sandbox.interpret({"stderr": "TimeoutError: Execution exceeded 30 seconds"})
    assert slow.timed_out and sandbox.TIME_LIMIT in slow.stderr
    warned = sandbox.interpret({"stderr": "UserWarning: something", "stdout": "ok"})
    assert not warned.crashed
    assert sandbox.last_json('x\n{"a": 1}\n') == {"a": 1}
    assert sandbox.last_json('{"a": 1}\nmore') is None


def test_inputs_pick_the_newest_file_of_a_name():
    files = [
        sandbox.InputFile("new", "a.xlsx", 10, "k1"),
        sandbox.InputFile("old", "a.xlsx", 10, "k0"),
        sandbox.InputFile("b", "b.csv", 10, "k2"),
    ]
    chosen, missing = sandbox.pick_inputs(files, None)
    assert [f.id for f in chosen] == ["new", "b"] and missing == []
    chosen, missing = sandbox.pick_inputs(files, ["outputs/a.xlsx", "c.pdf"])
    assert [f.id for f in chosen] == ["new"] and missing == ["c.pdf"]


def test_the_tools_are_described_for_the_model():
    tool = sandbox_tools.RUN_SKILL_SCRIPT
    keys = set(tool.parameters["properties"]["skill"]["enum"])
    assert keys >= {"xlsx-clean", "pdf-table", "calc-check"}
    assert "xlsx-clean 「엑셀 정리」: clean(" in tool.description
    assert sandbox_tools.RUN_CODE.read_only is False and tool.wants_context



async def test_run_code_keeps_the_code_that_ran_and_its_output_on_the_step(monkeypatch, chat):
    monkeypatch.setattr(sandbox_tools.settings, "code_model", "local/qwen3.8-27b")
    monkeypatch.setattr(sandbox_tools.settings, "code_fallback_model", "")
    monkeypatch.setattr(sandbox_tools.settings, "code_interpreter_api_key", "sbx-secret-123456")
    codes = [
        "import pandas as pd\ndf = pd.read_excel('지점실적.xlsx')\nprint(df['매출'])",
        "import pandas as pd\ndf = pd.read_excel('/mnt/data/지점실적.xlsx')\n"
        + "x = 1\n" * 4_000
        + "print(df.shape)",
    ]

    async def write_code(model, messages, api_key):
        return codes.pop(0)

    monkeypatch.setattr(sandbox_tools, "write_code", write_code)
    fake = FakeSandbox(
        [
            {"stderr": "Traceback:\nKeyError: '매출'"},
            {
                "stdout": "rows 12\nkey sbx-secret-123456\n" + "줄\n" * 5_000,
                "stderr": (
                    'File "/usr/local/lib/python3.12/site-packages/pandas/io/excel.py", '
                    "line 3\nwarn"
                ),
            },
        ]
    )
    _use(monkeypatch, fake)
    result = await sandbox_tools.run_code(
        {"task": "지점 매출"},
        ToolContext(user_id="user-1", session_id="chat-1", api_key="sk-user-key-999"),
    )

    trace = result.trace
    assert trace is not None and not result.failed
    # The repaired code, not the first draft, and which attempt wrote it.
    assert trace["codeAttempt"] == 2
    assert trace["code"].startswith("import pandas as pd\ndf = pd.read_excel('/mnt/data/지점")
    assert len(trace["code"]) < 6_100 and "자 생략)" in trace["code"]
    assert trace["stdout"].startswith("rows 12")
    assert len(trace["stdout"]) < 3_100 and "자 생략)" in trace["stdout"]
    assert "sbx-secret-123456" not in json.dumps(trace) and "[비밀값]" in trace["stdout"]
    # No paths outside /mnt/data; the file name stays.
    assert "/usr/local" not in trace["stderr"] and "…/excel.py" in trace["stderr"]
    # The detail line stays short; the trace is separate.
    assert result.detail == "시도 2회"


async def test_run_skill_script_keeps_its_command_and_output(monkeypatch, chat):
    fake = FakeSandbox([{"stdout": '[Sheet1] 정리\n{"ok": true}\n__KLOUDCHAT_EXIT__ 0\n'}])
    _use(monkeypatch, fake)
    result = await sandbox_tools.run_skill_script(
        {"skill": "xlsx-clean", "script": "clean", "args": {"file": "지점실적.xlsx"}},
        ToolContext(user_id="user-1", session_id="chat-1"),
    )
    assert result.trace is not None
    assert result.trace["code"].startswith("python3 skills/xlsx-clean/scripts/clean.py ")
    assert result.trace["stdout"].startswith("[Sheet1] 정리")
    assert "__KLOUDCHAT_EXIT__" not in result.trace["stdout"]
    assert "codeAttempt" not in result.trace and "stderr" not in result.trace
