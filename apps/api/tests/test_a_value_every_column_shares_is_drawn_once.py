"""A table row whose value columns all hold one value draws it once across them."""

from __future__ import annotations

import io

from pptx import Presentation
from pypdf import PdfReader

from app.services import deck_export, deck_type, revise

_ROWS = [
    ["Metric", "27B local", "Frontier cloud"],
    ["Tasks", "40 prompts, 3 raters", "40 prompts, 3 raters"],
    ["Cost /1k docs", "$4", "$38"],
    ["Offline", "Yes", "Yes"],
]
_SLIDE = {"layout": "table", "title": "Benchmark Setup", "rows": _ROWS}


def _table(data: bytes):
    slide = Presentation(io.BytesIO(data)).slides[0]
    return next(shape.table for shape in slide.shapes if shape.has_table)


def test_a_shared_value_is_one_merged_cell_in_the_pptx() -> None:
    table = _table(deck_export.to_pptx("Deck", [_SLIDE]))
    origin = table.cell(1, 1)
    assert origin.is_merge_origin and origin.span_width == 2
    assert table.cell(1, 2).is_spanned
    assert origin.text == "40 prompts, 3 raters"
    # each column's own values and short shared marks keep their cells
    assert not table.cell(2, 1).is_merge_origin and table.cell(2, 2).text == "$38"
    assert not table.cell(3, 1).is_merge_origin and table.cell(3, 2).text == "Yes"
    # the header never merges
    assert not table.cell(0, 1).is_merge_origin


def test_a_shared_value_is_written_once_in_the_pdf() -> None:
    text = PdfReader(io.BytesIO(deck_export.to_pdf("Deck", [_SLIDE]))).pages[0].extract_text()
    assert text.count("3 raters") == 1
    assert text.count("Yes") == 2


def test_the_shared_value_rule_keeps_short_marks_and_the_header() -> None:
    assert deck_type.shared_value(_ROWS[1], 3) == "40 prompts, 3 raters"
    assert deck_type.shared_value(["Offline", "✓", "✓"], 3) is None
    assert deck_type.shared_value(["Offline", "N/A", "N/A"], 3) is None
    assert deck_type.shared_value(["Plan", "same", "same", ""], 4) is None
    assert deck_type.shared_value(["Plan", "Shared", "Shared"], 4) is None
    # one value column is not a span
    assert deck_type.shared_value(["Plan", "Shared"], 2) is None
    # the spanned text counts half in each of its two columns, not whole in each
    text = "40 document-drafting prompts, three human raters"
    halves = [deck_type.em("Tasks"), deck_type.em(text) / 2, deck_type.em(text) / 2]
    weights = [
        max(3.0, min(22.0, max(deck_type.em(head), half))) + 2.0
        for head, half in zip(_ROWS[0], halves, strict=True)
    ]
    shares = deck_type.column_shares([_ROWS[0], ["Tasks", text, text]])
    assert shares == [w / sum(weights) for w in weights]


def test_a_split_shared_item_comes_back_as_one_spanned_cell() -> None:
    old = [{"layout": "bands", "title": "Setup",
            "bands": [["Tasks", "40 document-drafting prompts, three human raters"]]}]
    new = [{"layout": "table", "title": "Setup",
            "rows": [["Metric", "On-Prem LLM", "Cloud API"], ["Tasks", "40 prompts", "3 raters"]]}]
    rows = revise.spread_rows(new, old)[0]["rows"]
    table = _table(deck_export.to_pptx("Deck", [{**new[0], "rows": rows}]))
    assert table.cell(1, 1).is_merge_origin and table.cell(1, 1).text == "40 prompts, 3 raters"
