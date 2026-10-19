"""Phases with their own tasks and two flows side by side are drawn to their own shape, and
the slots keep only what the figure's description says."""

from __future__ import annotations

import io
import json

import pytest
from pptx import Presentation

from app.services import diagram, diagram_render, figure_kinds

_STAGES = {"kind": "stages", "stages": [
    {"name": "데이터 수집", "items": ["공개 데이터셋 3종", "자체 수집 문서"]},
    {"name": "전처리", "items": ["중복 제거", "라벨 검수"]},
    {"name": "평가", "items": ["자동 지표 평가", "사람 평가"]}],
    "back": {"from": "평가", "to": "라벨 검수", "note": "오류 유형 → 검수 기준"}}
_SAID = ("1) 데이터 수집: 공개 데이터셋 3종과 자체 수집 문서. 2) 전처리: 중복 제거, 라벨 검수. "
         "3) 평가: 자동 지표 평가와 사람 평가. 평가에서 발견된 오류 유형은 라벨 검수 기준에 "
         "되먹임된다.")
_LANES = {"kind": "lanes", "lanes": [
    {"name": "색인", "steps": ["문서 업로드", "청크 분할", "임베딩"]},
    {"name": "질의", "steps": ["질문", "질문 임베딩", "벡터 검색", "답변 생성"]}],
    "store": {"name": "벡터 DB", "from": "임베딩", "to": "벡터 검색", "note": "조회"}}


def _texts(blob: bytes) -> str:
    deck = Presentation(io.BytesIO(blob))
    return " ".join(shape.text_frame.text for slide in deck.slides for shape in slide.shapes
                    if shape.has_text_frame)


def test_stages_keep_their_tasks_and_an_arrow_back_to_an_earlier_stage():
    checked = figure_kinds.spec(_STAGES)
    assert [len(st["items"]) for st in checked["stages"]] == [2, 2, 2]
    assert checked["back"] == {"from": "평가", "to": "라벨 검수", "note": "오류 유형 → 검수 기준"}
    # an arrow forward, or to a name that is not there, is not feedback
    forward = {**_STAGES, "back": {"from": "중복 제거", "to": "평가"}}
    assert "back" not in figure_kinds.spec(forward)
    assert "back" not in figure_kinds.spec({**_STAGES, "back": {"from": "평가", "to": "배포"}})


def test_phases_without_their_tasks_are_a_line_of_steps():
    bare = {"kind": "stages", "stages": [{"name": n} for n in ("진단", "정의", "시범", "확산")]}
    checked = figure_kinds.spec(bare)
    assert checked["kind"] == "steps" and [s["name"] for s in checked["steps"]][0] == "진단"


def test_tasks_and_feedback_the_description_never_gives_are_dropped():
    said = "신규 협업 도구 도입 4단계: 현황 진단, 요구사항 정의, 시범 운영, 전사 확산."
    made = {"kind": "stages", "lang": "ko", "stages": [
        {"name": "현황 진단", "items": ["현행 도구 분석"]},
        {"name": "요구사항 정의", "items": ["KPI 측정", "사용자 인터뷰"]},
        {"name": "시범 운영", "items": ["파일럿 팀 선정"]},
        {"name": "전사 확산", "items": ["교육 실시"]}],
        "back": {"from": "전사 확산", "to": "현황 진단"}}
    checked = figure_kinds.grounded(figure_kinds.spec(made), "", said)
    assert checked["kind"] == "steps" and "back" not in checked
    # tasks the description names stay, give or take a word the writer added
    kept = figure_kinds.grounded(figure_kinds.spec({**_STAGES, "stages": [
        {"name": "데이터 수집", "items": ["공개 데이터셋 3종 활용"]}, *_STAGES["stages"][1:]]}),
        "", _SAID)
    assert kept["kind"] == "stages" and kept["stages"][0]["items"] == ["공개 데이터셋 3종 활용"]
    assert kept["back"]["from"] == "평가"


def test_lanes_share_a_store_only_between_a_writer_and_a_reader():
    checked = figure_kinds.spec(_LANES)
    assert checked["store"] == _LANES["store"]
    backwards = {**_LANES, "store": {"name": "벡터 DB", "from": "벡터 검색", "to": "임베딩"}}
    assert "store" not in figure_kinds.spec(backwards)
    assert figure_kinds.spec({**_LANES, "lanes": _LANES["lanes"][:1]}) is None
    # the store sits under the step that writes it and over the one that reads it
    first, second, _ = figure_kinds._lanes_plan(checked)
    assert first + 2 == second + 2


def test_a_comparison_of_three_is_not_cut_to_two():
    three = {"kind": "compare", "sides": ["외부 API", "자체 호스팅", "하이브리드"], "rows": [
        {"aspect": "비용", "before": "낮음", "after": "높음"},
        {"aspect": "데이터", "before": "외부", "after": "내부"}]}
    assert figure_kinds.spec(three) is None


@pytest.mark.parametrize("slots", [_STAGES, _LANES])
def test_each_kind_is_drawn_with_its_names_and_keeps_its_slots(slots):
    checked = figure_kinds.spec(slots)
    source = figure_kinds.to_source(checked)
    assert figure_kinds.from_source(source) == checked
    text = _texts(diagram_render.figure_deck(source, "#7A1F3D"))
    words = {"stages": ["데이터 수집", "공개 데이터셋 3종", "라벨 검수", "오류 유형 → 검수 기준"],
             "lanes": ["색인", "질의", "벡터 DB", "벡터 검색", "조회"]}[checked["kind"]]
    for word in words:
        assert word in text


def _answers(monkeypatch, answers: list[str]) -> list[list[dict]]:
    calls: list[list[dict]] = []

    async def complete(_model, messages, _key, max_tokens):
        calls.append(messages)
        return answers[len(calls) - 1], {"inputTokens": 3, "outputTokens": 2}

    monkeypatch.setattr(diagram, "_complete", complete)
    return calls


@pytest.mark.asyncio
async def test_stages_without_tasks_in_the_description_are_drawn_as_steps(monkeypatch):
    made = {"kind": "stages", "stages": [{"name": n, "items": ["지어낸 작업"]}
                                         for n in ("진단", "정의", "시범")]}
    _answers(monkeypatch, ["stages", "```json\n" + json.dumps(made, ensure_ascii=False) + "\n```"])
    source, _caption, _usage = await diagram.draw(description="도입 3단계: 진단, 정의, 시범.",
                                                  figure="flow", model="m", api_key="k")
    assert (figure_kinds.from_source(source) or {}).get("kind") == "steps"


@pytest.mark.asyncio
async def test_a_method_figure_is_never_a_talk_between_parties(monkeypatch):
    calls = _answers(monkeypatch, ["messages", "```mermaid\nflowchart LR\n  a --> b\n```"])
    source, _caption, _usage = await diagram.draw(description="교사와 학생 모델", figure="method",
                                                  model="m", api_key="k")
    assert figure_kinds.from_source(source) is None and len(calls) == 2
