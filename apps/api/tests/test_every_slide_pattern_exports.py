"""Every slide pattern leaves as a file, its own words drawn into the .pptx and the .pdf."""

from __future__ import annotations

import io
import json

import pytest
from pptx import Presentation

from app.services import deck_export, slide_patterns


def _slide(pattern: slide_patterns.Pattern) -> dict:
    example = json.loads(pattern.example.replace("{{", "{").replace("}}", "}"))
    return {"title": pattern.label, "layout": pattern.name, "accent": "#3b5bdb", **example}


def _words(slide: dict) -> list[str]:
    """The strings a pattern slide must show, by its shape."""
    out: list[str] = []
    for left, right in slide.get("items") or []:
        out += [left, right]
    for column in slide.get("columns") or []:
        out += [column["title"], *column["items"]]
    for value, label in slide.get("metrics") or []:
        out += [value, label]
    if slide.get("body"):
        out.append(slide["body"])
    return [word for word in out if word]


def _slides() -> list[dict]:
    return [_slide(pattern) for pattern in slide_patterns.PATTERNS]


def _text_of(shapes) -> str:
    return "\n".join(shape.text_frame.text for shape in shapes if shape.has_text_frame)


@pytest.mark.parametrize("look", ["editorial", "dark", "warm", "mono"])
def test_every_pattern_draws_its_words_into_the_pptx(look: str) -> None:
    slides = _slides()
    blob = deck_export.to_pptx("패턴", slides, tokens={"visualStyle": look})
    assert blob
    deck = Presentation(io.BytesIO(blob))
    assert len(deck.slides) == len(slide_patterns.PATTERNS)
    for pattern, slide, drawn in zip(slide_patterns.PATTERNS, slides, deck.slides, strict=True):
        if pattern.shape == "chart":
            frames = [shape for shape in drawn.shapes if shape.has_chart]
            assert frames, f"{pattern.name}: 차트가 없습니다"
            plot = frames[0].chart.plots[0]
            assert list(plot.categories) and len(list(plot.categories)) == len(
                slide["chart"]["categories"]
            ), pattern.name
            continue
        text = _text_of(drawn.shapes)
        for word in _words(slide):
            assert word in text, f"{pattern.name}: {word!r} 가 슬라이드에 없습니다"


def test_chart_patterns_keep_their_kind() -> None:
    from pptx.enum.chart import XL_CHART_TYPE

    expected = {
        "chart-pie": XL_CHART_TYPE.PIE,
        "chart-donut": XL_CHART_TYPE.DOUGHNUT,
        "chart-hbar": XL_CHART_TYPE.BAR_CLUSTERED,
        "chart-stacked": XL_CHART_TYPE.COLUMN_STACKED,
    }
    slides = [_slide(slide_patterns.BY_NAME[name]) for name in expected]
    # A chart without its kind takes the pattern's.
    del slides[0]["chart"]["kind"]
    deck = Presentation(io.BytesIO(deck_export.to_pptx("차트", slides)))
    for name, drawn in zip(expected, deck.slides, strict=True):
        frame = next(shape for shape in drawn.shapes if shape.has_chart)
        assert frame.chart.chart_type == expected[name], name


def test_every_pattern_leaves_as_a_pdf() -> None:
    blob = deck_export.to_pdf("패턴", _slides(), tokens={"visualStyle": "editorial"})
    assert blob.startswith(b"%PDF")
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(blob))
    assert len(reader.pages) == len(slide_patterns.PATTERNS)


def test_a_pattern_slide_without_its_field_falls_back_to_the_body() -> None:
    slide = {"title": "빈 카드", "layout": "cards-3", "body": "아직 카드가 없습니다"}
    deck = Presentation(io.BytesIO(deck_export.to_pptx("빈", [slide])))
    assert "아직 카드가 없습니다" in _text_of(deck.slides[0].shapes)
    assert deck_export.to_pdf("빈", [slide]).startswith(b"%PDF")
