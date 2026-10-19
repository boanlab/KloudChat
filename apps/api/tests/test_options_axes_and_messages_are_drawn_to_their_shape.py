"""Three or more options, a two-axis map and messages between parties are drawn to their own
shape from slots, like steps and comparisons."""

from __future__ import annotations

import io
import json

import pytest
from pptx import Presentation

from app.services import diagram, diagram_render, diagram_shapes, figure_kinds

_OPTIONS = {"kind": "options", "options": ["외부 API", "자체 호스팅", "하이브리드"], "rows": [
    {"aspect": "초기 비용", "cells": [{"text": "낮음", "mark": "good"},
                                   {"text": "GPU 서버 비용 높음", "mark": "bad"}, "중간"]},
    {"aspect": "데이터", "cells": ["외부로 나감", "내부에 머묾", "민감 업무만 내부"]}],
    "hot": "하이브리드"}
_AXES = {"kind": "axes", "x": {"name": "가격", "low": "낮음", "high": "높음"},
         "y": {"name": "기능 범위", "low": "좁음", "high": "넓음"},
         "items": [{"name": "메신저", "x": "low", "y": "low"},
                   {"name": "스위트", "x": "high", "y": "high"},
                   {"name": "우리 제품", "x": "mid", "y": "high"}], "hot": "우리 제품"}
_MESSAGES = {"kind": "messages", "parties": ["클라이언트", "서버"], "messages": [
    {"from": "클라이언트", "to": "서버", "name": "ClientHello", "note": "지원 암호 목록"},
    {"from": "서버", "to": "클라이언트", "name": "ServerHello"}]}


def _texts(blob: bytes) -> str:
    deck = Presentation(io.BytesIO(blob))
    return " ".join(shape.text_frame.text for slide in deck.slides for shape in slide.shapes
                    if shape.has_text_frame)


def test_options_keep_one_cell_per_option():
    checked = figure_kinds.spec(_OPTIONS)
    assert [len(row["cells"]) for row in checked["rows"]] == [3, 3]
    assert checked["rows"][0]["cells"][0] == {"text": "낮음", "mark": "good"}
    # a mark the text does not make is not kept, and a short row is padded with empty cells
    loose = {**_OPTIONS, "rows": [{"aspect": "a", "cells": [{"text": "x", "mark": "best"}, "y"]},
                                  *_OPTIONS["rows"]]}
    assert figure_kinds.spec(loose)["rows"][0]["cells"] == [{"text": "x"}, {"text": "y"},
                                                            {"text": ""}]


def test_slots_that_do_not_make_these_kinds_are_refused():
    # two options are a two-sided comparison
    assert figure_kinds.spec({**_OPTIONS, "options": ["가", "나"]}) is None
    # a map with fewer than three placed items
    assert figure_kinds.spec({**_AXES, "items": _AXES["items"][:2]}) is None
    # a message to a party that is not there, or to itself, is left out
    stray = {**_MESSAGES, "messages": [*_MESSAGES["messages"][:1],
                                       {"from": "서버", "to": "DB", "name": "조회"},
                                       {"from": "서버", "to": "서버", "name": "검증"}]}
    assert figure_kinds.spec(stray) is None


def test_an_item_the_text_does_not_place_is_left_off_the_map():
    unplaced = {**_AXES, "items": [*_AXES["items"], {"name": "새 도구", "x": "low"}]}
    assert [i["name"] for i in figure_kinds.spec(unplaced)["items"]] == ["메신저", "스위트",
                                                                       "우리 제품"]
    # a highlight names an item that is on the map
    assert "hot" not in figure_kinds.spec({**_AXES, "hot": "없는 제품"})


@pytest.mark.parametrize("slots", [_OPTIONS, _AXES, _MESSAGES])
def test_the_stored_source_carries_the_slots_and_shows_every_word(slots):
    checked = figure_kinds.spec(slots)
    source = figure_kinds.to_source(checked)
    assert figure_kinds.from_source(source) == checked
    # the fallback flowchart names every part, so the fact check reads them
    graph = diagram_shapes.parse(source)
    labels = " ".join([*(label for label, _ in graph.groups.values()),
                       *(n.label for n in graph.nodes.values())])
    for word in ("하이브리드", "GPU 서버 비용 높음", "우리 제품", "ServerHello", "지원 암호 목록"):
        if word in json.dumps(checked, ensure_ascii=False):
            assert word in labels


@pytest.mark.parametrize("slots", [_OPTIONS, _AXES, _MESSAGES])
def test_each_kind_is_drawn_with_its_names(slots):
    checked = figure_kinds.spec(slots)
    text = _texts(diagram_render.figure_deck(figure_kinds.to_source(checked), "#7A1F3D"))
    words = {"options": ["외부 API", "하이브리드", "초기 비용", "GPU 서버 비용 높음", "유리"],
             "axes": ["가격", "넓음", "메신저", "우리 제품"],
             "messages": ["클라이언트", "서버", "ClientHello", "지원 암호 목록"]}[checked["kind"]]
    for word in words:
        assert word in text


def _answers(monkeypatch, answers: list[str]) -> list[list[dict]]:
    calls: list[list[dict]] = []

    async def complete(_model, messages, _key, max_tokens):
        calls.append(messages)
        return answers[len(calls) - 1], {"inputTokens": 3, "outputTokens": 2}

    monkeypatch.setattr(diagram, "_complete", complete)
    return calls


@pytest.mark.asyncio
@pytest.mark.parametrize("label, slots", [("options", _OPTIONS), ("axes", _AXES),
                                          ("messages", _MESSAGES)])
async def test_the_shape_choice_writes_these_kinds_as_slots(monkeypatch, label, slots):
    filled = "```json\n" + json.dumps(slots, ensure_ascii=False) + "\n```\n그림: 비교"
    calls = _answers(monkeypatch, [label, filled])
    source, caption, _usage = await diagram.draw(description="…", figure="flow", model="m",
                                                 api_key="k")
    assert (figure_kinds.from_source(source) or {}).get("kind") == label and caption == "비교"
    assert len(calls) == 2 and f'"kind": "{label}"' in calls[1][0]["content"]
