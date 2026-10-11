"""A long flowchart is drawn in rows whose text reads at the size the figure is printed."""

from pptx import Presentation
from pptx.util import Pt

from app.services import diagram_render, flow_layout
from app.services import diagram_shapes as ds

#: Figure 1 of the persona sweep (R-12): eleven steps that came out 2449 × 425 on one line.
FIG1 = """flowchart LR
  sum[요약문] --> dec(원자적 주장 분해) --> list[주장 목록]
  list --> num(수치 정합성 검사기)
  list --> ret(원문 문장 검색)
  ret --> nli(NLI 함의 판정)
  nli --> gate{불확실성 게이트}
  gate -.->|불확실| ret
  gate --> conf(충돌 확인)
  conf -.->|충돌 시 NEI| score(주장별 점수)
  num --> score
  score --> avg(가중 평균) --> out[전체 사실 일관성]:::hot"""

_WORDS = [
    "요구 사항", "범위 정의", "이해관계자 면담", "현황 분석", "문제 정의", "대안 도출",
    "평가 기준 수립", "대안 평가", "상세 설계", "시범 적용", "효과 측정", "전사 확산",
    "운영 이관", "교육", "성과 보고",
]


def chain(n: int, extra: str = "") -> str:
    steps = " --> ".join(f"s{i}({_WORDS[i]})" for i in range(n))
    return f"flowchart LR\n  {steps}\n  s{n - 1}:::hot\n{extra}"


def _boxes(placed):
    return {n: (p.x - p.w / 2, p.y - p.h / 2, p.x + p.w / 2, p.y + p.h / 2)
            for n, p in placed.items()}


def _no_overlap(placed) -> bool:
    boxes = list(_boxes(placed).values())
    return all(
        a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1]
        for i, a in enumerate(boxes) for b in boxes[i + 1:]
    )


def _readable(graph, plan) -> None:
    """Every box at its floor: 10 pt, a decision 9 pt, in a frame printed at its own size."""
    for nid, p in plan.placed.items():
        node = graph.nodes[nid]
        size = ds._font_size(node.label, *ds.text_area(node, p.w, p.h, plan.icons)[:2])
        assert size >= (flow_layout.DECISION_PT if node.shape == "diamond" else flow_layout.BOX_PT)


def test_figure_1_folds_into_rows_that_read_at_the_a4_text_width():
    graph = ds.parse(FIG1)
    plan = flow_layout.plan(graph, flow_layout.DOCUMENT_WIDTH)
    assert len(plan.rows) == 3 and max(len(row) for row in plan.rows) == 4
    assert set(plan.placed) == set(graph.nodes)
    _readable(graph, plan)
    assert _no_overlap(plan.placed)
    # Printed at the text column (≈ 425 pt DOCX, 466 pt PDF) it is not shrunk: its own size.
    width, height = diagram_render.figure_size(graph)
    assert width <= 425 and 0.8 <= width / height <= 3
    # The step back (불확실 → 원문 문장 검색) is the one cut from the loop, not a solid edge.
    back = [graph.edges[i] for i in plan.back]
    assert [(e.a, e.b) for e in back] == [("gate", "ret")]


def test_no_line_of_figure_1_runs_across_the_figure():
    graph = ds.parse(FIG1)
    placed = flow_layout.plan(graph, flow_layout.DOCUMENT_WIDTH).placed
    for edge in graph.edges:
        a, b = placed[edge.a], placed[edge.b]
        how = flow_layout.kind(a, b)
        if how == "wrap":  # a row's end to the next row's start
            assert b.row == a.row + 1 and (a.last and b.col == 0 or a.col == b.col)
        elif how in ("forward", "arc", "lane"):
            assert a.row == b.row
        else:
            assert how == "mark"
    kinds = {(e.a, e.b): flow_layout.kind(placed[e.a], placed[e.b]) for e in graph.edges}
    assert kinds[("gate", "ret")] == "mark" and kinds[("num", "score")] == "mark"


def test_every_edge_is_drawn_as_a_line_or_a_note_naming_where_it_goes():
    graph = ds.parse(FIG1)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(500), Pt(500)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    width, height = diagram_render.figure_size(graph)
    ds.draw_pptx(slide, graph, left=18, top=18, width=width, height=height, accent="#2563EB")
    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert "불확실 → 원문 문장 검색" in texts and "→ 주장별 점수" in texts
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    lines = [s for s in slide.shapes if s._element.find(f".//{a}custGeom") is not None]
    marks = [t for t in texts if "→" in t]
    # A routed line per edge, and a stub plus a note per marked edge.
    assert len(lines) == len(graph.edges) and len(marks) == 2
    assert all(graph.nodes[n].label in texts for n in graph.nodes)
    # Boxes alike: no box gets an icon that would make its text smaller than the others'.
    assert not [p for p in prs.part.package.iter_parts() if str(p.partname).endswith(".svg")]


def test_chains_of_5_8_11_15_steps_fold_into_rows_of_three_to_five():
    shapes = {}
    for n in (5, 8, 11, 15):
        graph = ds.parse(chain(n))
        plan = flow_layout.plan(graph, flow_layout.DOCUMENT_WIDTH)
        _readable(graph, plan)
        assert _no_overlap(plan.placed)
        lengths = [len(row) for row in plan.rows]
        assert max(lengths) <= 5 and min(lengths) >= 2 and max(lengths) >= 3
        # Rows read left to right, in order, each starting at the left.
        order = [n for row in plan.rows for col in row for n in col]
        assert order == [f"s{i}" for i in range(n)]
        assert all(plan.placed[row[0][0]].col == 0 for row in plan.rows)
        shapes[n] = lengths
    assert len(shapes[5]) <= 2 and len(shapes[8]) == 2 and len(shapes[15]) >= 3


def test_a_short_flow_stays_on_one_line():
    graph = ds.parse(
        "flowchart LR\n  a[원본 문서] --> b(문단 분할) --> c(임베딩 생성) --> d[(벡터 저장소)]"
        " --> e[검색 결과]:::hot"
    )
    plan = flow_layout.plan(graph, flow_layout.DOCUMENT_WIDTH)
    assert len(plan.rows) == 1


def test_a_feedback_loop_is_a_short_arc_or_a_marker_never_a_line_across():
    graph = ds.parse(chain(8, "  s6 -.->|근거 부족| s2\n  s7 -.->|재시도| s0"))
    plan = flow_layout.plan(graph, flow_layout.DOCUMENT_WIDTH)
    back = sorted((graph.edges[i].a, graph.edges[i].b) for i in plan.back)
    assert back == [("s6", "s2"), ("s7", "s0")]
    for i in plan.back:
        e = graph.edges[i]
        a, b = plan.placed[e.a], plan.placed[e.b]
        how = flow_layout.kind(a, b)
        assert how in ("arc", "mark")
        if how == "arc":
            assert a.row == b.row and abs(a.col - b.col) <= 3


def test_a_loop_of_solid_edges_is_cut_where_the_source_closes_it():
    graph = ds.parse("flowchart LR\n  a(가) --> b(나) --> c(다)\n  c --> a")
    plan = flow_layout.plan(graph, flow_layout.DOCUMENT_WIDTH)
    assert [(graph.edges[i].a, graph.edges[i].b) for i in plan.back] == [("c", "a")]


def test_phases_written_as_rows_keep_one_phase_a_row_with_a_frame_each():
    source = """flowchart TB
  subgraph prep["데이터 준비"]
    direction LR
    a[원천 데이터] --> b(정제) --> c(라벨링)
  end
  subgraph train["학습"]
    direction LR
    d(사전 학습) --> e(미세 조정) --> f(검증)
  end
  subgraph serve["배포"]
    direction LR
    g(경량화) --> h(서빙) --> i[모니터링 지표]:::hot
  end
  prep --> train --> serve
  f -.->|성능 미달| e"""
    graph = ds.parse(source)
    assert graph.group_dirs == {"prep": "LR", "train": "LR", "serve": "LR"}
    assert flow_layout.applies(graph)
    plan = flow_layout.plan(graph, flow_layout.DOCUMENT_WIDTH)
    assert [[c[0] for c in row] for row in plan.rows] == [
        ["a", "b", "c"], ["d", "e", "f"], ["g", "h", "i"]]
    frames = flow_layout.frames(graph, plan.placed)
    assert [label for label, _ in frames] == ["데이터 준비", "학습", "배포"]


def test_a_phase_longer_than_a_row_gets_a_frame_on_each_row_it_runs_over():
    steps = " --> ".join(f"p{i}({_WORDS[i]})" for i in range(9))
    source = f'flowchart LR\n  subgraph all["분석 절차"]\n    {steps}\n  end\n  q(결론)\n  p8 --> q'
    graph = ds.parse(source)
    plan = flow_layout.plan(graph, flow_layout.DOCUMENT_WIDTH)
    assert len(plan.rows) >= 2
    labels = [label for label, _ in flow_layout.frames(graph, plan.placed)]
    assert labels[0] == "분석 절차" and all(x == "분석 절차 (이어서)" for x in labels[1:])
    assert len(labels) == len({plan.placed[f"p{i}"].row for i in range(9)})


def test_comparisons_and_hierarchies_keep_their_own_layout():
    compare = ds.parse(
        "flowchart LR\n"
        '  subgraph old["기존"]\n    direction TB\n    a1[검색: 키워드] ~~~ a2[근거: 없음]\n  end\n'
        '  subgraph new["제안"]\n    direction TB\n    b1[검색: 의미] ~~~ b2[근거: 인용]\n  end\n'
    )
    tree = ds.parse(
        "flowchart TB\n  r[정보 보안] --> c1[기밀성] & c2[무결성] & c3[가용성]\n  c1 --> d1[암호화]"
    )
    for graph in (compare, tree):
        assert not flow_layout.applies(graph)
        assert all(p.row == -1 for p in ds.layout(graph, 816, 300).values())


def test_a_deck_band_holds_the_rows_inside_its_box():
    graph = ds.parse(FIG1)
    for width, height in ((816, 260), (816, 200), (400, 330)):
        placed = ds.layout(graph, width, height)
        assert set(placed) == set(graph.nodes) and _no_overlap(placed)
        for x0, y0, x1, y1 in _boxes(placed).values():
            assert 0 <= x0 and x1 <= width + 0.01 and 0 <= y0 and y1 <= height + 0.01
    # A wide band takes more steps a row than the document column.
    wide = flow_layout.plan(graph, 816, 260)
    assert len(wide.rows) < len(flow_layout.plan(graph, flow_layout.DOCUMENT_WIDTH).rows)


def test_text_is_sized_to_wrap_between_words_not_inside_one():
    # 「라벨링」 in a narrow box: a size at which the word fits whole, not 「라벨/링」.
    size = ds._font_size("라벨링", 45, 52)
    assert ds._ems("라벨링") * size <= 45 - 8
    assert ds._lines("이해관계자 면담", 5.0) == 2 and ds._lines("이해관계자", 4.0) is None
