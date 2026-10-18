"""A card row with little to say is drawn larger so its cards are not mostly empty."""

import io

from pptx import Presentation

from app.services import deck_export, deck_type, fonts

_SHORT = [["카페", "창가에 2~4인 좌석"], ["회의", "공용 테이블과 화이트보드"],
          ["휴식", "안쪽의 1인 좌석"]]
_MID = [["카페 좌석", "창가에 햇빛이 드는 2~4인 카페 좌석을 두어 점심 시간에 마주 앉아 대화할 수 "
                   "있게 하고, 창가 화분으로 햇빛을 가린다."],
        ["공용 테이블", "가운데에 큰 공용 테이블과 화이트보드를 두어 팀 회의를 할 수 있게 한다."],
        ["1인 좌석", "안쪽에 조용히 혼자 쉬는 1인 좌석을 둔다."]]
_LONG = [["가", "긴 설명 " * 60], ["나", "짧다"], ["다", "짧다"]]


def _style(pairs) -> str:
    font = fonts.korean("gothic")
    span = (816 - 36) / 3
    return deck_export._card_style([tuple(p) for p in pairs], font, 1.0, span, 100 * deck_type.K)


def test_cards_take_the_size_their_words_allow():
    assert _style(_SHORT) == "marked"
    assert _style(_MID) == "large"
    assert _style(_LONG) == ""
    # without a measuring font the regular cards are kept
    assert deck_export._card_style([("가", "나")], None, 1.0, 200, 240) == ""


def _sizes(cards) -> tuple[set[str], set[float]]:
    data = deck_export.to_pptx("덱", [{"layout": "title", "title": "덱"},
                                      {"layout": "cards", "title": "세 공간", "cards": cards}])
    slide = Presentation(io.BytesIO(data)).slides[1]
    texts, sizes = set(), set()
    for shape in slide.shapes:
        if shape.has_text_frame:
            for run in (r for p in shape.text_frame.paragraphs for r in p.runs):
                texts.add(run.text)
                if run.font.size:
                    sizes.add(run.font.size.pt)
    return texts, sizes


def test_short_cards_are_numbered_and_drawn_larger():
    texts, sizes = _sizes(_SHORT)
    assert {"01", "02", "03"} <= texts
    assert {deck_type.TYPE["cardNameShort"], deck_type.TYPE["cardTextShort"]} <= sizes


def test_long_cards_keep_the_regular_sizes():
    texts, sizes = _sizes(_LONG)
    assert "01" not in texts
    assert deck_type.TYPE["cardNameShort"] not in sizes
