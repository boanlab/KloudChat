"""Steps, cycles, houses and two-sided comparisons are drawn to their own shape from slots."""

from __future__ import annotations

import io
import json

import pytest
from pptx import Presentation

from app.services import (
    deck_export,
    diagram,
    diagram_render,
    diagram_shapes,
    figure_facts,
    figure_kinds,
    figure_text,
)

_STEPS = {"kind": "steps", "steps": ["현황 진단", "요구사항 정의", "시범 운영", "전사 확산"]}
_CYCLE = {"kind": "cycle", "center": "PDCA", "stages": [
    {"name": "계획", "en": "Plan", "note": "목표와 개선안 수립"}, {"name": "실행", "en": "Do"},
    {"name": "점검", "en": "Check"}, {"name": "개선", "en": "Act"}]}
_HOUSE = {"kind": "house", "goal": "신뢰할 수 있는 데이터 활용",
          "pillars": ["데이터 품질 관리", "메타데이터 관리", "보안·개인정보 보호"],
          "base": [{"name": "거버넌스 조직", "items": ["데이터 위원회", "데이터 관리자"]},
                   {"name": "정책 · 표준", "note": "모든 영역이 따르는 규칙"}]}
_COMPARE = {"kind": "compare", "sides": ["기존 · 수작업 보고", "자동화 보고"], "rows": [
    {"aspect": "수집", "before": "엑셀로 수집", "after": "매일 자동 수집"},
    {"aspect": "작성", "before": "작성에 3일", "after": "즉시 생성"}],
    "hot": {"row": "작성", "badge": "90% 단축"}}


def _texts(blob: bytes) -> str:
    deck = Presentation(io.BytesIO(blob))
    return " ".join(
        shape.text_frame.text for slide in deck.slides for shape in slide.shapes
        if shape.has_text_frame
    )


# ── slots ─────────────────────────────────────────────────────────────


def test_slots_that_do_not_make_their_kind_are_refused():
    assert figure_kinds.spec({"kind": "steps", "steps": ["하나", "둘"]}) is None
    assert figure_kinds.spec({"kind": "quadrant", "points": []}) is None
    assert figure_kinds.spec({"kind": "house", "goal": "", "pillars": ["a", "b"]}) is None
    # A comparison row with one side empty is not a two-sided comparison.
    half = {**_COMPARE, "rows": [*_COMPARE["rows"], {"aspect": "비용", "before": "높음"}]}
    assert figure_kinds.spec(half) is None
    # A dash in a cell is a cell with nothing in it.
    dashed = {**_COMPARE, "rows": [*_COMPARE["rows"], {"aspect": "비용", "before": "-",
                                                      "after": "낮음"}]}
    assert figure_kinds.spec(dashed) is None


def test_a_long_slot_is_cut_at_a_space():
    note = "가" * 30 + " " + "나" * 20
    steps = figure_kinds.spec(
        {"kind": "steps", "steps": ["하나", "둘", {"name": "셋", "note": note}]}
    )
    assert steps["steps"][2]["note"] == "가" * 30


def test_a_note_that_only_repeats_its_name_is_dropped():
    steps = figure_kinds.spec({"kind": "steps", "steps": [
        {"name": "현황 진단", "note": "진단"}, {"name": "시범 운영", "note": "일부 부서 먼저"},
        {"name": "전사 확산", "note": "전사 확산"}]})
    assert [item.get("note") for item in steps["steps"]] == [None, "일부 부서 먼저", None]


def test_a_highlight_names_a_row_that_is_there():
    found = figure_kinds.spec({**_COMPARE, "hot": {"row": "없는 행", "badge": "x"}})
    assert "hot" not in found
    assert figure_kinds.spec(_COMPARE)["hot"] == {"row": "작성", "badge": "90% 단축"}


def test_text_the_figure_already_shows_is_not_written_twice():
    # A badge made of the row's own cells says nothing new.
    again = figure_kinds.spec({**_COMPARE, "hot": {"row": "작성", "badge": "3일 → 즉시"}})
    assert again["hot"] == {"row": "작성"}
    cycle = figure_kinds.spec({**_CYCLE, "center_note": "PDCA"})
    assert "center_note" not in cycle


def test_the_language_names_the_step_numbers():
    deck = diagram_render.figure_deck(figure_kinds.to_source(figure_kinds.spec(_STEPS, "en")))
    assert "Step 1" in _texts(deck) and "1단계" not in _texts(deck)


# ── stored source ─────────────────────────────────────────────────────


@pytest.mark.parametrize("slots", [_STEPS, _CYCLE, _HOUSE, _COMPARE])
def test_the_stored_source_is_a_flowchart_carrying_its_slots(slots):
    checked = figure_kinds.spec(slots)
    source = figure_kinds.to_source(checked)
    assert figure_kinds.from_source(source) == checked
    # Views that read mermaid still read a flowchart with every name in it.
    graph = diagram_shapes.parse(source)
    assert graph is not None
    labels = " ".join(node.label for node in graph.nodes.values())
    for name in [*(s if isinstance(s, str) else s["name"]
                   for s in slots.get("steps") or slots.get("stages") or [])]:
        assert name in labels
    assert source.splitlines()[-1].startswith(figure_kinds.MARK)


def test_a_plain_flowchart_carries_no_slots():
    assert figure_kinds.from_source("flowchart LR\n  a(수집) --> b(분석)") is None
    assert figure_kinds.from_source("flowchart LR\n%% kloudchat-figure: {broken") is None


def test_the_fact_check_taking_the_highlight_off_takes_it_off_the_slots():
    source = figure_kinds.to_source(figure_kinds.spec(_COMPARE))
    plain = source.replace(":::hot", "")
    assert "hot" in figure_kinds.from_source(source)
    assert "hot" not in figure_kinds.from_source(plain)
    # The fact check reads the comparison's cells as 「항목: 값」 nodes.
    assert figure_facts.check(source, context="").problems == []


# ── drawing ───────────────────────────────────────────────────────────


@pytest.mark.parametrize("slots", [_STEPS, _CYCLE, _HOUSE, _COMPARE])
def test_each_kind_is_drawn_with_its_names(slots):
    checked = figure_kinds.spec(slots)
    blob = diagram_render.figure_deck(figure_kinds.to_source(checked), "#7A1F3D")
    deck = Presentation(io.BytesIO(blob))
    width, height = figure_kinds.size(checked)
    # The slide is the figure's own size plus its margin.
    assert abs(deck.slide_width.pt - (width + 36)) < 1
    assert abs(deck.slide_height.pt - (height + 36)) < 1
    text = _texts(blob).replace("\n", " ")
    words = json.dumps(slots, ensure_ascii=False)
    for name in ("현황 진단", "계획", "신뢰할 수 있는 데이터 활용", "자동화 보고", "90% 단축"):
        if name in words:
            assert name in text


def test_a_cycle_stage_without_a_note_has_no_callout():
    blob = diagram_render.figure_deck(figure_kinds.to_source(figure_kinds.spec(_CYCLE)))
    text = _texts(blob)
    # Each stage is named once, on the ring; the one note stands beside it alone.
    assert text.count("계획") == 1 and text.count("실행") == 1
    assert "목표와 개선안 수립" in text


def test_each_part_of_a_cycle_takes_white_text():
    for colour in figure_kinds.hues("#2563EB", 6) + figure_kinds.hues("#F5C518", 6):
        c = colour.lstrip("#")
        assert figure_kinds._luminance(*(int(c[i:i + 2], 16) / 255 for i in (0, 2, 4))) <= 0.183


def test_a_label_breaks_at_spaces_before_it_breaks_a_word():
    lines, _size = figure_text.fit("데이터 생명주기 관리", 120, 60, 16, True)
    assert len(lines) == 2 and " ".join(lines) == "데이터 생명주기 관리"
    # A word wider than any line is the only thing broken inside.
    lines, size = figure_text.fit("CertificateVerify", 60, 40, 14)
    assert len(lines) > 1 and size == 9.0


def test_a_deck_draws_its_fixed_figure_as_shapes():
    import PIL.Image

    picture = io.BytesIO()
    PIL.Image.new("RGB", (640, 360), "white").save(picture, format="PNG")
    src = "data:image/png;base64," + __import__("base64").b64encode(picture.getvalue()).decode()
    source = figure_kinds.to_source(figure_kinds.spec(_CYCLE))
    slides = [{"id": "s1", "title": "품질 개선", "layout": "bullets", "bullets": [],
               "diagram": {"source": source, "caption": "PDCA"},
               "image": {"src": src, "diagram": True, "caption": "PDCA"}}]
    text = _texts(deck_export.to_pptx("덱", slides))
    assert "PDCA" in text and "계획" in text and "Check" in text


# ── writing the figure ────────────────────────────────────────────────


def _answers(monkeypatch, answers: list[str]) -> list[list[dict]]:
    calls: list[list[dict]] = []

    async def complete(_model, messages, _key, max_tokens):
        calls.append(messages)
        return answers[len(calls) - 1], {"inputTokens": 3, "outputTokens": 2}

    monkeypatch.setattr(diagram, "_complete", complete)
    return calls


_FILLED = "```json\n" + json.dumps(_CYCLE, ensure_ascii=False) + "\n```\n그림: PDCA 순환"
_MERMAID = "```mermaid\nflowchart LR\n  a(수집) --> b(분석)\n```\n그림: 분석 흐름"


@pytest.mark.asyncio
async def test_a_figure_with_a_fixed_shape_is_written_as_its_slots(monkeypatch):
    calls = _answers(monkeypatch, ["cycle", _FILLED])
    source, caption, usage = await diagram.draw(
        description="계획, 실행, 점검, 개선을 되풀이한다", figure="concept", model="m", api_key="k"
    )
    assert figure_kinds.from_source(source)["kind"] == "cycle" and caption == "PDCA 순환"
    assert len(calls) == 2 and usage == {"inputTokens": 6, "outputTokens": 4}
    # The fill asks for that one kind only.
    assert '"kind": "cycle"' in calls[1][0]["content"] and '"kind": "steps"' not in calls[1][0][
        "content"]


@pytest.mark.asyncio
async def test_any_other_shape_is_written_as_mermaid(monkeypatch):
    calls = _answers(monkeypatch, ["messages", _MERMAID])
    source, _caption, usage = await diagram.draw(
        description="클라이언트와 서버가 메시지를 주고받는다", figure="flow", model="m", api_key="k"
    )
    assert figure_kinds.from_source(source) is None and source.startswith("flowchart LR")
    # The mermaid prompt is the one it always was.
    assert calls[1] == diagram._messages("클라이언트와 서버가 메시지를 주고받는다", "flow", "ko")
    assert usage == {"inputTokens": 6, "outputTokens": 4}


@pytest.mark.asyncio
async def test_slots_that_do_not_make_the_kind_fall_back_to_mermaid(monkeypatch):
    calls = _answers(monkeypatch, ["steps", '```json\n{"kind": "steps", "steps": ["하나"]}\n```',
                                   _MERMAID])
    source, _caption, usage = await diagram.draw(
        description="수집해 분석한다", figure="flow", model="m", api_key="k"
    )
    assert figure_kinds.from_source(source) is None and len(calls) == 3
    assert usage == {"inputTokens": 9, "outputTokens": 6}


@pytest.mark.asyncio
async def test_a_method_figure_is_never_a_line_of_steps(monkeypatch):
    calls = _answers(monkeypatch, ["steps", _MERMAID])
    source, _caption, _usage = await diagram.draw(
        description="인코더와 디코더로 이루어진 모델", figure="method", model="m", api_key="k"
    )
    assert figure_kinds.from_source(source) is None and len(calls) == 2


@pytest.mark.asyncio
async def test_a_rewrite_of_a_broken_flowchart_skips_the_shape_choice(monkeypatch):
    calls = _answers(monkeypatch, [_MERMAID])
    await diagram.draw(description="…", figure="flow", model="m", api_key="k",
                       broken="flowchart LR\n  a(", error="문법 오류")
    assert len(calls) == 1


# ── a slide's figure strip ────────────────────────────────────────────


@pytest.mark.parametrize("slots", [_CYCLE, _HOUSE, _COMPARE])
def test_a_figure_for_a_slide_is_drawn_as_a_wide_strip(slots):
    full = figure_kinds.spec(slots)
    strip = figure_kinds.spec({**slots, "band": True})
    assert strip["band"] is True and "band" not in full
    width, height = figure_kinds.size(strip)
    # The strip fits a slide's figure band (about 5:1), lower than the full figure.
    assert width / height >= 3.2 and height < figure_kinds.size(full)[1]
    assert figure_kinds.from_source(figure_kinds.to_source(strip)) == strip


def test_a_wide_box_takes_the_strip_even_for_a_full_figure():
    from pptx.util import Pt

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    figure_kinds.draw(slide, figure_kinds.spec(_HOUSE), left=72, top=100, width=816,
                      height=160, accent="#0F766E")
    shapes = list(slide.shapes)
    # Drawn across most of the band, not shrunk into its middle (the full house would span
    # about a third of it), and inside the band.
    span = max(s.left + s.width for s in shapes) - min(s.left for s in shapes)
    assert span > Pt(600)
    assert min(s.top for s in shapes) >= Pt(99) and max(s.top + s.height for s in shapes) <= Pt(261)


@pytest.mark.asyncio
async def test_a_slide_figure_is_written_for_the_strip(monkeypatch):
    _answers(monkeypatch, ["cycle", _FILLED])
    source, _caption, _usage = await diagram.draw(
        description="계획, 실행, 점검, 개선을 되풀이한다", figure="concept", model="m",
        api_key="k", slide=True,
    )
    assert figure_kinds.from_source(source)["band"] is True


@pytest.mark.asyncio
async def test_the_slots_take_their_values_from_the_text_the_figure_goes_with(monkeypatch):
    from app.services import diagrams

    calls = _answers(monkeypatch, ["compare", "```json\n" + json.dumps(_COMPARE) + "\n```"])

    async def render_png(source, accent="#2563EB"):
        return b"png"

    monkeypatch.setattr(diagrams.diagram_render, "render_png", render_png)
    monkeypatch.setattr(diagram.diagram_render, "render_png", render_png)
    planned = diagrams.Planned(0, "compare", "기존과 자동화를 수집·작성으로 비교", "보고 방식",
                               context="수집: 엑셀 수작업 → 자동 수집\n작성: 3일 → 즉시")
    made, _spent, _png = await diagrams.make(planned, model="m", api_key="k", slide=True)
    assert figure_kinds.from_source(made["source"])["kind"] == "compare"
    # The shape choice reads the description; the fill also reads the slide's own text.
    assert "엑셀 수작업" not in calls[0][-1]["content"]
    assert "그림이 들어갈 글:\n수집: 엑셀 수작업 → 자동 수집" in calls[1][-1]["content"]
